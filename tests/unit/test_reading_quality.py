"""What a model reads is kept: one wrong item costs only itself, and a quote is cited where it is.

The reported reading: a long Word specification lost seven parts to "unusable"
answers (one item over a limit rejected the whole answer), six to "uncited" ones,
and 240 suggestions whose quotes a check too strict about typography refused.
"""

from __future__ import annotations

import io
import zipfile
from dataclasses import replace
from datetime import UTC, datetime
from types import SimpleNamespace
from typing import Any

from openpyxl import Workbook
from smb_kernel.documents.text_extractor import SafeDocumentTextExtractor

from knowledge_portal.application.ports.catalogue_extractor import (
    CatalogueProposal,
    ExtractionRequest,
    ExtractionSegment,
    ProposedChange,
)
from knowledge_portal.application.use_cases.catalogue_candidates import _same_workflow
from knowledge_portal.domain.architecture.candidates import (
    CandidateBasis,
    CandidateCitation,
    CandidateContent,
    CandidateKind,
    CandidateMatch,
    CatalogueCandidate,
    apply_candidate,
    classify,
    missing_systems,
    open_matches,
)
from knowledge_portal.domain.architecture.journeys import Activity, Journey
from knowledge_portal.domain.architecture.knowledge import (
    KnowledgeReleaseStatus,
    SystemDefinition,
)
from knowledge_portal.domain.architecture.products import (
    OfferingPlan,
    OrderType,
    PlanCharacteristic,
    ProductOffering,
)
from knowledge_portal.infrastructure.architecture.catalogue_tables import CatalogueTableReader
from knowledge_portal.infrastructure.architecture.knowledge_yaml import seed_knowledge
from knowledge_portal.infrastructure.architecture.located_extractor import (
    LocatedDocumentExtractor,
    _without_hidden_sheets,
)
from knowledge_portal.infrastructure.llm.catalogue_extraction import (
    StructuredCatalogueExtractor,
    slug,
)
from knowledge_portal.infrastructure.llm.prompts.catalogue_extraction_prompt import (
    DETAILED_SYSTEM_PROMPT,
    LEAN_SYSTEM_PROMPT,
)
from tests.unit.test_document_reading_coverage import (
    MAINTAINER,
    _docx,
    _p,
    _Recording,
    _setup,
)

DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def _item(**values: Any) -> dict[str, Any]:
    """An answered item as the model writes it: JSON, not yet checked."""
    base: dict[str, Any] = {
        "kind": "system",
        "system": "Order Hub",
        "name": "Order Hub",
        "name_ar": None,
        "aliases": [],
        "triggers": [],
        "target_system": None,
        "component": None,
        "technology": None,
        "domain": None,
        "parent_domain": None,
        "offering": None,
        "journey": None,
        "text": None,
        "evidence_numbers": [1],
        "quote": "The Order Hub receives partner orders",
        "basis": "stated",
        "reasoning": None,
        "relationship_kind": None,
        "channel_kind": None,
        "interface_style": None,
        "portfolio_level": None,
        "realisation_layer": None,
        "realised_by": [],
    }
    return base | values


class _Answering:
    """A client that answers each call with raw items, as a transport parses them."""

    model = "answering"

    def __init__(self, *answers: list[Any]) -> None:
        self.answers = list(answers)

    def parse(self, **_: Any) -> Any:
        return SimpleNamespace(changes=self.answers.pop(0) if self.answers else [])


def _read(texts: list[str], *items: Any) -> Any:
    request = ExtractionRequest(
        "Reference",
        tuple(
            ExtractionSegment(number, f"paragraph {number}", text)
            for number, text in enumerate(texts, start=1)
        ),
        (),
    )
    extractor = StructuredCatalogueExtractor(_Answering(list(items)), supports_images=False)
    return extractor.propose(request)


def test_one_item_the_model_got_wrong_costs_only_itself() -> None:
    proposal = _read(
        ["The Order Hub receives partner orders. BCRM keeps the customer records."],
        # Eleven evidence numbers, and a key left out: both mended.
        _item(evidence_numbers=[1] * 11),
        {key: value for key, value in _item(name="BCRM", system="BCRM").items() if key != "aliases"}
        | {"quote": "BCRM keeps the customer records"},
        # Not a kind at all: left out, and counted.
        _item(kind="database"),
    )

    assert [change.content.name for change in proposal.changes] == ["Order Hub", "BCRM"]
    assert proposal.warnings == ("Incomplete suggestions were left out: 1 item.",)


def test_a_quote_with_other_typography_or_another_passage_is_cited_where_it_is() -> None:
    proposal = _read(
        [
            "Intro text.",
            "The “Order Hub” receives partner orders — every day.",
            "BCRM keeps the customer",
            "records for every account.",
        ],
        # Straight quotes and a hyphen for the document's curly quotes and dash, citing
        # the wrong passage: kept, and cited where the words are.
        _item(evidence_numbers=[1], quote='The "Order Hub" receives partner orders - every day'),
        # A sentence across two passages: cited at both.
        _item(
            name="BCRM",
            system="BCRM",
            evidence_numbers=[3, 4],
            quote="BCRM keeps the customer records",
        ),
    )

    hub, bcrm = proposal.changes
    assert hub.locations == ("paragraph 2",)
    assert hub.quote == 'The "Order Hub" receives partner orders - every day'
    assert bcrm.locations == ("paragraph 3", "paragraph 4")
    assert proposal.warnings == ()


def test_a_paraphrase_is_replaced_by_the_whole_sentence_even_after_an_initial() -> None:
    proposal = _read(
        ["Speed Booster add-ons are offered only for B. Pro+ bundles in the portal."],
        _item(quote="Speed Booster offered only for B. Pro+ bundles"),
    )

    assert proposal.changes[0].quote == (
        "Speed Booster add-ons are offered only for B. Pro+ bundles in the portal."
    )


def test_a_quote_no_passage_says_is_left_out_and_counted() -> None:
    proposal = _read(
        ["The Order Hub receives partner orders."],
        _item(),
        _item(name="Ghost", system="Ghost", quote="A ghost ledger reconciles money"),
    )

    assert [change.content.name for change in proposal.changes] == ["Order Hub"]
    assert proposal.warnings == (
        "1 suggestion(s) were left out: what they quote was not found in the document.",
    )


def _cell(text: str, merge: str | None = None, span: int = 1) -> str:
    properties = ""
    if merge is not None or span > 1:
        properties = "<w:tcPr>"
        properties += f'<w:gridSpan w:val="{span}"/>' if span > 1 else ""
        if merge is not None:
            properties += f'<w:vMerge w:val="{merge}"/>' if merge else "<w:vMerge/>"
        properties += "</w:tcPr>"
    paragraph = f"<w:p><w:r><w:t>{text}</w:t></w:r></w:p>" if text else "<w:p/>"
    return f"<w:tc>{properties}{paragraph}</w:tc>"


def test_a_word_cell_merged_down_or_across_keeps_its_value_and_columns() -> None:
    table = (
        "<w:tbl>"
        f"<w:tr>{_cell('Scenario')}{_cell('Details', span=2)}</w:tr>"
        f"<w:tr>{_cell('Faulty CPE', merge='restart')}{_cell('Replace')}{_cell('Same day')}</w:tr>"
        f"<w:tr>{_cell('', merge='')}{_cell('Return')}{_cell('Next day')}</w:tr>"
        "</w:tbl>"
    )
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("[Content_Types].xml", "<Types />")
        archive.writestr(
            "word/document.xml",
            f'<w:document xmlns:w="{W}"><w:body>{table}</w:body></w:document>',
        )

    located = LocatedDocumentExtractor(SafeDocumentTextExtractor()).extract(DOCX, buffer.getvalue())

    assert [item.text for item in located] == [
        "Scenario | Details | Details",
        "Scenario: Faulty CPE | Details: Replace | Details: Same day",
        "Scenario: Faulty CPE | Details: Return | Details: Next day",
    ]


# Plans -------------------------------------------------------------------------------------


def test_an_offerings_tiers_are_read_as_its_plans() -> None:
    offering = {
        "plans": [
            {
                "name": "PO1",
                "characteristics": [
                    {"name": "Offer Description", "value": "Unlimited 100Mbps over GSM"},
                    {"name": "IN Offer ID", "value": "860"},
                    {"name": "in offer id", "value": "861"},
                ],
                "description": None,
            },
            {"name": "PO1", "characteristics": [], "description": "Again"},
            {"name": "PO2", "characteristics": [], "description": None},
        ]
    }
    proposal = _read(
        ["Business Pro Plus has a Backup GSM Bandwidth PO per speed."],
        _item(
            kind="product_offering",
            system="",
            name="Business Pro Plus",
            offering=offering,
            quote="Business Pro Plus has a Backup GSM Bandwidth PO per speed",
        ),
    )

    (change,) = proposal.changes
    product = change.content.product
    assert product is not None
    assert [plan.name for plan in product.plans] == ["PO1", "PO2"]
    assert [(fact.name, fact.value) for fact in product.plans[0].characteristics] == [
        ("Offer Description", "Unlimited 100Mbps over GSM"),
        ("IN Offer ID", "860"),
    ]


def test_an_accepted_reading_without_plans_keeps_the_offerings_plans() -> None:
    planned = ProductOffering(
        "bpp",
        "Business Pro Plus",
        plans=(OfferingPlan("PO1", (PlanCharacteristic("Speed", "100 Mbps"),)),),
    )
    release = replace(seed_knowledge(), status=KnowledgeReleaseStatus.DRAFT, products=(planned,))
    reading = CandidateContent(
        CandidateKind.PRODUCT,
        "bpp",
        name="Business Pro Plus",
        product=ProductOffering("bpp", "Business Pro Plus", proposition="For business."),
    )

    (offering,) = apply_candidate(reading, release).products

    assert offering.proposition == "For business."
    assert offering.plans == planned.plans


# Channels and journeys ---------------------------------------------------------------------


def _offering_reading(*channels: str) -> ProposedChange:
    product = ProductOffering(
        "bpp",
        "Business Pro Plus",
        order_types=(OrderType("NEW", "New Activation", channels=channels),),
    )
    return ProposedChange(
        CandidateContent(CandidateKind.PRODUCT, "bpp", name="Business Pro Plus", product=product),
        ("paragraph 2",),
        "Applicable Channels: BCRM, B2B, SMB App",
        source_name="Business Pro Plus",
    )


class _Scripted:
    supports_images = False
    model = "scripted"
    prompt_version = "scripted-v1"

    def __init__(self, *changes: ProposedChange) -> None:
        self.changes = changes

    def propose(self, request: ExtractionRequest) -> CatalogueProposal:
        return CatalogueProposal(self.changes, self.model, self.prompt_version)


def test_channels_an_offering_is_ordered_through_are_suggested_as_channels() -> None:
    content = _docx(_p("Ordering") + _p("The journey of an order is described in the annex."))
    proposer, draft, _ = _setup(content, _Recording())
    proposer._extractor = _Scripted(_offering_reading("BCRM", "Partner Shop", "bcrm"))

    run = proposer.execute(draft.id, "doc", lambda: None)

    channels = [
        item.content.channel
        for item in proposer._candidates.list(draft.id)
        if item.content.kind is CandidateKind.CHANNEL
    ]
    # Once each; a channel named like a system enters through it.
    assert [(item.name, item.entry_system_id) for item in channels if item] == [
        ("BCRM", "bcrm"),
        ("Partner Shop", None),
    ]
    assert (
        "The document mentions journeys, but none was read: a journey is read where the "
        "document sets out the activities of an order type in order."
    ) in run.warnings


# Embedded documents ------------------------------------------------------------------------

OFFICE = "urn:schemas-microsoft-com:office:office"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
P = "http://schemas.openxmlformats.org/presentationml/2006/main"
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
TYPES = '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"/>'
RELATIONSHIPS = "http://schemas.openxmlformats.org/package/2006/relationships"


def _pptx(text: str) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("[Content_Types].xml", TYPES)
        archive.writestr(
            "ppt/presentation.xml",
            f'<p:presentation xmlns:p="{P}" xmlns:r="{R}"><p:sldIdLst>'
            '<p:sldId id="256" r:id="rId1"/></p:sldIdLst></p:presentation>',
        )
        archive.writestr(
            "ppt/_rels/presentation.xml.rels",
            f'<Relationships xmlns="{RELATIONSHIPS}">'
            f'<Relationship Id="rId1" Type="{R}/slide" Target="slides/slide1.xml"/>'
            "</Relationships>",
        )
        archive.writestr(
            "ppt/slides/slide1.xml",
            f'<p:sld xmlns:p="{P}" xmlns:a="{A}"><p:cSld><p:spTree><p:sp><p:txBody>'
            f"<a:p><a:r><a:t>{text}</a:t></a:r></a:p></p:txBody></p:sp></p:spTree></p:cSld>"
            "</p:sld>",
        )
    return buffer.getvalue()


def _with_embedded() -> bytes:
    def embedded(rel_id: str) -> str:
        return f'<w:p><w:r><w:object><o:OLEObject r:id="{rel_id}"/></w:object></w:r></w:p>'

    body = (
        _p("Annexure", "Heading1")
        + _p("Workflows", "Heading2")
        + embedded("rId1")
        + embedded("rId2")
        + embedded("rId3")
        + embedded("rId4")
        + embedded("rId5")
    )
    relationships = (
        f'<Relationship Id="rId1" Type="{R}/package" Target="embeddings/flow.pptx"/>'
        f'<Relationship Id="rId2" Type="{R}/oleObject" Target="embeddings/oleObject1.bin"/>'
        f'<Relationship Id="rId3" Type="{R}/package" Target="embeddings/spec.docx"/>'
        f'<Relationship Id="rId4" Type="{R}/package" Target="embeddings/flows.xlsx"/>'
        f'<Relationship Id="rId5" Type="{R}/package" Target="embeddings/flows2.xlsx"/>'
    )
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("[Content_Types].xml", TYPES)
        archive.writestr(
            "word/document.xml",
            f'<w:document xmlns:w="{W}" xmlns:o="{OFFICE}" xmlns:r="{R}"><w:body>{body}</w:body>'
            "</w:document>",
        )
        archive.writestr(
            "word/_rels/document.xml.rels",
            f'<Relationships xmlns="{RELATIONSHIPS}">{relationships}</Relationships>',
        )
        archive.writestr("word/embeddings/flow.pptx", _pptx("Step 1: BCRM captures the order."))
        archive.writestr("word/embeddings/oleObject1.bin", b"legacy OLE package")
        archive.writestr(
            "word/embeddings/spec.docx", _docx(_p("CWOM orchestrates the fulfilment."))
        )
        # Two versions of one workflow book, the same word for word.
        archive.writestr("word/embeddings/flows.xlsx", _workbook())
        archive.writestr("word/embeddings/flows2.xlsx", _workbook())
    return buffer.getvalue()


def test_office_files_a_word_document_embeds_are_read_and_their_citations_open() -> None:
    extractor = _Recording()
    proposer, draft, reader = _setup(_with_embedded(), extractor)

    run = proposer.execute(draft.id, "doc", lambda: None)

    read = [
        (item.location, item.section, item.text)
        for item in extractor.requests[-1].segments
        if item.location.startswith("embedded")
    ]
    # Documents and decks first; the second workbook's sheet repeats the first's.
    assert [location for location, _, _ in read] == [
        "embedded 3, paragraph 1",
        "embedded 1, slide 1",
        "embedded 4, sheet 1",
        "embedded 4, sheet 1, row 1",
        "embedded 4, sheet 1, row 2",
    ]
    assert read[:2] == [
        (
            "embedded 3, paragraph 1",
            ("Annexure", "Workflows", "spec.docx"),
            "CWOM orchestrates the fulfilment.",
        ),
        (
            "embedded 1, slide 1",
            ("Annexure", "Workflows", "flow.pptx"),
            "Step 1: BCRM captures the order.",
        ),
    ]
    assert (
        "1 embedded sheet(s) were left out: an embedded workbook read before has the same "
        "sheet, word for word."
    ) in run.warnings
    assert (
        "1 embedded object(s) were not read: only Word, PowerPoint and Excel files are, each "
        "within the upload limits."
    ) in run.warnings

    passage = reader.passage(draft.id, "doc", "embedded 3, paragraph 1", MAINTAINER)
    assert passage.passage.text == "CWOM orchestrates the fulfilment."


def _workbook() -> bytes:
    book = Workbook()
    flows = book.active
    assert flows is not None
    flows.title = "Provisioning WF"
    flows.append(["Step", "System"])
    flows.append(["1", "BCRM captures the order"])
    filler = book.create_sheet("Sheet1")
    filler.sheet_state = "hidden"
    for row in range(1, 50):
        filler.cell(row, 7, "filler")
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def _sheets(content: bytes) -> list[str]:
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        return sorted(name for name in archive.namelist() if name.startswith("xl/worksheets/"))


def test_an_embedded_workbook_is_read_without_its_hidden_sheets() -> None:
    """Each embedded workflow book hid a sheet filled to row 1,048,576 that refused the book."""
    slim = _without_hidden_sheets(_workbook())

    assert _sheets(slim) == ["xl/worksheets/sheet1.xml"]
    located = LocatedDocumentExtractor(SafeDocumentTextExtractor()).extract(XLSX, slim)
    assert [item.text for item in located if item.heading] == ["Provisioning WF"]
    assert any("BCRM captures the order" in item.text for item in located)
    assert _without_hidden_sheets(b"not a workbook") == b"not a workbook"


def test_a_dependency_the_model_gave_no_words_for_is_said_by_its_kind() -> None:
    """The reported reading left out 104 stated dependencies whose description was null."""
    proposal = _read(
        ["SMB App calls CRMGW APIs. B2B uses TIBCO."],
        _item(
            kind="relationship",
            system="SMB App",
            name=None,
            target_system="CRMGW",
            relationship_kind="calls_api",
            quote="SMB App calls CRMGW APIs",
        ),
        _item(
            kind="relationship",
            system="B2B",
            name=None,
            target_system="TIBCO",
            quote="B2B uses TIBCO",
        ),
        # No target: never a dependency on a system called "item".
        _item(kind="relationship", system="B2B", name=None, quote="B2B uses TIBCO"),
    )

    assert [change.content.text for change in proposal.changes] == [
        "Calls the API of CRMGW",
        "Depends on TIBCO",
    ]
    assert proposal.warnings == ("Incomplete suggestions were left out: 1 relationship.",)


# Workflow books ----------------------------------------------------------------------------


def _workflow_book() -> bytes:
    """An SDD's workflow sheet: a title merged across, a header, numbered stages, a note."""
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = "Provisioning WF-P1"
    sheet.append(["WF-1", "Order fulfillment | Business Pro Plus New Request"])
    sheet.merge_cells("B1:E1")
    sheet.append(["Seq.", "Description", "Section", "System", "OLA"])
    sheet.append(["1", "CPE platform selection", "System", "BCRM", "3 Min"])
    sheet.append(["2", "Domain name provisioning", "System", "CWOM", None])
    sheet.append(["3", "Customer appointment", "Field", "WFMS", None])
    sheet.append(["Note: 1", "Templates are shown in the portal", None, None, None])
    sheet.append([None])
    sheet.append(["Version", "Owner"])
    sheet.append(["1.0", "OM team"])
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def test_a_workbooks_rows_are_read_under_their_header_and_title() -> None:
    located = LocatedDocumentExtractor(SafeDocumentTextExtractor()).extract(XLSX, _workflow_book())

    by_place = {item.location: item for item in located}
    title = by_place["sheet 1, row 1"]
    assert (title.heading, title.text) == (
        True,
        "WF-1 Order fulfillment | Business Pro Plus New Request",
    )
    stage = by_place["sheet 1, row 3"]
    assert stage.text == (
        "Seq.: 1 | Description: CPE platform selection | Section: System | System: BCRM | "
        "OLA: 3 Min"
    )
    assert stage.heading_path == ("Provisioning WF-P1", title.text)
    assert ("System", "BCRM") in stage.cells
    # Two labels are not a header: that table reads as the extractor wrote it.
    assert by_place["sheet 1, row 9"].text == "A9=1.0 | B9=OM team"


def test_a_workflow_table_is_read_as_a_journey_without_the_model() -> None:
    located = LocatedDocumentExtractor(SafeDocumentTextExtractor()).extract(XLSX, _workflow_book())
    segments = tuple(
        ExtractionSegment(
            number, item.location, item.text, section=item.heading_path, cells=item.cells
        )
        for number, item in enumerate(located, start=1)
    )

    reading = CatalogueTableReader().read(ExtractionRequest("SDD", segments, ()))

    (change,) = reading.changes
    journey = change.content.journey
    assert journey is not None
    assert journey.name == "WF-1 Order fulfillment — Business Pro Plus New Request"
    assert [(item.number, item.name, item.performing_system_id) for item in journey.activities] == [
        ("1", "CPE platform selection", "bcrm"),
        ("2", "Domain name provisioning", "cwom"),
        ("3", "Customer appointment", "wfms"),
    ]
    assert change.reader is not None and change.reader[1] == "catalogue-tables-v10"


# One workflow read twice; combined system names ------------------------------------------


def _journey_reading(name: str, *steps: tuple[str, str]) -> ProposedChange:
    journey = Journey(
        slug(name),
        name,
        activities=tuple(
            Activity(str(number), step, performing_system_id=performer)
            for number, (step, performer) in enumerate(steps, start=1)
        ),
    )
    return ProposedChange(
        CandidateContent(CandidateKind.JOURNEY, journey.id, name=name, journey=journey),
        ("paragraph 2",),
        "Workflows",
        source_name=name,
    )


STAGES = (
    ("Waiter clearance", "wfms"),
    ("ONT/OLT configuration", "om-e2eso"),
    ("HSI configuration", "b2b-portal"),
)


def test_one_workflow_read_twice_is_one_journey_and_two_workflows_stay_two() -> None:
    proposer, draft, _ = _setup(_docx(_p("Workflows")), _Recording())
    proposer._extractor = _Scripted(
        ProposedChange(
            CandidateContent(CandidateKind.SYSTEM, "e2eso", name="E2ESO"),
            ("paragraph 1",),
            "Workflows",
            source_name="E2ESO",
        ),
        _journey_reading("WF-3-B Order fulfillment — GPON Provisioning order", *STAGES),
        # Another version of the book writes it "WF-3B".
        _journey_reading("WF-3B Order fulfillment — GPON Provisioning order", *STAGES[:2]),
        # The model names only its start, from fewer steps.
        _journey_reading("WF-3B Order fulfillment", STAGES[0]),
        # The same number, another workflow.
        _journey_reading("WF-3B Up/Downgrade", ("Cease service", "wfms")),
    )

    proposer.execute(draft.id, "doc", lambda: None)

    journeys = [
        item.content.journey
        for item in proposer._candidates.list(draft.id)
        if item.content.journey is not None
    ]
    assert [journey.name for journey in journeys] == [
        "WF-3-B Order fulfillment — GPON Provisioning order",
        "WF-3B Up/Downgrade",
    ]
    first = journeys[0]
    # "OM-E2ESO" is E2ESO, which the catalogue knows as an alias of its activation
    # systems; "B2B-PORTAL" is one name, its parts too long to be acronyms.
    assert [
        (step.performing_system_id, step.supporting_system_ids) for step in first.activities
    ] == [
        ("wfms", ()),
        ("service-activation", ()),
        ("b2b-portal", ()),
    ]


def _coded_book() -> bytes:
    """A workflow sheet as the SDD's later book version writes it: a title across three
    cells above a header, and a coded title row whose stages reuse the header above."""
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = "P3"
    sheet.append(["WF-3A", "Order fulfillment", "PORT IN Shift, Main Order"])
    sheet.append(["Seq.", "Description", "Section", "System", "Remarks"])
    sheet.append(["1", "Number portability check", "System", "NPS", "WFMS will send to OM"])
    sheet.append(["2", "Order capture", "System", "CWOM", None])
    sheet.append(["3", "Line installation", "Field", "WFMS", None])
    sheet.append(["WF-3", "Migration of Legacy product DEL, PBX", "Oder Type : Migrate"])
    sheet.append(["1", "Cease legacy line", "System", "CWOM", None])
    sheet.append(["2", "Provision new line", "System", "E2ESO", None])
    sheet.append(["3", "Notify customer", "System", "CWOM", None])
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def test_coded_workflow_titles_each_head_their_journey_and_rows_stay_open() -> None:
    located = LocatedDocumentExtractor(SafeDocumentTextExtractor()).extract(XLSX, _coded_book())
    segments = tuple(
        ExtractionSegment(
            number, item.location, item.text, section=item.heading_path, cells=item.cells
        )
        for number, item in enumerate(located, start=1)
    )

    reading = CatalogueTableReader().read(ExtractionRequest("SDD", segments, ()))

    journeys = [change.content.journey for change in reading.changes]
    assert [(journey.name, len(journey.activities)) for journey in journeys if journey] == [
        ("WF-3A Order fulfillment PORT IN Shift, Main Order", 3),
        ("WF-3 Migration of Legacy product DEL, PBX Oder Type : Migrate", 3),
    ]
    # Their stages still go to the model, which reads the flows their remarks state.
    assert reading.read == set()
    assert reading.consumed == set()


def test_only_a_reading_with_room_asks_for_every_offering_and_prose_journeys() -> None:
    for rule in (
        "Every offering the document names",
        "A process the document describes step by step",
        "A team, group or function named in a System column",
    ):
        assert rule in DETAILED_SYSTEM_PROMPT
        assert rule not in LEAN_SYSTEM_PROMPT
    # Asked to propose every name, a small model proposed tickets and phases as systems.
    assert "Propose every system the passages name" not in DETAILED_SYSTEM_PROMPT


# Cleaning what is read ---------------------------------------------------------------------

COPIED = (
    "Backup GSM Bandwidth offers are defined per GPON speed profile and provisioned by CWOM "
    "through E2ESO once the order is validated."
)


class _Quoting(_Recording):
    """Records what it is asked to read, and reads one system from the first copy."""

    def propose(self, request: ExtractionRequest) -> CatalogueProposal:
        self.requests.append(request)
        copy = next(item for item in request.segments if item.text == COPIED)
        return CatalogueProposal(
            (
                ProposedChange(
                    CandidateContent(
                        CandidateKind.SYSTEM, "bundle-orchestrator", name="Bundle Orchestrator"
                    ),
                    (copy.location,),
                    "provisioned by CWOM through E2ESO",
                    source_name="Bundle Orchestrator",
                ),
            ),
            "quoting",
            "quoting-v1",
        )


def test_copied_blocks_are_read_once_and_cited_at_every_copy_and_history_is_left_out() -> None:
    body = (
        _p("Revision History", "Heading1")
        + _p("Version 1.2 updated the CPE table and the exit charges for the bundle.")
        + _p("PSM", "Heading1")
        + _p(COPIED)
        + _p("ECM", "Heading1")
        + _p(COPIED)
        + _p(
            '{"orderId": "1", "items": [{"spec": "CFSS_INTERNET_CPE_HE", "action": "add"}], '
            '"state": "acknowledged", "channel": "B2B", "notes": [], "relatedParty": '
            '[{"id": "2", "role": "customer"}], "billingAccount": {"id": "3"}}'
        )
    )
    extractor = _Quoting()
    proposer, draft, _ = _setup(_docx(body), extractor)

    run = proposer.execute(draft.id, "doc", lambda: None)

    read = [item.text for item in extractor.requests[-1].segments]
    assert read.count(COPIED) == 1
    assert not any("Version 1.2" in text or text.startswith("{") for text in read)
    (system,) = [
        item
        for item in proposer._candidates.list(draft.id)
        if item.content.kind is CandidateKind.SYSTEM
    ]
    assert [citation.location for citation in system.citations] == ["paragraph 4", "paragraph 6"]
    assert "1 repeated passage(s) were read once and cited wherever they appear." in run.warnings
    assert any(warning.startswith("3 passage(s) with no architecture") for warning in run.warnings)


def _system_reading(name: str) -> CatalogueCandidate:
    return CatalogueCandidate(
        "s1",
        "draft",
        "doc",
        CandidateContent(CandidateKind.SYSTEM, slug(name), name=name),
        (CandidateCitation("row 1", name),),
        "model",
        "v1",
        datetime(2026, 10, 12, tzinfo=UTC),
    )


def test_a_name_spelled_almost_like_a_draft_system_is_asked_about_never_merged() -> None:
    systems = (
        SystemDefinition("fortiguard", "FortiGuard", ()),
        SystemDefinition("e2eso", "E2ESO", ()),
        SystemDefinition("fortimanager-7-2", "FortiManager 7.2", ()),
        SystemDefinition("cwom", "CWOM", ()),
    )
    release = replace(
        seed_knowledge(), status=KnowledgeReleaseStatus.DRAFT, systems=systems, relationships=()
    )

    def asked(name: str) -> list[str]:
        return [match.system_id for match in open_matches(_system_reading(name), release)]

    assert asked("FortiGaurd") == ["fortiguard"]
    assert asked("ESESO") == ["e2eso"]
    # Another version, not a slip; and too short a name to tell.
    assert asked("FortiManager 7.4") == []
    assert asked("CWOX") == []
    assert classify(_system_reading("FortiGaurd").content, release) is CandidateMatch.NEW
    # Asked once, on the system's own suggestion: a capability naming it is not asked again.
    capability = replace(
        _system_reading("FortiGaurd"),
        content=CandidateContent(
            CandidateKind.CAPABILITY,
            "fortigaurd",
            capability_id="licence-check",
            name="Licence check",
            triggers=("check licence",),
        ),
    )
    assert open_matches(capability, release) == ()


def test_workflow_titles_match_however_their_parts_are_separated() -> None:
    def journey(name: str) -> Journey:
        return Journey(
            slug(name),
            name,
            activities=(Activity("1", "Waiter clearance"), Activity("2", "ONT configuration")),
        )

    assert _same_workflow(
        journey("WF-3A Order fulfillment — Business Pro/ PRO+ with 5G Backup"),
        journey("WF-3A Order fulfillment | Business Pro/ PRO+ with 5G Backup New Request"),
    )
    assert _same_workflow(
        journey("WF-2-A Up/Downgrade Change plan (upgrade) from Office Presence"),
        journey("WF-2-A Up/Downgrade - Change plan (upgrade) from Office Presence"),
    )
    assert not _same_workflow(
        journey("WF-1 Order fulfillment — GPON"), journey("WF-1 Up/Downgrade — GPON")
    )


# What waits for what ------------------------------------------------------------------------


def test_a_name_confirmed_as_an_alias_is_never_waited_for() -> None:
    """The reported draft: OM-E2ESO, confirmed as ESESO, still kept 24 suggestions waiting."""
    systems = (SystemDefinition("eseso", "ESESO", ("OM-E2ESO",)),)
    release = replace(
        seed_knowledge(), status=KnowledgeReleaseStatus.DRAFT, systems=systems, relationships=()
    )
    journey = Journey(
        "wf-1",
        "WF-1 Order fulfillment",
        activities=(
            Activity("1", "Configure CPE", performing_system_id="om-e2eso"),
            Activity("2", "Configure ONT", performing_system_id="om-enrm-asap"),
        ),
    )
    link = CandidateContent(
        CandidateKind.RELATIONSHIP, "eseso", target_system_id="rtf", text="Sends orders"
    )

    assert missing_systems(
        CandidateContent(CandidateKind.JOURNEY, "wf-1", name=journey.name, journey=journey), release
    ) == ("om-enrm-asap",)
    assert missing_systems(link, release) == ("rtf",)
    assert missing_systems(CandidateContent(CandidateKind.SYSTEM, "nec", name="NEC"), release) == ()


def _performer_book(*performers: str) -> bytes:
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = "Provisioning WF-P1"
    sheet.append(["WF-1", "Order fulfillment"])
    sheet.append(["Seq.", "Description", "Section", "System"])
    for number, performer in enumerate(performers, start=1):
        sheet.append([str(number), f"Stage {number}", "System", performer])
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def test_a_workflow_stage_names_its_performer_whole_and_a_bare_system_cell_names_none() -> None:
    located = LocatedDocumentExtractor(SafeDocumentTextExtractor()).extract(
        XLSX, _performer_book("E2ESO-ASAP ENG/UCAAS", "System", "WFMS")
    )
    segments = tuple(
        ExtractionSegment(
            number, item.location, item.text, section=item.heading_path, cells=item.cells
        )
        for number, item in enumerate(located, start=1)
    )

    (change,) = CatalogueTableReader().read(ExtractionRequest("SDD", segments, ())).changes

    journey = change.content.journey
    assert journey is not None
    assert [step.performing_system_id for step in journey.activities] == [
        "e2eso-asap-eng-ucaas",
        None,
        "wfms",
    ]


def test_a_link_from_a_group_of_channels_or_systems_is_left_out() -> None:
    proposal = _read(
        ["Channels submit the request to RTF. NPS sends the request to RTF."],
        _item(
            kind="relationship",
            system="Channels",
            name=None,
            target_system="RTF",
            quote="Channels submit the request to RTF",
        ),
        _item(
            kind="relationship",
            system="NPS",
            name=None,
            target_system="RTF",
            quote="NPS sends the request to RTF",
        ),
    )

    assert [change.source_name for change in proposal.changes] == ["NPS"]
    assert proposal.warnings == ("Incomplete suggestions were left out: 1 relationship.",)


def test_an_item_about_no_system_is_left_out_not_given_a_system_called_item() -> None:
    # The reported reading: a spreadsheet row whose "Section" said only "System".
    proposal = _read(
        ["The Order Hub receives partner orders. Online Billing is automated."],
        _item(),
        _item(
            kind="capability",
            system="",
            name="Online Billing",
            triggers=["month end"],
            quote="Online Billing is automated",
        ),
    )

    assert [change.content.system_id for change in proposal.changes] == ["order-hub"]
    assert proposal.warnings == ("Incomplete suggestions were left out: 1 capability.",)


def _link(
    source: str, target: str, basis: CandidateBasis = CandidateBasis.STATED
) -> ProposedChange:
    return ProposedChange(
        CandidateContent(
            CandidateKind.RELATIONSHIP, slug(source), target_system_id=slug(target), text="REST API"
        ),
        ("embedded 1, slide 4",),
        f"{source} | {target} | REST API",
        source_name=source,
        target_name=target,
        basis=basis,
        rationale="Implied by the flow." if basis is CandidateBasis.INFERRED else None,
    )


def test_a_system_only_an_interface_table_names_is_suggested_once_where_it_is_named() -> None:
    content = _docx(_p("Interfaces") + _p("The interfaces are listed in the annex."))
    proposer, draft, _ = _setup(content, _Recording())
    proposer._extractor = _Scripted(
        _link("vCPE DO", "FortiGuard"),
        _link("vCPE DO", "BCRM"),
        _link("NRM", "vCPE DO"),
        # Neither a placeholder nor an inferred link names a system.
        _link("Unknown", "BCRM"),
        _link("Ghost", "BCRM", CandidateBasis.INFERRED),
    )

    proposer.execute(draft.id, "doc", lambda: None)

    systems = {
        item.content.name: item
        for item in proposer._candidates.list(draft.id)
        if item.content.kind is CandidateKind.SYSTEM
    }
    # NRM is in the draft already.
    assert sorted(systems) == ["FortiGuard", "vCPE DO"]
    assert systems["vCPE DO"].content.system_id == "vcpe-do"
    assert [citation.location for citation in systems["vCPE DO"].citations] == [
        "embedded 1, slide 4"
    ]
