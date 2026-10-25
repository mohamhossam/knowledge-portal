"""Edge cases of the organisation catalogue, catalogue files and AI suggestions."""

from __future__ import annotations

import io
import zipfile
from dataclasses import replace
from datetime import UTC, datetime
from threading import RLock
from typing import Any

import psycopg
import pytest
from fastapi.testclient import TestClient
from smb_kernel.documents.text_extractor import SafeDocumentTextExtractor
from smb_kernel.time.fixed import FixedClock

from knowledge_portal.application.errors import PersistenceError
from knowledge_portal.application.ports.architecture_knowledge import ArchitectureQuery
from knowledge_portal.application.ports.catalogue_extractor import (
    CatalogueExtractionUnsupportedError,
    ExtractionRequest,
    ExtractionSegment,
    KnownSystem,
)
from knowledge_portal.application.ports.catalogue_file import CatalogueFileFormat
from knowledge_portal.application.ports.identity import Actor
from knowledge_portal.application.use_cases.architecture_index import BuildArchitectureIndex
from knowledge_portal.application.use_cases.architecture_knowledge import (
    KnowledgeNotFoundError,
    ManageArchitectureKnowledge,
)
from knowledge_portal.application.use_cases.catalogue_candidates import (
    DecideCatalogueCandidate,
    ProposeCatalogueChanges,
)
from knowledge_portal.application.use_cases.organisation_catalogue import (
    ManageOrganisationCatalogue,
)
from knowledge_portal.application.use_cases.resolve_architecture_knowledge import (
    ResolveArchitectureKnowledge,
)
from knowledge_portal.domain.architecture.candidates import (
    CandidateCitation,
    CandidateContent,
    CandidateDecisionConflictError,
    CandidateKind,
    CandidateNotFoundError,
    CatalogueCandidate,
)
from knowledge_portal.domain.architecture.knowledge import (
    InvalidKnowledgeError,
    KnowledgeConflictError,
    KnowledgeDocumentVersion,
    KnowledgeReleaseStatus,
)
from knowledge_portal.domain.document.value_objects import DocumentVersionId
from knowledge_portal.domain.identity.errors import AuthorizationDeniedError
from knowledge_portal.domain.organisation.catalogue import (
    OrganisationNotFoundError,
    Person,
    Product,
    Squad,
    SquadResource,
    SquadRole,
    ValueStream,
)
from knowledge_portal.infrastructure.architecture.catalogue_files import CatalogueFileAdapter
from knowledge_portal.infrastructure.architecture.embeddings import FakeEmbeddings
from knowledge_portal.infrastructure.architecture.evidence_index import InMemoryEvidenceIndex
from knowledge_portal.infrastructure.architecture.located_extractor import (
    LocatedDocumentExtractor,
)
from knowledge_portal.infrastructure.architecture.reasoning import FakeArchitectureReasoner
from knowledge_portal.infrastructure.architecture.tokenizer import FakeWordTokenizer
from knowledge_portal.infrastructure.documents.library_worker import OfflineDocumentScanner
from knowledge_portal.infrastructure.llm.catalogue_extraction import (
    FakeCatalogueExtractor,
    slug,
)
from knowledge_portal.infrastructure.llm.catalogue_matching import FakeSystemMatcher
from knowledge_portal.infrastructure.persistence.document_storage import (
    InMemoryDocumentStorage,
)
from knowledge_portal.infrastructure.persistence.in_memory_catalogue_candidates import (
    InMemoryCatalogueCandidates,
)
from knowledge_portal.infrastructure.persistence.in_memory_organisation import (
    InMemoryOrganisationRepository,
)
from knowledge_portal.infrastructure.persistence.in_memory_transaction import (
    InMemoryTransactionManager,
)
from knowledge_portal.infrastructure.persistence.postgres_catalogue_candidates import (
    PostgresCatalogueCandidates,
)
from knowledge_portal.infrastructure.persistence.postgres_organisation import (
    PostgresOrganisationRepository,
)
from knowledge_portal.infrastructure.persistence.postgres_store import PostgresStore
from tests.catalogue_seed import seed_knowledge, seeded_repository

NOW = datetime(2026, 9, 29, 9, 0, tzinfo=UTC)
MAINTAINER = Actor("amina", frozenset({"knowledge_maintainer"}))
READER = Actor("ravi", frozenset({"knowledge_reader"}))
OWNER = {"X-Fake-Actor-Id": "fake-owner"}
ADAPTER = CatalogueFileAdapter()


def test_organisation_use_case_covers_every_record_kind() -> None:
    releases = seeded_repository()
    manage = ManageOrganisationCatalogue(InMemoryOrganisationRepository(FixedClock(NOW)), releases)
    manage.save_person(Person("layla", "Layla"), None, MAINTAINER)
    manage.save_value_stream(ValueStream("retail", "Retail", "layla"), None, MAINTAINER)
    manage.save_product(Product("fibre", "retail", "Fibre", system_ids=("bcrm",)), None, MAINTAINER)
    manage.save_squad(
        Squad("sales", "Sales", "retail", None, (SquadResource("bcrm", SquadRole.SYSTEM_CONTACT),)),
        None,
        MAINTAINER,
    )
    # A link that lapsed from the active release does not block unrelated edits.
    releases.save(
        replace(
            seed_knowledge(),
            id="draft",
            status=KnowledgeReleaseStatus.DRAFT,
            published_at=None,
            published_by=None,
        ),
        None,
        "amina",
        "create_draft",
    )
    renamed = manage.save_product(
        Product("fibre", "retail", "Business fibre", system_ids=("bcrm",)), 1, MAINTAINER
    )
    assert renamed.product("fibre").name == "Business fibre"
    assert manage.ownership("bcrm", READER).products[0].name == "Business fibre"

    with pytest.raises(AuthorizationDeniedError):
        manage.audit(READER)
    with pytest.raises(OrganisationNotFoundError):
        manage.remove_product("missing", 1, MAINTAINER)
    manage.remove_squad("sales", 1, MAINTAINER)
    manage.remove_product("fibre", 2, MAINTAINER)
    assert manage.remove_value_stream("retail", 1, MAINTAINER).value_streams == ()
    with pytest.raises(OrganisationNotFoundError):
        manage.view(MAINTAINER).value_stream("retail")


def test_organisation_routes_update_and_remove_each_kind(client: TestClient) -> None:
    def call(method: str, path: str, body: dict[str, Any]) -> Any:
        return client.request(method, f"/organisation{path}", json=body, headers=OWNER)

    call("POST", "/people", {"person": {"id": "sam", "name": "Sam"}})
    assert call("PUT", "/people/sam", {"person": {"id": "sam", "name": "Sam R"}}).status_code == 422
    assert (
        call(
            "PUT",
            "/people/sam",
            {"expected_revision": 1, "person": {"id": "sam", "name": "Sam R", "team": "Agile"}},
        ).json()["people"][0]["team"]
        == "Agile"
    )
    call("POST", "/value-streams", {"value_stream": {"id": "vs", "name": "Retail"}})
    updated = call(
        "PUT",
        "/value-streams/vs",
        {
            "expected_revision": 1,
            "value_stream": {"id": "vs", "name": "Retail", "lead_person_id": "sam"},
        },
    )
    assert updated.json()["value_streams"][0]["lead_person_id"] == "sam"
    product = {"id": "p", "value_stream_id": "vs", "name": "Fibre", "system_ids": ["bcrm"]}
    assert call("POST", "/products", {"product": product}).status_code == 201
    assert (
        call("PUT", "/products/p", {"expected_revision": 1, "product": product}).status_code == 200
    )
    blocked = call("DELETE", "/value-streams/vs", {"expected_revision": 2})
    assert blocked.status_code == 422
    assert call("DELETE", "/products/p", {"expected_revision": 2}).json()["products"] == []
    assert call("DELETE", "/value-streams/vs", {"expected_revision": 2}).status_code == 200
    audit = client.get("/organisation/audit", headers=OWNER).json()
    assert audit[0]["action"] == "remove_value_stream"


def _workbook_with(parts: dict[str, bytes]) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        for name, data in parts.items():
            archive.writestr(name, data)
    return buffer.getvalue()


@pytest.mark.parametrize(
    ("file_format", "content", "message"),
    [
        (CatalogueFileFormat.YAML, b"\xff\xfe", "UTF-8"),
        (CatalogueFileFormat.YAML, b"systems: [\n", "YAML file could not be parsed"),
        (CatalogueFileFormat.YAML, b"- a list\n", "object at the top level"),
        (CatalogueFileFormat.YAML, b"systems: nope\n", "'systems' must be a list"),
        (
            CatalogueFileFormat.YAML,
            b"systems:\n  - {id: a, name: A, capabilities: nope}\n",
            "capabilities must be a list",
        ),
        (
            CatalogueFileFormat.YAML,
            b"systems:\n  - {id: a, name: A, capabilities: [x]}\n",
            "capability 1 must be an object",
        ),
        (
            CatalogueFileFormat.YAML,
            b"systems:\n  - {id: a, name: A, aliases: x}\n",
            "aliases must be a list",
        ),
        (
            CatalogueFileFormat.YAML,
            b"systems:\n  - {id: a, name: [A]}\n",
            "name must be text",
        ),
        (
            CatalogueFileFormat.JSON,
            b'{"systems": [{"id": "a", "name": "A"}], "dependencies": [{"source_system_id": "a",'
            b' "target_system_id": "a", "description": "self"}]}',
            "dependencies entry 1: A system cannot depend on itself",
        ),
        (
            CatalogueFileFormat.JSON,
            b'{"systems": [{"id": "a", "name": "A", "capabilities": [{"id": "c", "name": "C"}]}]}',
            "capability 1: A capability needs",
        ),
        (
            CatalogueFileFormat.XLSX,
            _workbook_with({"[Content_Types].xml": b"<Types/>"}),
            "could not be opened",
        ),
    ],
)
def test_catalogue_file_errors_say_where(
    file_format: CatalogueFileFormat, content: bytes, message: str
) -> None:
    with pytest.raises(InvalidKnowledgeError, match=message):
        ADAPTER.read(file_format, content)


def test_workbook_rows_are_checked_against_the_systems_sheet() -> None:
    from openpyxl import Workbook

    workbook = Workbook()
    workbook.remove(workbook.worksheets[0])
    systems = workbook.create_sheet("Systems")
    for row in (("system_id", "name"), (7.0, "Seven"), ("7", "Again")):
        systems.append(list(row))
    buffer = io.BytesIO()
    workbook.save(buffer)
    with pytest.raises(InvalidKnowledgeError, match="system '7' is listed twice"):
        ADAPTER.read(CatalogueFileFormat.XLSX, buffer.getvalue())

    workbook = Workbook()
    workbook.remove(workbook.worksheets[0])
    workbook.create_sheet("Systems").append(["system_id", "name"])
    constraints = workbook.create_sheet("Constraints")
    constraints.append(["system_id", "constraint"])
    constraints.append(["erp", "Batch only"])
    buffer = io.BytesIO()
    workbook.save(buffer)
    with pytest.raises(InvalidKnowledgeError, match="Constraints row 2: system 'erp'"):
        ADAPTER.read(CatalogueFileFormat.XLSX, buffer.getvalue())


def test_file_import_refuses_published_empty_and_oversized_files() -> None:
    repository = seeded_repository()
    manage = ManageArchitectureKnowledge(
        repository,
        ADAPTER,
        InMemoryEvidenceIndex(FakeEmbeddings(), FakeWordTokenizer()),
        10,
        OfflineDocumentScanner(),
    )
    with pytest.raises(KnowledgeConflictError, match="Only a draft"):
        manage.preview_file_import(seed_knowledge().id, "c.yaml", b"systems: []", MAINTAINER)
    draft = manage.create_draft(MAINTAINER, "Next version")
    with pytest.raises(InvalidKnowledgeError, match="empty"):
        manage.preview_file_import(draft.id, "c.yaml", b"", MAINTAINER)
    with pytest.raises(InvalidKnowledgeError, match="too large"):
        manage.preview_file_import(draft.id, "c.yaml", b"x" * 11, MAINTAINER)
    assert manage.export_file(draft.id, CatalogueFileFormat.JSON, MAINTAINER).startswith(b"{")


def _proposer(
    extractor: Any = None,
) -> tuple[
    ProposeCatalogueChanges,
    DecideCatalogueCandidate,
    Any,
    InMemoryCatalogueCandidates,
    ManageArchitectureKnowledge,
]:
    repository = seeded_repository()
    manage = ManageArchitectureKnowledge(
        repository,
        ADAPTER,
        InMemoryEvidenceIndex(FakeEmbeddings(), FakeWordTokenizer()),
        10_000,
        OfflineDocumentScanner(),
    )
    storage = InMemoryDocumentStorage()
    candidates = InMemoryCatalogueCandidates()
    documents = SafeDocumentTextExtractor()
    proposer = ProposeCatalogueChanges(
        manage,
        storage,
        LocatedDocumentExtractor(documents),
        documents,
        extractor or FakeCatalogueExtractor(),
        FakeSystemMatcher(),
        candidates,
        FixedClock(NOW),
    )
    draft = manage.create_draft(MAINTAINER, "Next version")
    version = KnowledgeDocumentVersion(
        "doc", "Doc", "doc.txt", "text/plain", "sum", "doc", "amina", NOW
    )
    image = replace(version, id="img", filename="d.png", mime_type="image/png", storage_key="img")
    storage.put(DocumentVersionId("doc"), b"Nothing architectural here.\n")
    drafted = draft.updated(documents=(version, image))
    repository.save(drafted, draft.revision, "amina", "upload_document")
    decide = DecideCatalogueCandidate(
        manage, repository, candidates, InMemoryTransactionManager(RLock())
    )
    return proposer, decide, drafted, candidates, manage


class _NoVision(FakeCatalogueExtractor):
    supports_images = False


def test_extraction_guards_and_empty_documents() -> None:
    proposer, decide, draft, _, _ = _proposer(_NoVision())
    with pytest.raises(KnowledgeNotFoundError):
        proposer.document(draft.id, "missing", MAINTAINER)
    with pytest.raises(CatalogueExtractionUnsupportedError):
        proposer.document(draft.id, "img", MAINTAINER)

    stages: list[str] = []
    run = proposer.execute(draft.id, "doc", lambda: None, lambda step: stages.append(step.stage))

    assert run.warnings == ("No catalogue content was found in this document.",)
    assert stages == ["opening", "matching", "saving"]
    assert decide.overview(draft.id, MAINTAINER).runs == (run,)
    with pytest.raises(KnowledgeNotFoundError):
        proposer.execute(draft.id, "gone", lambda: None)
    with pytest.raises(KnowledgeConflictError, match="drafts only"):
        proposer.execute(seed_knowledge().id, "doc", lambda: None)


def _candidate(draft_id: str, content: CandidateContent, id_: str = "c1") -> CatalogueCandidate:
    return CatalogueCandidate(
        id_, draft_id, "doc", content, (CandidateCitation("line 1", "quote"),), "m", "v", NOW
    )


def test_decisions_check_revision_ownership_and_undo_on_a_lost_draft() -> None:
    _, decide, draft, candidates, manage = _proposer()
    system = CandidateContent(CandidateKind.SYSTEM, "hub", name="Hub")
    orphan = CandidateContent(CandidateKind.CONSTRAINT, "nowhere", text="Only nightly")
    candidates.replace_proposals(
        _run(draft.id),
        (_candidate(draft.id, system), _candidate(draft.id, orphan, "c2")),
    )

    with pytest.raises(KnowledgeConflictError, match="reload"):
        decide.decide(draft.id, "c1", draft.revision + 5, True, MAINTAINER, NOW)
    with pytest.raises(CandidateNotFoundError):
        decide.decide(draft.id, "unknown", draft.revision, True, MAINTAINER, NOW)
    with pytest.raises(InvalidKnowledgeError, match="cannot change what kind"):
        decide.decide(
            draft.id, "c1", draft.revision, True, MAINTAINER, NOW, content=replace(orphan)
        )

    class _LosingRepository:
        def save(self, *_: object) -> None:
            raise PersistenceError("lost")

    # A save that loses rolls the decision marks back with it: they share one
    # transaction, so no candidate is left decided against an unsaved draft.
    rollback = InMemoryTransactionManager(RLock())
    rollback.enroll(candidates)
    losing = DecideCatalogueCandidate(
        manage,
        _LosingRepository(),  # type: ignore[arg-type]
        candidates,
        rollback,
    )
    with pytest.raises(PersistenceError):
        losing.decide(draft.id, "c1", draft.revision, True, MAINTAINER, NOW)
    assert candidates.get("c1") == _candidate(draft.id, system)

    updated, remaining = decide.accept_all(draft.id, draft.revision, MAINTAINER, NOW)
    assert remaining == 1
    assert any(item.id == "hub" for item in updated.systems)
    same, still = decide.accept_all(draft.id, updated.revision, MAINTAINER, NOW)
    assert (same.revision, still) == (updated.revision, 1)
    with pytest.raises(CandidateDecisionConflictError):
        candidates.save_decision(candidates.get("c1") or _candidate(draft.id, system))


def test_an_accepted_fact_keeps_where_it_was_read_after_its_suggestion_is_gone() -> None:
    from knowledge_portal.domain.architecture.provenance import AcceptedCitation

    _, decide, draft, candidates, _ = _proposer()
    candidates.replace_proposals(
        _run(draft.id),
        (_candidate(draft.id, CandidateContent(CandidateKind.SYSTEM, "hub", name="Hub")),),
    )

    updated = decide.decide(draft.id, "c1", draft.revision, True, MAINTAINER, NOW)
    # A later reading of the document replaces the draft's suggestions (ADR-0120).
    candidates.replace_proposals(_run(draft.id), ())

    hub = next(item for item in updated.systems if item.id == "hub")
    (entry,) = hub.accepted
    assert entry.citations == (AcceptedCitation("line 1", "quote", "doc", "Doc"),)


def test_accepting_the_ready_ones_takes_only_the_set_the_person_was_shown() -> None:
    _, decide, draft, candidates, _ = _proposer()
    shown = CandidateContent(CandidateKind.SYSTEM, "hub", name="Hub")
    # Arrived after the page listed what it would accept: a reading finishing meanwhile.
    arrived = CandidateContent(CandidateKind.SYSTEM, "gateway", name="Gateway")
    candidates.replace_proposals(
        _run(draft.id), (_candidate(draft.id, shown), _candidate(draft.id, arrived, "c2"))
    )

    updated, remaining = decide.accept_all(
        draft.id, draft.revision, MAINTAINER, NOW, candidate_ids=("c1",)
    )

    assert {item.id for item in updated.systems} >= {"hub"}
    assert "gateway" not in {item.id for item in updated.systems}
    # What is left counts every suggestion still waiting, the unlisted one included.
    assert remaining == 1
    assert (candidates.get("c2") or _candidate(draft.id, arrived)).status.value == "proposed"

    # Naming nothing accepts nothing, and the draft is unchanged.
    same, still = decide.accept_all(draft.id, updated.revision, MAINTAINER, NOW, candidate_ids=())
    assert (same.revision, still) == (updated.revision, 1)


def _run(release_id: str) -> Any:
    from knowledge_portal.application.ports.catalogue_candidates import ExtractionRun

    return ExtractionRun("r", release_id, "doc", "m", "v", 2, (), NOW)


def test_fake_extractor_reads_labelled_lines_and_slugs_names() -> None:
    proposal = FakeCatalogueExtractor().propose(
        ExtractionRequest(
            "Doc",
            (
                ExtractionSegment(
                    1,
                    "p1",
                    "Capability: Orphan (x)\nSystem: Order Hub\nConstraint: Nightly only\n"
                    "Order Hub depends on customer records.\nnoise",
                ),
            ),
            (KnownSystem("bcrm", "BCRM", ("customer records",)),),
        )
    )

    kinds = [item.content.kind for item in proposal.changes]
    assert kinds == [CandidateKind.SYSTEM, CandidateKind.CONSTRAINT, CandidateKind.RELATIONSHIP]
    assert proposal.changes[2].content.target_system_id == "bcrm"
    assert proposal.changes[2].content.text == "Depends on"
    assert slug("   ") == "item"


def test_fake_extractor_skips_table_rows() -> None:
    landscape_row = "System: Order Store | ID: SYS-STORE | Function: Order records"
    component_row = "Component: Fibre Access | Code: PO_FIBRE | Type: Access"
    proposal = FakeCatalogueExtractor().propose(
        ExtractionRequest(
            "Doc",
            (
                ExtractionSegment(
                    1,
                    "row 1",
                    landscape_row,
                    cells=(("System", "Order Store"), ("ID", "SYS-STORE")),
                    read=True,
                ),
                ExtractionSegment(
                    2,
                    "row 2",
                    component_row,
                    cells=(("Component", "Fibre Access"), ("Code", "PO_FIBRE")),
                ),
                # A row whose cells were not kept still reads as a row, not prose.
                ExtractionSegment(3, "row 3", f"{landscape_row}\n{component_row}"),
            ),
            (),
        )
    )

    assert proposal.changes == ()


def test_ai_mapping_of_a_published_release_records_organisation_ownership() -> None:
    repository = seeded_repository()
    index = InMemoryEvidenceIndex(FakeEmbeddings(), FakeWordTokenizer())
    manage = ManageArchitectureKnowledge(
        repository, ADAPTER, index, 10_000, OfflineDocumentScanner()
    )
    draft = manage.create_draft(MAINTAINER, "Next version")
    built = BuildArchitectureIndex(
        manage,
        index,
        InMemoryDocumentStorage(),
        LocatedDocumentExtractor(SafeDocumentTextExtractor()),
        FakeWordTokenizer(),
    ).execute(draft.id, draft.revision, "amina", fence=lambda: None)
    manage.publish(draft.id, built.revision, MAINTAINER, NOW, "Reviewed")
    organisation = InMemoryOrganisationRepository(FixedClock(NOW))
    organisation.change(
        lambda current: (
            current.put_person(Person("bea", "Bea"), None)
            .put_value_stream(ValueStream("retail", "Retail"), None)
            .put_squad(
                Squad(
                    "sales",
                    "Sales",
                    "retail",
                    None,
                    (SquadResource("bcrm", SquadRole.SYSTEM_CONTACT, "bea"),),
                ),
                None,
            )
        ),
        "amina",
        "seed",
        "sales",
    )
    resolver = ResolveArchitectureKnowledge(
        repository,
        index,
        FakeArchitectureReasoner(),
        organisation,
    )

    result = resolver.match(
        ArchitectureQuery(("Update BCRM customer records",), ("BCRM", "Legacy mainframe"))
    )

    assert result.knowledge_version == draft.id
    by_id = {item.id: item for item in result.systems}
    assert [item.name for item in by_id["bcrm"].squads] == ["Sales"]
    assert [item.name for item in by_id["bcrm"].value_streams] == ["Retail"]
    unknown = next(item for item in result.systems if not item.catalogued)
    assert (unknown.name, unknown.squads) == ("Legacy mainframe", ())
    with pytest.raises(KnowledgeConflictError, match="unavailable"):
        resolver.match(ArchitectureQuery(("x",), release_id="missing"))


class _BrokenConnector:
    """A database that refuses every connection."""

    def acquire(self) -> Any:
        raise psycopg.OperationalError("database is down")

    def release(self, connection: Any) -> None:
        return None

    def connection(self, timeout_seconds: float | None = None) -> Any:
        raise psycopg.OperationalError("database is down")

    def close(self) -> None:
        return None


def test_postgres_adapters_report_database_failures_as_persistence_errors() -> None:
    store = PostgresStore(_BrokenConnector())
    organisation = PostgresOrganisationRepository(store, FixedClock(NOW))
    candidates = PostgresCatalogueCandidates(store)
    candidate = _candidate("draft", CandidateContent(CandidateKind.SYSTEM, "hub", name="Hub"))
    failing: list[Any] = [
        organisation.load,
        lambda: organisation.change(lambda current: current, "a", "b", "c"),
        lambda: organisation.audit(5),
        lambda: candidates.replace_proposals(_run("draft"), (candidate,)),
        lambda: candidates.runs("draft"),
        lambda: candidates.list("draft"),
        lambda: candidates.get("c1"),
        lambda: candidates.save_decision(candidate),
    ]
    for call in failing:
        with pytest.raises(PersistenceError):
            call()


def test_extraction_adapter_builds_every_kind_and_skips_unusable_items() -> None:
    from knowledge_portal.infrastructure.llm.catalogue_extraction import (
        ChangeOutput,
        ExtractionOutput,
        StructuredCatalogueExtractor,
    )

    def change(**values: Any) -> ChangeOutput:
        base: dict[str, Any] = {
            "kind": "system",
            "system": "Hub",
            "name": None,
            "aliases": [],
            "triggers": [],
            "target_system": None,
            "component": None,
            "technology": None,
            "domain": None,
            "parent_domain": None,
            "offering": None,
            "journey": None,
            "text": None,
            "evidence_numbers": [1],
            "quote": "Hub takes orders",
            "basis": "stated",
            "reasoning": None,
            "relationship_kind": None,
            "channel_kind": None,
            "interface_style": None,
            "portfolio_level": None,
            "realisation_layer": None,
            "realised_by": [],
            "source_level": None,
            "source_version": None,
            "conflict_sides": [],
            "affected_offerings": [],
        }
        return ChangeOutput(**(base | values))

    class Client:
        model = "m"

        def parse(self, **_: Any) -> ExtractionOutput:
            return ExtractionOutput(
                changes=[
                    change(aliases=["Hub", "Order hub", " "]),
                    change(kind="capability", name="Order taking", triggers=["take order"]),
                    change(kind="constraint", text="Nightly only"),
                    change(kind="capability", name="No phrases", triggers=[]),
                    change(quote="   "),
                ]
            )

    request = ExtractionRequest(
        "Doc", (ExtractionSegment(1, "p1", "The Hub takes orders at night."),), ()
    )
    client: Any = Client()
    proposal = StructuredCatalogueExtractor(client, supports_images=False).propose(request)

    system, capability, constraint = (item.content for item in proposal.changes)
    assert (system.name, system.aliases) == ("Hub", ("Order hub",))
    assert (capability.system_id, capability.capability_id) == ("hub", "order-taking")
    assert constraint.text == "Nightly only"
    assert proposal.warnings == (
        "1 suggestion(s) were left out: what they quote was not found in the document.",
        # A capability with no phrases cannot be suggested.
        "Incomplete suggestions were left out: 1 capability.",
    )


def test_diff_names_every_changed_system_field() -> None:
    from knowledge_portal.domain.architecture.diff import diff_releases

    base = seed_knowledge()
    first = base.systems[0]
    changed = replace(first, aliases=(*first.aliases, "New alias"), constraints=("Batch only",))
    diff = diff_releases(base, replace(base, systems=(changed, *base.systems[1:])))

    assert diff.changes[0].fields == ("aliases", "constraints")
