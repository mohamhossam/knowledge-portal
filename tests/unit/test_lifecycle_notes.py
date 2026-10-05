"""What happens to an offering over its life, as its sources tell it (requirement-portal
ADR-0101, step 4)."""

from __future__ import annotations

import io
from dataclasses import replace

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from knowledge_portal.application.ports.catalogue_file import CatalogueFileFormat
from knowledge_portal.domain.architecture.channels import Channel
from knowledge_portal.domain.architecture.diff import ChangedItem, diff_releases
from knowledge_portal.domain.architecture.knowledge import (
    ArchitectureKnowledge,
    InvalidKnowledgeError,
)
from knowledge_portal.domain.architecture.lifecycle import (
    LifecycleNote,
    NoteBlock,
    NoteBlockKind,
)
from knowledge_portal.domain.architecture.products import (
    OrderType,
    ProductOffering,
    SourceConfidence,
    merge_offerings,
)
from knowledge_portal.infrastructure.architecture.catalogue_files import CatalogueFileAdapter
from knowledge_portal.infrastructure.architecture.knowledge_yaml import seed_knowledge

ADAPTER = CatalogueFileAdapter()
OWNER = {"X-Fake-Actor-Id": "fake-owner"}
SEED = seed_knowledge()
WEB = Channel("web", "B2B Web")
MATRIX = LifecycleNote(
    "LC-UD",
    "Up / Downgrade scenario matrix",
    kind="Change",
    summary="Which workflow each transition runs.",
    order_types=("UPDOWN",),
    blocks=(
        NoteBlock(
            NoteBlockKind.TABLE,
            columns=("From", "To", "Workflow"),
            rows=(("Pro (w/ 5G)", "Pro (w/o 5G)", "WF-1"), ("Office Presence", "Pro")),
            caption="Workflow content is in annexure images.",
            confidence=SourceConfidence.CONFIRMED,
            source="SDD §10",
        ),
        NoteBlock(NoteBlockKind.TEXT, text="Transitions run on the summary order path."),
    ),
    confidence=SourceConfidence.CONFIRMED,
    source="SDD §11.1.3",
)
RENEWAL = LifecycleNote(
    "LC-REN",
    "Renewal",
    kind="Commercial",
    channels=("web",),
    blocks=(
        NoteBlock(
            NoteBlockKind.LIST,
            title="v8.2 carry-over (not in SDD, re-verify)",
            items=("Inherit remaining tenure", "Fresh 24-month commitment"),
            confidence=SourceConfidence.INFERRED,
        ),
    ),
)
OFFERING = ProductOffering(
    "pro",
    "Business Pro",
    order_types=(OrderType("NEW", "New Activation"), OrderType("UPDOWN", "Up / Downgrade")),
    lifecycle_notes=(MATRIX, RENEWAL),
)


def _release(*offerings: ProductOffering, **changes: object) -> ArchitectureKnowledge:
    fields: dict[str, object] = {"channels": (WEB,), "products": offerings or (OFFERING,)}
    return replace(SEED, **{**fields, **changes})  # type: ignore[arg-type]


def test_a_short_table_row_is_padded_to_the_columns() -> None:
    assert MATRIX.blocks[0].rows[1] == ("Office Presence", "Pro", "")


def test_a_note_concerns_its_order_types_and_channels_or_every_one() -> None:
    assert MATRIX.concerns("updown") and not MATRIX.concerns("NEW")
    assert RENEWAL.concerns("NEW", "web") and not RENEWAL.concerns("NEW", "shop")
    assert RENEWAL.concerns("NEW", None)


@pytest.mark.parametrize(
    ("build", "message"),
    [
        (lambda: NoteBlock("chart"), "text, a list or a table"),  # type: ignore[arg-type]
        (lambda: NoteBlock(NoteBlockKind.TEXT), "A text block needs its text"),
        (lambda: NoteBlock(NoteBlockKind.LIST), "A list needs at least one item"),
        (lambda: NoteBlock(NoteBlockKind.TABLE, columns=("", " ")), "needs its column heads"),
        (
            lambda: NoteBlock(NoteBlockKind.TABLE, columns=("A",), rows=(("1", "2"),)),
            "a table row has 2 cells but the table 1 columns",
        ),
        (lambda: NoteBlock(NoteBlockKind.TABLE, columns=tuple("ABCDEFGHIJK")), "at most 10"),
        (lambda: NoteBlock(NoteBlockKind.TEXT, text="x", columns=("A",)), "Only a table"),
        (lambda: LifecycleNote("x", "Empty"), "a note says something"),
        (
            lambda: replace(OFFERING, lifecycle_notes=(MATRIX, MATRIX)),
            "lifecycle note ids must be unique",
        ),
        (
            lambda: replace(OFFERING, lifecycle_notes=(replace(MATRIX, order_types=("CEASE",)),)),
            "order type 'CEASE' is not one of the offering's order types",
        ),
        (lambda: _release(channels=()), "Renewal names channel 'web'"),
    ],
)
def test_what_a_lifecycle_note_refuses(build: object, message: str) -> None:
    with pytest.raises(InvalidKnowledgeError, match=f"(?i){message}"):
        build()  # type: ignore[operator]


@pytest.mark.parametrize("file_format", list(CatalogueFileFormat))
def test_every_catalogue_file_format_round_trips_lifecycle_notes(
    file_format: CatalogueFileFormat,
) -> None:
    release = _release()

    content = ADAPTER.read(file_format, ADAPTER.write(file_format, release))

    assert content.products == release.products


def test_a_workbook_row_must_follow_its_table() -> None:
    book = load_workbook(io.BytesIO(ADAPTER.write(CatalogueFileFormat.XLSX, _release())))
    sheet = book["LifecycleBlocks"]
    sheet.insert_rows(2)
    for column, value in enumerate(("pro", "LC-UD", "row", None, None, "Orphan"), 1):
        sheet.cell(row=2, column=column, value=value)
    changed = io.BytesIO()
    book.save(changed)

    with pytest.raises(InvalidKnowledgeError, match="a table row must follow its table"):
        ADAPTER.read(CatalogueFileFormat.XLSX, changed.getvalue())


def test_a_second_reading_adds_notes_but_never_replaces_one() -> None:
    later = replace(
        OFFERING,
        lifecycle_notes=(
            replace(MATRIX, title="Another reading"),
            LifecycleNote("LC-CEASE", "Cessation", summary="Blocked while activating."),
        ),
    )

    merged = merge_offerings(OFFERING, later)

    assert [item.title for item in merged.lifecycle_notes] == [
        "Up / Downgrade scenario matrix",
        "Renewal",
        "Cessation",
    ]


def test_the_diff_names_lifecycle_notes() -> None:
    after = replace(OFFERING, lifecycle_notes=(MATRIX,))

    (change,) = [
        item
        for item in diff_releases(_release(), _release(after)).changes
        if item.item is ChangedItem.PRODUCT
    ]

    assert change.fields == ("lifecycle_notes",)


def test_a_draft_saves_lifecycle_notes_through_the_api(client: TestClient) -> None:
    draft = client.post(
        "/architecture-knowledge/releases", json={"name": "Next version"}, headers=OWNER
    ).json()
    url = f"/architecture-knowledge/releases/{draft['id']}"
    body = {"systems": draft["systems"], "relationships": draft["relationships"]}
    note = {
        "id": "LC-CEASE",
        "title": "Cessation",
        "order_types": ["NEW"],
        "blocks": [
            {"kind": "table", "columns": ["When", "Then"], "rows": [["Activating", "Blocked"]]},
            {"kind": "list", "items": ["Cancel the in-flight order first"]},
        ],
    }
    offering = {
        "id": "pro",
        "name": "Business Pro",
        "order_types": [{"code": "NEW", "name": "New Activation"}],
        "lifecycle_notes": [note],
    }

    saved = client.put(
        url,
        json={**body, "expected_revision": draft["revision"], "products": [offering]},
        headers=OWNER,
    )

    assert saved.status_code == 200, saved.text
    (stored,) = saved.json()["products"][0]["lifecycle_notes"]
    assert stored["blocks"][0]["rows"] == [["Activating", "Blocked"]]
    assert stored["blocks"][1]["items"] == ["Cancel the in-flight order first"]

    refused = client.put(
        url,
        json={
            **body,
            "expected_revision": saved.json()["revision"],
            "products": [{**offering, "lifecycle_notes": [{**note, "blocks": [{"kind": "list"}]}]}],
        },
        headers=OWNER,
    )
    assert refused.status_code == 422
    assert "A list needs at least one item" in refused.text
