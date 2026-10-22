"""Excel, YAML and JSON catalogue files, and the changes view reviewed before publishing."""

from __future__ import annotations

import io
from dataclasses import replace
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from knowledge_portal.application.ports.catalogue_file import CatalogueFileFormat
from knowledge_portal.domain.architecture.diff import ChangedItem, ChangeKind, diff_releases
from knowledge_portal.domain.architecture.knowledge import (
    InvalidKnowledgeError,
    KnowledgeCapability,
    LandscapeDomain,
    SystemDefinition,
    SystemRelationship,
)
from knowledge_portal.infrastructure.architecture.catalogue_files import CatalogueFileAdapter
from knowledge_portal.infrastructure.architecture.knowledge_yaml import (
    knowledge_from_yaml,
    seed_knowledge,
)

ADAPTER = CatalogueFileAdapter()
OWNER = {"X-Fake-Actor-Id": "fake-owner"}


@pytest.mark.parametrize("file_format", list(CatalogueFileFormat))
def test_every_format_round_trips_the_catalogue(file_format: CatalogueFileFormat) -> None:
    seed = replace(
        seed_knowledge(),
        systems=(
            replace(seed_knowledge().systems[0], constraints=("Read-only on weekends",)),
            *seed_knowledge().systems[1:],
        ),
    )

    content = ADAPTER.read(file_format, ADAPTER.write(file_format, seed))

    assert content.systems == seed.systems
    assert content.relationships == seed.relationships


CATALOGUE = Path(__file__).resolve().parents[2] / "catalogues" / "smb-architecture.yaml"


@pytest.mark.parametrize("file_format", list(CatalogueFileFormat))
def test_every_format_keeps_delivery_drops_and_cross_cutting_domains(
    file_format: CatalogueFileFormat,
) -> None:
    catalogue = knowledge_from_yaml(CATALOGUE.read_text(encoding="utf-8"))
    orders = {item.code: item for item in catalogue.products[0].order_types}
    assert orders["NEW"].delivery_drop == "P1"
    assert orders["UPDOWNGRD"].delivery_drop == "P2 / P3"
    assert "Delivered in drop" not in (orders["UPDOWNGRD"].description or "")
    crossing = [item.id for item in catalogue.landscape_domains if item.cross_cutting]
    assert crossing == ["enterprise", "integration"]

    content = ADAPTER.read(file_format, ADAPTER.write(file_format, catalogue))

    assert content.products[0].order_types == catalogue.products[0].order_types
    assert content.landscape_domains == catalogue.landscape_domains


def _workbook(**sheets: list[tuple[object, ...]]) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.worksheets[0])
    for name, rows in sheets.items():
        sheet = workbook.create_sheet(name)
        for row in rows:
            sheet.append(list(row))
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_excel_reads_separated_lists_and_ignores_blank_rows() -> None:
    content = ADAPTER.read(
        CatalogueFileFormat.XLSX,
        _workbook(
            Systems=[
                ("system_id", "name", "name_ar", "aliases"),
                ("crm", "CRM", None, "customer records; accounts"),
                (None, None, None, None),
                ("billing", "Billing", "الفوترة", None),
            ],
            Capabilities=[
                ("system_id", "capability_id", "name", "triggers"),
                ("crm", "lookup", "Customer lookup", "find customer;search account"),
            ],
            Relationships=[
                ("source_system_id", "target_system_id", "description"),
                ("crm", "billing", "Sends invoices"),
            ],
        ),
    )

    assert content.systems == (
        SystemDefinition(
            "crm",
            "CRM",
            ("customer records", "accounts"),
            (
                KnowledgeCapability(
                    "lookup", "Customer lookup", ("find customer", "search account")
                ),
            ),
        ),
        SystemDefinition("billing", "Billing", name_ar="الفوترة"),
    )
    assert content.relationships == (SystemRelationship("crm", "billing", "Sends invoices"),)


@pytest.mark.parametrize(
    ("sheets", "message"),
    [
        ({"Other": [("x",)]}, "needs a 'Systems' sheet"),
        ({"Systems": [("id", "title")]}, "Systems row 1 must have the headers"),
        (
            {"Systems": [("system_id", "name"), ("crm", None)]},
            "Systems row 2: name is required",
        ),
        (
            {
                "Systems": [("system_id", "name"), ("crm", "CRM")],
                "Capabilities": [
                    ("system_id", "capability_id", "name", "triggers"),
                    ("erp", "post", "Post", "post"),
                ],
            },
            "Capabilities row 2: system 'erp' is not listed",
        ),
        (
            {
                "Systems": [("system_id", "name"), ("crm", "CRM")],
                "Capabilities": [
                    ("system_id", "capability_id", "name", "triggers"),
                    ("crm", "lookup", "Lookup", None),
                ],
            },
            "Capabilities row 2: A capability needs non-empty matching phrases",
        ),
    ],
)
def test_excel_errors_name_the_sheet_and_row(
    sheets: dict[str, list[tuple[object, ...]]], message: str
) -> None:
    with pytest.raises(InvalidKnowledgeError, match=message):
        ADAPTER.read(CatalogueFileFormat.XLSX, _workbook(**sheets))


@pytest.mark.parametrize("column", ["flows", "channels", "milestones", "statuses", "fallout"])
def test_excel_tracking_ignores_a_column_named_for_one_of_its_lists(column: str) -> None:
    # The tracking sheets append to these lists; a Tracking cell under the same name was
    # spread over the list and the import crashed instead of reading the workbook.
    content = ADAPTER.read(
        CatalogueFileFormat.XLSX,
        _workbook(
            Systems=[("system_id", "name"), ("crm", "CRM"), ("om", "Order management")],
            Products=[("product_id", "name"), ("fibre", "Fibre")],
            OrderTypes=[("product_id", "code", "name"), ("fibre", "NEW", "New connection")],
            Tracking=[("product_id", "order_types", column), ("fibre", "NEW", "stray")],
            TrackingFlows=[
                ("product_id", "from_system_id", "to_system_id", "label"),
                ("fibre", "crm", "om", "Order placed"),
            ],
            TrackingChannels=[("product_id", "channel_id"), ("fibre", "web")],
            TrackingEvents=[
                ("product_id", "kind", "label"),
                ("fibre", "milestone", "Installed"),
                ("fibre", "status", "In progress"),
                ("fibre", "fallout", "Address not found"),
            ],
        ),
    )

    tracking = content.products[0].tracking
    assert tracking is not None
    assert tracking.order_types == ("NEW",)
    assert [flow.label for flow in tracking.flows] == ["Order placed"]
    assert [channel.channel_id for channel in tracking.channels] == ["web"]
    assert [event.label for event in tracking.milestones] == ["Installed"]
    assert [event.label for event in tracking.statuses] == ["In progress"]
    assert [case.trigger for case in tracking.fallout] == ["Address not found"]


def test_text_formats_report_syntax_and_shape_errors() -> None:
    with pytest.raises(InvalidKnowledgeError, match="JSON file could not be parsed: line 1"):
        ADAPTER.read(CatalogueFileFormat.JSON, b"{nope")
    with pytest.raises(InvalidKnowledgeError, match="systems entry 1: name is required"):
        ADAPTER.read(CatalogueFileFormat.YAML, b"systems:\n  - id: crm\n")
    with pytest.raises(InvalidKnowledgeError, match="not an .xlsx"):
        ADAPTER.read(CatalogueFileFormat.XLSX, b"plain text")
    with pytest.raises(InvalidKnowledgeError, match=r"\.xlsx"):
        CatalogueFileFormat.from_filename("catalogue.csv")


def test_template_has_instructions_and_every_sheet() -> None:
    workbook = load_workbook(io.BytesIO(ADAPTER.template()))

    assert workbook.sheetnames == [
        "Instructions",
        "Systems",
        "Components",
        "Capabilities",
        "Constraints",
        "Relationships",
        "Domains",
        "LandscapeDomains",
        "Products",
        "OrderTypes",
        "OfferingComponents",
        "Responsibilities",
        "ProductPoints",
        "Journeys",
        "Activities",
        "FlowRules",
        "ActivityIntegrations",
        "Channels",
        "Realisation",
        "NFRs",
        "Tracking",
        "TrackingFlows",
        "TrackingChannels",
        "TrackingEvents",
        "LifecycleNotes",
        "LifecycleBlocks",
        "Sources",
        "Conflicts",
        "ConflictScopes",
        "Questions",
        "Decisions",
        "Boundaries",
        "ChangeHistory",
        "ChangeItems",
        "Portfolio",
        "Plans",
        "PlanCharacteristics",
        "BusinessRules",
        "Concepts",
        "Schemes",
        "Vocabulary",
        "Interfaces",
        "RealisationRecords",
        "MasteryNotes",
        "Acceptances",
    ]
    assert ADAPTER.read(CatalogueFileFormat.XLSX, ADAPTER.template()).systems == ()


def test_diff_lists_added_changed_and_removed_items() -> None:
    base = seed_knowledge()
    first, second, *rest = base.systems
    changed_first = replace(
        first,
        name="Renamed",
        capabilities=(replace(first.capabilities[0], triggers=("new phrase",)),),
    )
    draft = replace(
        base,
        id="draft",
        systems=(changed_first, *rest, SystemDefinition("new", "New system")),
        relationships=base.relationships[1:],
    )

    diff = diff_releases(base, draft)

    found = {(item.item, item.change, item.key): item for item in diff.changes}
    assert found[(ChangedItem.SYSTEM, ChangeKind.CHANGED, first.id)].fields == ("name",)
    assert (ChangedItem.SYSTEM, ChangeKind.ADDED, "new") in found
    assert (ChangedItem.SYSTEM, ChangeKind.REMOVED, second.id) in found
    capability_key = f"{first.id}/{first.capabilities[0].id}"
    assert found[(ChangedItem.CAPABILITY, ChangeKind.CHANGED, capability_key)].fields == (
        "triggers",
    )
    assert sum(item.item is ChangedItem.RELATIONSHIP for item in diff.changes) >= 1
    assert diff_releases(base, base).empty


def test_diff_names_a_landscape_domain_made_cross_cutting() -> None:
    base = replace(seed_knowledge(), landscape_domains=(LandscapeDomain("integration", "Bus"),))
    draft = replace(
        base,
        id="draft",
        landscape_domains=(LandscapeDomain("integration", "Bus", cross_cutting=True),),
    )

    changes = diff_releases(base, draft).changes

    assert [(item.item, item.key, item.fields) for item in changes] == [
        (ChangedItem.LANDSCAPE_DOMAIN, "integration", ("cross_cutting",))
    ]


def test_changes_and_template_endpoints_are_for_maintainers(client: TestClient) -> None:
    draft = client.post(
        "/architecture-knowledge/releases", json={"name": "Next version"}, headers=OWNER
    ).json()
    workbook = _workbook(
        Systems=[("system_id", "name"), ("crm", "CRM")],
    )

    preview = client.post(
        f"/architecture-knowledge/releases/{draft['id']}/catalogue-file/preview",
        files={"file": ("catalogue.xlsx", workbook, "application/octet-stream")},
        headers=OWNER,
    )
    applied = client.post(
        f"/architecture-knowledge/releases/{draft['id']}/catalogue-file",
        data={"expected_revision": draft["revision"]},
        files={"file": ("catalogue.xlsx", workbook, "application/octet-stream")},
        headers=OWNER,
    )
    changes = client.get(
        f"/architecture-knowledge/releases/{draft['id']}/changes", headers=OWNER
    ).json()

    assert preview.status_code == 200
    assert {"item": "system", "change": "added", "key": "crm", "label": "CRM", "fields": []} in (
        preview.json()["changes"]
    )
    assert applied.json()["systems"][0]["id"] == "crm"
    assert changes["changes"] == preview.json()["changes"]
    template = client.get("/architecture-knowledge/catalogue-template.xlsx", headers=OWNER)
    assert template.status_code == 200
    assert 'filename="catalogue-template.xlsx"' in template.headers["content-disposition"]
    reader = {"X-Fake-Actor-Id": "fake-reviewer"}
    assert (
        client.get("/architecture-knowledge/catalogue-template.xlsx", headers=reader).status_code
        == 403
    )
