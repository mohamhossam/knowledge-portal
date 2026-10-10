"""Business capability concepts: the scheme, its links, the backfill and their review (ADR-0114)."""

from __future__ import annotations

import io
import json
from dataclasses import replace
from datetime import UTC, datetime
from pathlib import Path

import pytest
import yaml
from openpyxl import load_workbook

from knowledge_portal.application.ports.catalogue_file import CatalogueFileFormat
from knowledge_portal.domain.architecture.candidates import (
    CandidateBasis,
    CandidateCitation,
    CandidateContent,
    CandidateDependencyError,
    CandidateKind,
    CandidateMatch,
    CatalogueCandidate,
    apply_candidate,
    classify,
    needs_one_by_one,
)
from knowledge_portal.domain.architecture.concept_backfill import propose_concepts
from knowledge_portal.domain.architecture.concepts import (
    BusinessCapability,
    CapabilityRef,
    concept_id_for,
)
from knowledge_portal.domain.architecture.diff import ChangedItem, ChangeKind, diff_releases
from knowledge_portal.domain.architecture.knowledge import (
    ArchitectureKnowledge,
    CapabilityDomain,
    InvalidKnowledgeError,
    KnowledgeCapability,
    SystemDefinition,
)
from knowledge_portal.domain.architecture.products import (
    ComponentResponsibility,
    OfferingComponent,
    ProductOffering,
)
from knowledge_portal.infrastructure.architecture.catalogue_files import CatalogueFileAdapter
from knowledge_portal.infrastructure.architecture.knowledge_yaml import (
    knowledge_from_yaml,
    seed_knowledge,
)

ROOT = Path(__file__).resolve().parents[2]
CATALOGUE = ROOT / "catalogues" / "smb-architecture.yaml"
GOLDEN = ROOT / "tests" / "fixtures" / "golden" / "impact_golden_set.json"
ADAPTER = CatalogueFileAdapter()
AT = datetime(2026, 10, 10, 12, tzinfo=UTC)

WIFI = BusinessCapability(
    "cap-wifi",
    "Managed Wi-Fi",
    ("WiFi", "access point"),
    "Wireless access the operator manages.",
    domain_id="service",
)


def _release(**values: object) -> ArchitectureKnowledge:
    fields: dict[str, object] = {
        "id": "draft",
        "revision": 1,
        "systems": (
            SystemDefinition(
                "cwom",
                "CWOM",
                capabilities=(
                    KnowledgeCapability("wifi", "Managed Wi-Fi", ("wireless",), "service"),
                    KnowledgeCapability("orders", "Order orchestration", ("orchestrate",)),
                ),
            ),
            SystemDefinition(
                "e2eso",
                "E2ESO",
                capabilities=(
                    KnowledgeCapability("wifi-config", "managed wi-fi", ("WLAN",), "service"),
                ),
            ),
        ),
        "relationships": (),
        "capability_domains": (CapabilityDomain("service", "Service"),),
        "products": (
            ProductOffering(
                "bpp",
                "Business Pro Plus",
                components=(
                    OfferingComponent(
                        "ap",
                        "Access point",
                        description="Managed WiFi for the office.",
                        responsibilities=(
                            ComponentResponsibility("cwom", "PROVISIONS", "Sets up"),
                        ),
                    ),
                ),
            ),
        ),
    }
    fields.update(values)
    return ArchitectureKnowledge(**fields)  # type: ignore[arg-type]


# The scheme ---------------------------------------------------------------------------------


def test_a_concept_needs_distinct_labels_and_one_place_in_the_scheme() -> None:
    with pytest.raises(InvalidKnowledgeError, match="each label must differ"):
        BusinessCapability("cap-a", "Wi-Fi", ("wifi",))
    with pytest.raises(InvalidKnowledgeError, match="only a top concept names its domain"):
        BusinessCapability("cap-a", "Wi-Fi", broader_id="cap-b", domain_id="service")
    with pytest.raises(InvalidKnowledgeError, match="broader than itself"):
        BusinessCapability("cap-a", "Wi-Fi", broader_id="cap-a")
    with pytest.raises(InvalidKnowledgeError, match="Concept label must not be blank"):
        BusinessCapability("cap-a", " ")
    assert concept_id_for("Managed Wi-Fi access points", {"cap-managed-wi-fi-access-points"}) == (
        "cap-managed-wi-fi-access-points-2"
    )


@pytest.mark.parametrize(
    ("concepts", "message"),
    [
        ((WIFI, replace(WIFI, pref_label="Other")), "Concept ids must be unique"),
        ((WIFI, BusinessCapability("cap-b", "Access point")), "already names Managed Wi-Fi"),
        ((replace(WIFI, domain_id="nowhere"),), "which is not in the catalogue"),
        (
            (WIFI, BusinessCapability("cap-b", "Indoor Wi-Fi", broader_id="cap-missing")),
            "narrows 'cap-missing'",
        ),
        (
            (
                BusinessCapability("cap-a", "A", broader_id="cap-b"),
                BusinessCapability("cap-b", "B", broader_id="cap-a"),
            ),
            "narrower than itself",
        ),
        (
            (
                WIFI,
                BusinessCapability("cap-b", "B", broader_id="cap-wifi"),
                BusinessCapability("cap-c", "C", broader_id="cap-b"),
                BusinessCapability("cap-d", "D", broader_id="cap-c"),
            ),
            "at most 3 levels",
        ),
    ],
)
def test_the_release_refuses_a_broken_scheme(
    concepts: tuple[BusinessCapability, ...], message: str
) -> None:
    with pytest.raises(InvalidKnowledgeError, match=message):
        _release(business_capabilities=concepts)


def test_links_name_concepts_in_the_release_or_say_why_none_fits() -> None:
    release = _release(business_capabilities=(WIFI,))
    system = release.systems[0]
    linked = replace(system.capabilities[0], concept_id="cap-missing")
    with pytest.raises(InvalidKnowledgeError, match="linked to concept 'cap-missing'"):
        release.updated(systems=(replace(system, capabilities=(linked,)),))
    offering = release.products[0]
    part = replace(offering.components[0], capability_ids=("cap-missing",))
    with pytest.raises(InvalidKnowledgeError, match="Access point is linked to concept"):
        release.updated(products=(replace(offering, components=(part,)),))
    with pytest.raises(InvalidKnowledgeError, match="cannot also be marked"):
        KnowledgeCapability("x", "X", ("x",), concept_id="cap-wifi", unlinked_reason="None fits")
    with pytest.raises(InvalidKnowledgeError, match="cannot also be marked"):
        OfferingComponent("x", "X", capability_ids=("cap-wifi",), unlinked_reason="None fits")
    with pytest.raises(InvalidKnowledgeError, match="more than once"):
        OfferingComponent("x", "X", capability_ids=("cap-wifi", "cap-wifi"))


def test_a_concept_sits_in_its_top_concept_s_domain() -> None:
    narrower = BusinessCapability("cap-indoor", "Indoor Wi-Fi", broader_id="cap-wifi")
    release = _release(business_capabilities=(WIFI, narrower))

    assert [item.id for item in release.concept_path("cap-indoor")] == ["cap-wifi", "cap-indoor"]
    domain = release.concept_domain("cap-indoor")
    assert domain is not None and domain.id == "service"
    assert release.concept_path("cap-missing") == ()


# The backfill -------------------------------------------------------------------------------


def test_the_backfill_merges_one_name_across_systems_and_reuses_a_concept_s_label() -> None:
    release = _release(business_capabilities=(BusinessCapability("cap-orch", "Orchestrate"),))

    proposals = {item.concept.pref_label: item for item in propose_concepts(release)}

    wifi = proposals["Managed Wi-Fi"]
    assert wifi.capabilities == (
        CapabilityRef("cwom", "wifi"),
        CapabilityRef("e2eso", "wifi-config"),
    )
    assert wifi.concept.id == "cap-managed-wi-fi"
    assert wifi.concept.domain_id == "service"
    assert wifi.concept.alt_labels == ("wireless", "WLAN")
    assert not wifi.existing
    # "orchestrate" already names a concept, so it is not taken as another label.
    assert proposals["Order orchestration"].concept.alt_labels == ()


def test_the_backfill_proposes_a_concept_for_every_capability_of_the_seed() -> None:
    seed = seed_knowledge()
    draft = replace(seed, status=seed.status.DRAFT, published_at=None, published_by=None)
    proposals = propose_concepts(draft)

    merged = draft
    for proposal in proposals:
        content = CandidateContent(
            CandidateKind.CONCEPT,
            proposal.concept.id,
            concept=proposal.concept,
            capability_refs=proposal.capabilities,
        )
        assert classify(content, merged) is CandidateMatch.NEW
        merged = apply_candidate(content, merged)

    capabilities = [item for system in merged.systems for item in system.capabilities]
    assert len(merged.business_capabilities) == len(capabilities) == 31
    assert all(item.concept_id for item in capabilities)
    assert all(item.domain_id for item in merged.business_capabilities)
    assert propose_concepts(merged) == ()


# Concept suggestions ------------------------------------------------------------------------


def _concept(concept: BusinessCapability, *refs: CapabilityRef) -> CandidateContent:
    return CandidateContent(
        CandidateKind.CONCEPT, concept.id, concept=concept, capability_refs=refs
    )


def test_accepting_a_concept_links_only_capabilities_nobody_has_decided() -> None:
    marked = replace(
        _release().systems[1],
        capabilities=(replace(_release().systems[1].capabilities[0], unlinked_reason="Config"),),
    )
    release = _release(systems=(_release().systems[0], marked))
    content = _concept(WIFI, CapabilityRef("cwom", "wifi"), CapabilityRef("e2eso", "wifi-config"))

    assert classify(content, release) is CandidateMatch.NEW
    merged = apply_candidate(content, release)

    assert merged.systems[0].capabilities[0].concept_id == "cap-wifi"
    assert merged.systems[1].capabilities[0].concept_id is None
    assert classify(content, merged) is CandidateMatch.ALREADY_PRESENT
    # A second suggestion for the same concept adds labels and fills gaps, never moves it.
    other = replace(
        WIFI, id="cap-other", alt_labels=("wireless LAN",), domain_id=None, exact_match="ODA-1"
    )
    assert classify(_concept(other), merged) is CandidateMatch.UPDATES_EXISTING
    again = apply_candidate(_concept(other), merged).business_capabilities[0]
    assert again.id == "cap-wifi"
    assert again.alt_labels == ("WiFi", "access point", "wireless LAN")
    assert (again.domain_id, again.exact_match) == ("service", "ODA-1")


def test_a_concept_suggestion_waits_for_what_it_names() -> None:
    release = _release()
    assert classify(_concept(WIFI, CapabilityRef("ghost", "x")), release) is (
        CandidateMatch.NEEDS_SYSTEM
    )
    narrower = BusinessCapability("cap-indoor", "Indoor Wi-Fi", broader_id="cap-wifi")
    assert classify(_concept(narrower), release) is CandidateMatch.NEEDS_CONCEPT
    with pytest.raises(CandidateDependencyError, match="before the narrower"):
        apply_candidate(_concept(narrower), release)
    elsewhere = replace(WIFI, domain_id="nowhere")
    assert classify(_concept(elsewhere), release) is CandidateMatch.NEEDS_DOMAIN
    with pytest.raises(CandidateDependencyError, match="Add capability domain"):
        apply_candidate(_concept(elsewhere), release)


def test_component_links_are_inferred_and_decided_one_by_one() -> None:
    release = _release(business_capabilities=(WIFI,))
    link = CandidateContent(
        CandidateKind.COMPONENT_LINK, "bpp", component_id="ap", concept_ids=("WiFi",)
    )
    assert classify(link, release) is CandidateMatch.NEW
    assert classify(replace(link, component_id="ghost"), release) is CandidateMatch.NEEDS_OFFERING
    unknown = replace(link, concept_ids=("cap-ghost",))
    assert classify(unknown, release) is CandidateMatch.NEEDS_CONCEPT
    with pytest.raises(CandidateDependencyError, match="cap-ghost"):
        apply_candidate(unknown, release)

    merged = apply_candidate(link, release)
    assert merged.products[0].components[0].capability_ids == ("cap-wifi",)
    assert classify(link, merged) is CandidateMatch.ALREADY_PRESENT

    candidate = CatalogueCandidate(
        "c1",
        "draft",
        "component_links",
        link,
        (CandidateCitation("Catalogue › Business Pro Plus › component ap", "Managed WiFi"),),
        "model",
        "capability-links-v1",
        AT,
        basis=CandidateBasis.INFERRED,
        rationale="It is the Wi-Fi access point.",
    )
    assert needs_one_by_one(candidate, release)
    with pytest.raises(InvalidKnowledgeError, match="at least one concept"):
        CandidateContent(CandidateKind.COMPONENT_LINK, "bpp", component_id="ap")
    with pytest.raises(InvalidKnowledgeError, match="Only a component link names concepts"):
        CandidateContent(CandidateKind.SYSTEM, "x", name="X", concept_ids=("cap-wifi",))
    with pytest.raises(InvalidKnowledgeError, match="Only a concept suggestion holds"):
        CandidateContent(CandidateKind.SYSTEM, "x", name="X", concept=WIFI)


def test_accepting_a_component_link_overrides_none_fits() -> None:
    offering = _release().products[0]
    marked = replace(offering.components[0], unlinked_reason="Hardware only")
    release = _release(
        business_capabilities=(WIFI,), products=(replace(offering, components=(marked,)),)
    )
    link = CandidateContent(
        CandidateKind.COMPONENT_LINK, "bpp", component_id="ap", concept_ids=("cap-wifi",)
    )
    part = apply_candidate(link, release).products[0].components[0]
    assert (part.capability_ids, part.unlinked_reason) == (("cap-wifi",), None)


def test_a_document_s_reading_of_an_offering_keeps_its_component_links() -> None:
    offering = _release().products[0]
    linked = replace(offering.components[0], capability_ids=("cap-wifi",))
    release = _release(
        business_capabilities=(WIFI,), products=(replace(offering, components=(linked,)),)
    )
    reading = replace(offering, components=(replace(offering.components[0], name="AP"),))
    content = CandidateContent(CandidateKind.PRODUCT, "bpp", product=reading)

    merged = apply_candidate(content, release)

    assert merged.products[0].components[0].name == "AP"
    assert merged.products[0].components[0].capability_ids == ("cap-wifi",)


# Review before publishing -------------------------------------------------------------------


def test_the_diff_reports_concepts_and_links() -> None:
    base = _release(business_capabilities=(WIFI,))
    system = base.systems[0]
    offering = base.products[0]
    draft = base.updated(
        business_capabilities=(
            replace(WIFI, definition="Wi-Fi in the office."),
            BusinessCapability("cap-orch", "Orchestration"),
        ),
        systems=(
            replace(
                system,
                capabilities=(
                    replace(system.capabilities[0], concept_id="cap-wifi"),
                    system.capabilities[1],
                ),
            ),
            base.systems[1],
        ),
        products=(
            replace(
                offering,
                components=(replace(offering.components[0], capability_ids=("cap-wifi",)),),
            ),
        ),
    )

    diff = diff_releases(base, draft)
    changes = {(item.item, item.change, item.key): item for item in diff.changes}

    assert (ChangedItem.CONCEPT, ChangeKind.ADDED, "cap-orch") in changes
    assert changes[(ChangedItem.CONCEPT, ChangeKind.CHANGED, "cap-wifi")].fields == ("definition",)
    assert changes[(ChangedItem.CAPABILITY, ChangeKind.CHANGED, "cwom/wifi")].fields == ("concept",)
    assert changes[(ChangedItem.PRODUCT, ChangeKind.CHANGED, "bpp")].fields == ("capability_links",)


# Catalogue files ----------------------------------------------------------------------------


def _linked_release() -> ArchitectureKnowledge:
    base = _release(business_capabilities=(WIFI,))
    system = base.systems[0]
    offering = base.products[0]
    return base.updated(
        systems=(
            replace(
                system,
                capabilities=(
                    replace(system.capabilities[0], concept_id="cap-wifi"),
                    replace(system.capabilities[1], unlinked_reason="Shared by every offering"),
                ),
            ),
            base.systems[1],
        ),
        products=(
            replace(
                offering,
                components=(replace(offering.components[0], capability_ids=("cap-wifi",)),),
            ),
        ),
    )


@pytest.mark.parametrize("file_format", list(CatalogueFileFormat))
def test_every_catalogue_file_carries_concepts_and_links(file_format: CatalogueFileFormat) -> None:
    release = _linked_release()

    content = ADAPTER.read(file_format, ADAPTER.write(file_format, release))

    assert content.business_capabilities == release.business_capabilities
    assert content.systems == release.systems
    assert content.products[0].components[0].capability_ids == ("cap-wifi",)


def test_a_file_names_a_concept_it_does_not_list_by_where() -> None:
    raw = yaml.safe_load(ADAPTER.write(CatalogueFileFormat.YAML, _linked_release()))
    raw["business_capabilities"] = []
    with pytest.raises(InvalidKnowledgeError, match="cwom, capability wifi: concept 'cap-wifi'"):
        ADAPTER.read(CatalogueFileFormat.YAML, yaml.safe_dump(raw).encode())
    raw = yaml.safe_load(ADAPTER.write(CatalogueFileFormat.YAML, _linked_release()))
    raw["business_capabilities"][0]["alt_labels"] = ["Managed Wi-Fi"]
    with pytest.raises(InvalidKnowledgeError, match="business_capabilities entry 1"):
        ADAPTER.read(CatalogueFileFormat.YAML, yaml.safe_dump(raw).encode())


def test_the_workbook_has_a_concepts_sheet_and_link_columns() -> None:
    workbook = load_workbook(io.BytesIO(ADAPTER.write(CatalogueFileFormat.XLSX, _linked_release())))
    concepts = list(workbook["Concepts"].iter_rows(values_only=True))
    assert concepts[1][:3] == ("cap-wifi", "Managed Wi-Fi", "WiFi; access point")
    capabilities = workbook["Capabilities"]
    headers = [cell.value for cell in capabilities[1]]
    assert headers[-2:] == ["concept_id", "unlinked_reason"]


# The committed catalogue --------------------------------------------------------------------


def test_the_committed_catalogue_links_every_component_and_holds_the_golden_concepts() -> None:
    catalogue = knowledge_from_yaml(CATALOGUE.read_text(encoding="utf-8"))
    golden = json.loads(GOLDEN.read_text(encoding="utf-8"))

    parts = [part for offering in catalogue.products for part in offering.components]
    capabilities = [item for system in catalogue.systems for item in system.capabilities]
    assert parts and all(part.capability_ids or part.unlinked_reason for part in parts)
    assert all(item.concept_id or item.unlinked_reason for item in capabilities)
    assert {item["id"] for item in golden["concepts"]} <= {
        item.id for item in catalogue.business_capabilities
    }
    assert all(item.domain_id for item in catalogue.business_capabilities)
    assert propose_concepts(catalogue) == ()
