"""A release, or the blank template, written as the Excel workbook."""

from __future__ import annotations

import io
import json
from collections.abc import Sequence
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font

from knowledge_portal.domain.architecture.change_requests import ChangeRequestRecord
from knowledge_portal.domain.architecture.channels import Channel
from knowledge_portal.domain.architecture.concepts import BusinessCapability
from knowledge_portal.domain.architecture.governance import KnowledgeSource, SourceConflict
from knowledge_portal.domain.architecture.interfaces import SystemInterface
from knowledge_portal.domain.architecture.journeys import Journey
from knowledge_portal.domain.architecture.knowledge import (
    ArchitectureKnowledge,
    CapabilityDomain,
    LandscapeDomain,
    SystemDefinition,
    SystemRelationship,
)
from knowledge_portal.domain.architecture.lifecycle import LifecycleNote, NoteBlockKind
from knowledge_portal.domain.architecture.portfolio import PortfolioNode
from knowledge_portal.domain.architecture.products import OfferingComponent, ProductOffering
from knowledge_portal.domain.architecture.realisations import RealisationRecord
from knowledge_portal.domain.architecture.tracking import OrderTracking
from knowledge_portal.domain.architecture.vocabularies import ConceptScheme, VocabularyTerm
from knowledge_portal.infrastructure.architecture.catalogue_files.fields import _LIST_SEPARATOR
from knowledge_portal.infrastructure.architecture.catalogue_files.mapping_writing import (
    _acceptance_rows,
    _iso,
)
from knowledge_portal.infrastructure.architecture.catalogue_files.workbook_sheets import (
    _CELLS,
    _HEADERS,
    _INSTRUCTIONS,
    ACCEPTANCES,
    ACTIVITIES,
    ACTIVITY_INTEGRATIONS,
    BOUNDARIES,
    BUSINESS_RULES,
    CAPABILITIES,
    CHANGE_HISTORY,
    CHANGE_ITEMS,
    CHANNELS,
    COMPONENTS,
    CONCEPTS,
    CONFLICT_SCOPES,
    CONFLICTS,
    CONSTRAINTS,
    DECISIONS,
    DOMAINS,
    FLOW_RULES,
    INSTRUCTIONS,
    INTERFACES,
    JOURNEYS,
    LANDSCAPE,
    LIFECYCLE_BLOCKS,
    LIFECYCLE_NOTES,
    MASTERY_NOTES,
    NFRS,
    OFFERING_COMPONENTS,
    ORDER_TYPES,
    PLAN_CHARACTERISTICS,
    PLANS,
    PORTFOLIO,
    PRODUCT_POINTS,
    PRODUCTS,
    QUESTIONS,
    REALISATION,
    REALISATION_RECORDS,
    RELATIONSHIPS,
    RESPONSIBILITIES,
    SCHEMES,
    SOURCES,
    SYSTEMS,
    TRACKING,
    TRACKING_CHANNELS,
    TRACKING_EVENTS,
    TRACKING_FLOWS,
    VOCABULARY,
)


def _append(sheet: Any, values: Sequence[object], *, header: bool = False) -> None:
    sheet.append(list(values))
    if header:
        for cell in sheet[sheet.max_row]:
            cell.font = Font(bold=True)


def _yes(value: bool | None) -> str | None:
    return None if value is None else "yes" if value else "no"


def _confidence(item: Any) -> str | None:
    return item.confidence.value if item.confidence else None


def _write_offering(sheets: dict[str, Any], product: ProductOffering) -> None:
    _append(
        sheets[PRODUCTS],
        (
            product.id,
            product.name,
            product.code,
            product.family,
            product.version,
            product.lifecycle,
            product.proposition,
            f"{_LIST_SEPARATOR} ".join(product.rules),
            _confidence(product),
            product.source,
            f"{_LIST_SEPARATOR} ".join(product.sources),
            product.primary_source,
            product.portfolio_node_id,
            product.purpose,
            f"{_LIST_SEPARATOR} ".join(product.former_ids),
            _yes(product.deprecated) if product.deprecated else None,
            product.replaced_by,
            product.family_id,
            product.lifecycle_id,
        ),
    )
    _write_plans_and_rules(sheets, product)
    _write_open_points(sheets, product)
    for order in product.order_types:
        _append(
            sheets[ORDER_TYPES],
            (
                product.id,
                order.code,
                order.name,
                _yes(order.enabled),
                order.description,
                _confidence(order),
                order.source,
                f"{_LIST_SEPARATOR} ".join(order.channels),
                order.delivery_drop,
                order.stage.value if order.stage else None,
            ),
        )
    for part in product.components:
        _write_offering_component(sheets, product.id, part)
    if product.tracking is not None:
        _write_tracking(sheets, product.id, product.tracking)
    for note in product.lifecycle_notes:
        _write_note(sheets, product.id, note)
    for nfr in product.nfrs:
        _append(
            sheets[NFRS],
            (
                product.id,
                nfr.quality,
                nfr.coverage.value,
                nfr.statement,
                _confidence(nfr),
                nfr.source,
                nfr.quality_id,
            ),
        )
    for kind, points in (("value", product.values), ("audience", product.audiences)):
        for point in points:
            _append(
                sheets[PRODUCT_POINTS],
                (product.id, kind, point.name, point.description, _confidence(point), point.source),
            )


def _write_plans_and_rules(sheets: dict[str, Any], product: ProductOffering) -> None:
    for plan in product.plans:
        _append(
            sheets[PLANS],
            (product.id, plan.name, plan.description, _confidence(plan), plan.source),
        )
        for fact in plan.characteristics:
            _append(sheets[PLAN_CHARACTERISTICS], (product.id, plan.name, fact.name, fact.value))
    for rule in product.business_rules:
        _append(
            sheets[BUSINESS_RULES],
            (
                product.id,
                rule.id,
                rule.statement,
                rule.kind,
                _confidence(rule),
                rule.source,
                rule.kind_id,
            ),
        )


def _write_open_points(sheets: dict[str, Any], product: ProductOffering) -> None:
    """An offering's questions and decisions, and what it is bounded by or does not use."""
    for question in product.questions:
        _append(
            sheets[QUESTIONS],
            (
                product.id,
                question.id,
                question.text,
                question.impact,
                _confidence(question),
                question.source,
                f"{_LIST_SEPARATOR} ".join(question.order_types),
                question.title,
            ),
        )
    for decision in product.decisions:
        _append(
            sheets[DECISIONS],
            (
                product.id,
                decision.id,
                decision.title,
                decision.text,
                _confidence(decision),
                decision.source,
            ),
        )
    for kind, texts in (("boundary", product.boundaries), ("not_used", product.not_used)):
        for text in texts:
            _append(sheets[BOUNDARIES], (product.id, kind, text))


def _write_offering_component(
    sheets: dict[str, Any], product_id: str, part: OfferingComponent
) -> None:
    _append(
        sheets[OFFERING_COMPONENTS],
        (
            product_id,
            part.id,
            part.name,
            part.code,
            part.kind,
            _yes(part.mandatory),
            _yes(part.customer_visible),
            part.description,
            part.commercial_spec,
            part.technical_spec,
            part.technical_details,
            _confidence(part),
            part.source,
            f"{_LIST_SEPARATOR} ".join(part.capability_ids),
            part.unlinked_reason,
            part.kind_id,
            _yes(part.hub) if part.hub else None,
        ),
    )
    for duty in part.responsibilities:
        _append(
            sheets[RESPONSIBILITIES],
            (
                product_id,
                part.id,
                duty.system_id,
                duty.role,
                duty.description,
                f"{_LIST_SEPARATOR} ".join(duty.order_types),
                _confidence(duty),
                duty.source,
                duty.role_id,
            ),
        )
    for item in part.realisation:
        _append(
            sheets[REALISATION],
            (
                product_id,
                part.id,
                item.layer.value,
                item.name,
                _confidence(item),
                item.source,
                item.record_id,
            ),
        )


def _write_note(sheets: dict[str, Any], product_id: str, note: LifecycleNote) -> None:
    joined = f"{_LIST_SEPARATOR} ".join
    _append(
        sheets[LIFECYCLE_NOTES],
        (
            product_id,
            note.id,
            note.title,
            note.kind,
            note.summary,
            joined(note.order_types),
            joined(note.channels),
            _confidence(note),
            note.source,
        ),
    )
    blank = (None,) * len(_CELLS)

    def row(kind: str, title: str | None, text: str | None, cells: tuple[object, ...]) -> None:
        padded = (*cells, *blank)[: len(_CELLS)]
        _append(
            sheets[LIFECYCLE_BLOCKS],
            (product_id, note.id, kind, title, text, *padded, None, None, None),
        )

    for block in note.blocks:
        if block.kind is NoteBlockKind.TEXT:
            cells: tuple[object, ...] = ()
            text = block.text
        elif block.kind is NoteBlockKind.LIST:
            cells, text = (), None
        else:
            cells, text = block.columns, block.caption
        _append(
            sheets[LIFECYCLE_BLOCKS],
            (
                product_id,
                note.id,
                block.kind.value,
                block.title,
                text,
                *(*cells, *blank)[: len(_CELLS)],
                _confidence(block),
                block.source,
                "yes" if block.to_verify else None,
            ),
        )
        for item in block.items:
            row("item", None, item, ())
        for table_row in block.rows:
            row("row", None, None, table_row)


def _write_tracking(sheets: dict[str, Any], product_id: str, tracking: OrderTracking) -> None:
    _append(
        sheets[TRACKING],
        (
            product_id,
            f"{_LIST_SEPARATOR} ".join(tracking.order_types),
            tracking.scope_note,
            tracking.not_applicable_note,
            _confidence(tracking),
            tracking.source,
        ),
    )
    for flow in tracking.flows:
        _append(
            sheets[TRACKING_FLOWS],
            (
                product_id,
                flow.from_system_id,
                flow.to_system_id,
                flow.label,
                flow.interface,
                _confidence(flow),
                flow.source,
                flow.interface_id,
            ),
        )
    for channel in tracking.channels:
        _append(
            sheets[TRACKING_CHANNELS],
            (
                product_id,
                channel.channel_id,
                channel.correlation_key,
                channel.ui_system_id,
                channel.story,
                channel.read_system_id,
                channel.read_interface,
                channel.ui_note,
                _confidence(channel),
                channel.source,
            ),
        )
    for kind, events in (("milestone", tracking.milestones), ("status", tracking.statuses)):
        for event in events:
            _append(
                sheets[TRACKING_EVENTS],
                (
                    product_id,
                    kind,
                    event.label,
                    event.detail,
                    event.system_id,
                    _confidence(event),
                    event.source,
                ),
            )
    for case in tracking.fallout:
        _append(
            sheets[TRACKING_EVENTS],
            (
                product_id,
                "fallout",
                case.trigger,
                case.handling,
                case.system_id,
                _confidence(case),
                case.source,
            ),
        )


def _write_journey(sheets: dict[str, Any], journey: Journey) -> None:
    joined = f"{_LIST_SEPARATOR} ".join
    _append(
        sheets[JOURNEYS],
        (
            journey.id,
            journey.name,
            journey.product_id,
            journey.order_type_code,
            journey.description,
            _confidence(journey),
            journey.source,
            f"{_LIST_SEPARATOR} ".join(journey.former_ids),
        ),
    )
    _write_activities(sheets, journey)
    for rule in journey.flow_rules:
        _append(
            sheets[FLOW_RULES],
            (
                journey.id,
                rule.kind.value,
                rule.from_activity,
                rule.to_activity,
                rule.condition,
                rule.branch,
                rule.parallel_group,
                rule.rejoin_at,
                _confidence(rule),
                rule.source,
                rule.channel_id,
            ),
        )
    for link in journey.integrations:
        _append(
            sheets[ACTIVITY_INTEGRATIONS],
            (
                journey.id,
                link.from_activity,
                link.to_activity,
                link.interaction,
                link.interface,
                link.payload,
                link.timing.value if link.timing else None,
                link.correlation_key,
                _confidence(link),
                link.source,
                link.from_system_id,
                link.to_system_id,
                link.via_system_id,
                link.purpose,
                link.style,
                link.tmf_equivalent,
                joined(link.open_api_ids),
                link.interface_id,
            ),
        )


def _write_activities(sheets: dict[str, Any], journey: Journey) -> None:
    joined = f"{_LIST_SEPARATOR} ".join
    for step in journey.activities:
        _append(
            sheets[ACTIVITIES],
            (
                journey.id,
                step.number,
                step.name,
                step.phase,
                step.track,
                step.performing_system_id,
                joined(step.supporting_system_ids),
                step.system_function,
                step.mode,
                _yes(step.customer_visible),
                step.description,
                joined(step.component_ids),
                step.input,
                step.output,
                step.etom,
                _confidence(step),
                step.source,
                joined(step.channels),
                _yes(step.channel_entry) if step.channel_entry else None,
                step.performer,
                step.point_of_no_return,
                step.role,
                step.etom_id,
                step.role_id,
                step.mode_id,
                step.phase_id,
            ),
        )


def _workbook(release: ArchitectureKnowledge | None) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.worksheets[0])
    instructions = workbook.create_sheet(INSTRUCTIONS)
    for line in _INSTRUCTIONS:
        instructions.append(list(line))
    instructions["A1"].font = Font(bold=True)
    instructions.column_dimensions["A"].width = 90
    sheets = {name: workbook.create_sheet(name) for name in _HEADERS}
    for name, sheet in sheets.items():
        _append(sheet, _HEADERS[name], header=True)
        sheet.freeze_panes = "A2"
        for column in "ABCDEFGHIJKLMNOPQRS":
            sheet.column_dimensions[column].width = 32
    if release is not None:
        _write_release(sheets, release)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _write_release(sheets: dict[str, Any], release: ArchitectureKnowledge) -> None:
    """Every row of the release, each sheet in the order the catalogue lists its entries."""
    for system in release.systems:
        _write_system(sheets, system)
    _write_relationships(sheets, release.relationships)
    _write_domains(sheets[DOMAINS], release.capability_domains)
    _write_domains(sheets[LANDSCAPE], release.landscape_domains)
    _write_concepts(sheets, release.business_capabilities)
    _write_schemes(sheets, release.schemes)
    _write_terms(sheets, release.vocabulary)
    _write_interfaces(sheets, release.interfaces)
    _write_realisations(sheets, release.realisations)
    _write_portfolio(sheets, release.portfolio)
    for product in release.products:
        _write_offering(sheets, product)
    for journey in release.journeys:
        _write_journey(sheets, journey)
    _write_channels(sheets, release.channels)
    _write_sources(sheets, release.sources)
    for conflict in release.conflicts:
        _write_conflict(sheets, conflict)
    for record in release.change_history:
        _write_change_request(sheets, record)
    _write_acceptances(sheets, release)


def _write_system(sheets: dict[str, Any], system: SystemDefinition) -> None:
    _append(
        sheets[SYSTEMS],
        (
            system.id,
            system.name,
            f"{_LIST_SEPARATOR} ".join(system.aliases),
            system.description,
            system.landscape_domain_id,
            system.owner,
            _yes(system.external) if system.external else None,
            system.roadmap,
            system.placement_from,
            system.placement_reason,
            _confidence(system),
            system.source,
            f"{_LIST_SEPARATOR} ".join(system.masters),
            f"{_LIST_SEPARATOR} ".join(system.reads),
            _yes(system.deprecated) if system.deprecated else None,
            system.replaced_by,
        ),
    )
    for capability in system.capabilities:
        _append(
            sheets[CAPABILITIES],
            (
                system.id,
                capability.id,
                capability.name,
                f"{_LIST_SEPARATOR} ".join(capability.triggers),
                capability.domain_id,
                capability.component_id,
                capability.concept_id,
                capability.unlinked_reason,
                _confidence(capability),
                capability.source,
            ),
        )
    for component in system.components:
        _append(
            sheets[COMPONENTS],
            (
                system.id,
                component.id,
                component.name,
                f"{_LIST_SEPARATOR} ".join(component.aliases),
                component.technology,
                component.description,
                _confidence(component),
                component.source,
                component.technology_id,
            ),
        )
    for constraint in system.constraints:
        _append(sheets[CONSTRAINTS], (system.id, constraint))
    for note in system.mastery_notes:
        _append(sheets[MASTERY_NOTES], (system.id, note.entity_id, note.note))


def _write_relationships(
    sheets: dict[str, Any], relationships: Sequence[SystemRelationship]
) -> None:
    for relationship in relationships:
        _append(
            sheets[RELATIONSHIPS],
            (
                relationship.source_system_id,
                relationship.target_system_id,
                relationship.description,
                relationship.kind.value,
                _confidence(relationship),
                relationship.source,
            ),
        )


def _write_domains(sheet: Any, tree: Sequence[CapabilityDomain | LandscapeDomain]) -> None:
    for domain in tree:
        crossing = (
            (
                _yes(domain.cross_cutting) if domain.cross_cutting else None,
                domain.layout.value if domain.layout else None,
            )
            if isinstance(domain, LandscapeDomain)
            else ()
        )
        _append(
            sheet,
            (
                domain.id,
                domain.name,
                domain.parent_id,
                domain.description,
                *crossing,
            ),
        )


def _write_concepts(sheets: dict[str, Any], concepts: Sequence[BusinessCapability]) -> None:
    for concept in concepts:
        _append(
            sheets[CONCEPTS],
            (
                concept.id,
                concept.pref_label,
                f"{_LIST_SEPARATOR} ".join(concept.alt_labels),
                concept.definition,
                concept.broader_id,
                concept.domain_id,
                concept.exact_match,
                _confidence(concept),
                concept.source,
                _yes(concept.deprecated) if concept.deprecated else None,
                concept.replaced_by,
                f"{_LIST_SEPARATOR} ".join(concept.close_match),
            ),
        )


def _write_schemes(sheets: dict[str, Any], schemes: Sequence[ConceptScheme]) -> None:
    for scheme in schemes:
        _append(
            sheets[SCHEMES],
            (
                scheme.id,
                scheme.name,
                scheme.definition,
                f"{_LIST_SEPARATOR} ".join(item.value for item in scheme.fields),
            ),
        )


def _write_terms(sheets: dict[str, Any], terms: Sequence[VocabularyTerm]) -> None:
    for term in terms:
        _append(
            sheets[VOCABULARY],
            (
                term.id,
                str(term.scheme),
                term.pref_label,
                f"{_LIST_SEPARATOR} ".join(term.alt_labels),
                term.notation,
                term.definition,
                term.broader_id,
                term.exact_match,
                _confidence(term),
                term.source,
                _yes(term.deprecated) if term.deprecated else None,
                term.replaced_by,
                f"{_LIST_SEPARATOR} ".join(term.close_match),
                term.rank,
            ),
        )


def _write_interfaces(sheets: dict[str, Any], interfaces: Sequence[SystemInterface]) -> None:
    for interface in interfaces:
        _append(
            sheets[INTERFACES],
            (
                interface.id,
                interface.name,
                interface.system_id,
                interface.style.value,
                f"{_LIST_SEPARATOR} ".join(interface.consumer_ids),
                f"{_LIST_SEPARATOR} ".join(interface.open_api_ids),
                f"{_LIST_SEPARATOR} ".join(interface.entity_ids),
                interface.description,
                _confidence(interface),
                interface.source,
                f"{_LIST_SEPARATOR} ".join(interface.relays),
            ),
        )


def _write_realisations(sheets: dict[str, Any], realisations: Sequence[RealisationRecord]) -> None:
    for realised in realisations:
        _append(
            sheets[REALISATION_RECORDS],
            (
                realised.id,
                realised.layer.value,
                realised.name,
                f"{_LIST_SEPARATOR} ".join(realised.aliases),
                f"{_LIST_SEPARATOR} ".join(realised.system_ids),
                f"{_LIST_SEPARATOR} ".join(realised.realised_by),
                realised.description,
                _confidence(realised),
                realised.source,
                f"{_LIST_SEPARATOR} ".join(realised.rules),
            ),
        )


def _write_portfolio(sheets: dict[str, Any], portfolio: Sequence[PortfolioNode]) -> None:
    for node in portfolio:
        _append(
            sheets[PORTFOLIO],
            (
                node.id,
                node.name,
                node.level,
                node.parent_id,
                node.description,
                _confidence(node),
                node.source,
                _yes(node.deprecated) if node.deprecated else None,
                node.replaced_by,
                node.level_id,
            ),
        )


def _write_channels(sheets: dict[str, Any], channels: Sequence[Channel]) -> None:
    for channel in channels:
        _append(
            sheets[CHANNELS],
            (
                channel.id,
                channel.name,
                channel.kind,
                channel.entry_system_id,
                channel.description,
                _confidence(channel),
                channel.source,
                channel.kind_id,
            ),
        )


def _write_sources(sheets: dict[str, Any], sources: Sequence[KnowledgeSource]) -> None:
    for source in sources:
        _append(
            sheets[SOURCES],
            (
                source.id,
                source.level.value,
                source.title,
                source.short,
                source.version,
                source.file,
                _yes(source.supplied),
                source.authority,
                source.scope,
                source.boundary,
            ),
        )


def _write_conflict(sheets: dict[str, Any], conflict: SourceConflict) -> None:
    _append(
        sheets[CONFLICTS],
        (
            conflict.id,
            conflict.title,
            conflict.a.source_id,
            conflict.a.reference,
            conflict.a.statement,
            conflict.b.source_id,
            conflict.b.reference,
            conflict.b.statement,
            conflict.difference,
            conflict.impact,
            conflict.decision,
            _confidence(conflict),
            conflict.source,
        ),
    )
    for scope in conflict.scope:
        _append(
            sheets[CONFLICT_SCOPES],
            (
                conflict.id,
                scope.product_id,
                f"{_LIST_SEPARATOR} ".join(scope.order_types),
                scope.question_id,
            ),
        )


def _write_change_request(sheets: dict[str, Any], record: ChangeRequestRecord) -> None:
    trace = record.trace
    _append(
        sheets[CHANGE_HISTORY],
        (
            record.id,
            record.title,
            record.origin.value,
            record.product_id,
            record.requester,
            record.reason,
            record.priority,
            record.target_date,
            _iso(record.created_at),
            _iso(record.applied_at),
            f"{_LIST_SEPARATOR} ".join(record.gaps),
            trace.requirement_id if trace else None,
            trace.breakdown_revision if trace else None,
            trace.approval_id if trace else None,
            trace.approved_by if trace else None,
            _iso(trace.approved_at) if trace else None,
            trace.epic_id if trace else None,
            trace.epic_name if trace else None,
            f"{_LIST_SEPARATOR} ".join(f"{item.id}: {item.name}" for item in trace.features)
            if trace
            else None,
            trace.export_schema if trace else None,
            trace.knowledge_version if trace else None,
        ),
    )
    for item in record.items:
        _append(
            sheets[CHANGE_ITEMS],
            (record.id, item.kind, item.summary, item.status.value, item.feature_id),
        )


def _write_acceptances(sheets: dict[str, Any], release: ArchitectureKnowledge) -> None:
    for row in _acceptance_rows(release):
        _append(
            sheets[ACCEPTANCES],
            (
                row["entity"],
                json.dumps(row["key"], ensure_ascii=False),
                row["candidate_id"],
                row["model"],
                row["prompt_version"],
                row["decided_by"],
                row["decided_at"],
                json.dumps(row["citations"], ensure_ascii=False) if row.get("citations") else None,
            ),
        )
