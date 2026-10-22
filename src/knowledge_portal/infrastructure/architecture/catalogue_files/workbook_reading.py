"""The Excel workbook read into catalogue content: one sheet per kind, each row gathered
into the YAML/JSON shape where the mapping readers can read it."""

from __future__ import annotations

import io
import json
import zipfile
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from smb_kernel.documents.extraction_base import MAX_OFFICE_PARTS, MAX_OFFICE_UNCOMPRESSED_BYTES

from knowledge_portal.application.ports.catalogue_file import CatalogueContent
from knowledge_portal.domain.architecture.acceptance import with_recorded_acceptances
from knowledge_portal.domain.architecture.knowledge import (
    CapabilityDomain,
    InvalidKnowledgeError,
    KnowledgeCapability,
    LandscapeDomain,
    MasteryNote,
    SystemComponent,
    SystemDefinition,
    SystemRelationship,
)
from knowledge_portal.infrastructure.architecture.catalogue_files.fields import (
    _flag,
    _kind,
    _located,
    _optional_text,
    _retirement,
    _split,
    _text,
    _trust,
)
from knowledge_portal.infrastructure.architecture.catalogue_files.mapping_reading import (
    _acceptances,
    _channels,
    _concepts,
    _conflicts,
    _content,
    _cross_cutting,
    _history,
    _interfaces,
    _journeys,
    _portfolio,
    _realisations,
    _schemes,
    _sources,
    _terms,
)
from knowledge_portal.infrastructure.architecture.catalogue_files.offering_reading import _offerings
from knowledge_portal.infrastructure.architecture.catalogue_files.offering_sheets import (
    _sheet_offerings,
)
from knowledge_portal.infrastructure.architecture.catalogue_files.workbook_sheets import (
    ACCEPTANCES,
    ACTIVITIES,
    ACTIVITY_INTEGRATIONS,
    CAPABILITIES,
    CHANGE_HISTORY,
    CHANGE_ITEMS,
    CHANNELS,
    COMPONENTS,
    CONCEPTS,
    CONFLICT_SCOPES,
    CONFLICTS,
    CONSTRAINTS,
    DOMAINS,
    FLOW_RULES,
    INTERFACES,
    JOURNEYS,
    LANDSCAPE,
    MASTERY_NOTES,
    PORTFOLIO,
    REALISATION_RECORDS,
    RELATIONSHIPS,
    SCHEMES,
    SOURCES,
    SYSTEMS,
    VOCABULARY,
    _rows,
)


def _key_cell(value: object, where: str) -> list[object]:
    """A workbook key: the JSON list the export wrote, such as ["bcrm", "quotes"]."""
    try:
        key = json.loads(_text(value, where, "key"))
    except json.JSONDecodeError as exc:
        raise InvalidKnowledgeError(f"{where}: key must be a list as exported.") from exc
    if not isinstance(key, list):
        raise InvalidKnowledgeError(f"{where}: key must be a list as exported.")
    return key


def _check_archive(content: bytes) -> None:
    """Refuse archive bombs before openpyxl inflates anything."""
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            parts = archive.infolist()
    except zipfile.BadZipFile as exc:
        raise InvalidKnowledgeError("The Excel file is damaged or not an .xlsx workbook.") from exc
    if len(parts) > MAX_OFFICE_PARTS or (
        sum(item.file_size for item in parts) > MAX_OFFICE_UNCOMPRESSED_BYTES
    ):
        raise InvalidKnowledgeError("The Excel file is too large to import.")


def _read_workbook(content: bytes) -> CatalogueContent:
    _check_archive(content)
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except (InvalidFileException, KeyError, OSError, ValueError) as exc:
        raise InvalidKnowledgeError("The Excel file could not be opened.") from exc
    try:
        systems = _sheet_systems(workbook)
        _sheet_capabilities(workbook, systems)
        _sheet_components(workbook, systems)
        _sheet_constraints(workbook, systems)
        _sheet_mastery_notes(workbook, systems)
        relationships = _sheet_relationships(workbook)
        domains = _sheet_domains(workbook, DOMAINS, CapabilityDomain)
        landscape = _cross_cutting(
            _sheet_domains(workbook, LANDSCAPE, LandscapeDomain),
            _rows(workbook, LANDSCAPE),
            f"{LANDSCAPE} row",
        )
        offerings = _offerings(_sheet_offerings(workbook))
        journeys = _journeys(_sheet_journeys(workbook))
        channels = _channels(_sheet_channels(workbook))
        sources = _sources(_sheet_sources(workbook))
        conflicts = _conflicts(_sheet_conflicts(workbook))
        history = (
            _history(_sheet_history(workbook)) if CHANGE_HISTORY in workbook.sheetnames else None
        )
        portfolio = _portfolio(_sheet_portfolio(workbook))
        concepts = _concepts(_sheet_concepts(workbook))
        terms = _terms(_sheet_terms(workbook))
        schemes = _schemes(_sheet_schemes(workbook))
        interfaces = _interfaces(_sheet_interfaces(workbook))
        realisations = _realisations(_sheet_realisations(workbook))
        acceptances = _acceptances(_sheet_acceptances(workbook))
    finally:
        workbook.close()
    body = _content(
        _system_definitions(systems),
        relationships,
        domains,
        landscape,
        offerings,
        journeys,
        channels,
        sources,
        conflicts,
        history,
        portfolio,
        concepts,
        terms,
        interfaces,
        realisations,
        schemes=schemes,
    )
    return with_recorded_acceptances(body, acceptances)


# The Systems sheet first, then the sheets whose rows each belong to one of its systems.
def _sheet_systems(workbook: Any) -> dict[str, dict[str, Any]]:
    systems: dict[str, dict[str, Any]] = {}
    for number, cells in _rows(workbook, SYSTEMS):
        where = f"{SYSTEMS} row {number}"
        system_id = _text(cells.get("system_id"), where, "system_id")
        if system_id in systems:
            raise InvalidKnowledgeError(f"{where}: system {system_id!r} is listed twice.")
        systems[system_id] = {
            "where": where,
            "name": _text(cells.get("name"), where, "name"),
            "aliases": _split(cells.get("aliases"), where, "aliases"),
            "description": _optional_text(cells.get("description"), where, "description"),
            "landscape_domain_id": _optional_text(
                cells.get("landscape_domain_id"), where, "landscape_domain_id"
            ),
            "owner": _optional_text(cells.get("owner"), where, "owner"),
            "external": bool(_flag(cells.get("external"), where, "external", False)),
            "roadmap": _optional_text(cells.get("roadmap"), where, "roadmap"),
            "placement_from": _optional_text(cells.get("placement_from"), where, "placement_from"),
            "placement_reason": _optional_text(
                cells.get("placement_reason"), where, "placement_reason"
            ),
            "confidence": _trust(cells.get("confidence"), where),
            "source": _optional_text(cells.get("source"), where, "source"),
            "masters": _split(cells.get("masters"), where, "masters"),
            "reads": _split(cells.get("reads"), where, "reads"),
            **_retirement(cells, where),
            "mastery_notes": [],
            "capabilities": [],
            "constraints": [],
            "components": [],
        }
    return systems


def _owning_system(
    systems: dict[str, dict[str, Any]], cells: dict[str, object], where: str
) -> dict[str, Any]:
    system_id = _text(cells.get("system_id"), where, "system_id")
    if system_id not in systems:
        raise InvalidKnowledgeError(
            f"{where}: system {system_id!r} is not listed on the {SYSTEMS} sheet."
        )
    return systems[system_id]


def _sheet_capabilities(workbook: Any, systems: dict[str, dict[str, Any]]) -> None:
    for number, cells in _rows(workbook, CAPABILITIES):
        where = f"{CAPABILITIES} row {number}"
        try:
            _owning_system(systems, cells, where)["capabilities"].append(
                KnowledgeCapability(
                    _text(cells.get("capability_id"), where, "capability_id"),
                    _text(cells.get("name"), where, "name"),
                    _split(cells.get("triggers"), where, "triggers"),
                    _optional_text(cells.get("domain_id"), where, "domain_id"),
                    _optional_text(cells.get("component_id"), where, "component_id"),
                    _optional_text(cells.get("concept_id"), where, "concept_id"),
                    _optional_text(cells.get("unlinked_reason"), where, "unlinked_reason"),
                    confidence=_trust(cells.get("confidence"), where),
                    source=_optional_text(cells.get("source"), where, "source"),
                )
            )
        except InvalidKnowledgeError as exc:
            raise _located(where, exc) from exc


def _sheet_components(workbook: Any, systems: dict[str, dict[str, Any]]) -> None:
    for number, cells in _rows(workbook, COMPONENTS):
        where = f"{COMPONENTS} row {number}"
        try:
            _owning_system(systems, cells, where)["components"].append(
                SystemComponent(
                    id=_text(cells.get("component_id"), where, "component_id"),
                    name=_text(cells.get("name"), where, "name"),
                    description=_optional_text(cells.get("description"), where, "description"),
                    aliases=_split(cells.get("aliases"), where, "aliases"),
                    technology=_optional_text(cells.get("technology"), where, "technology"),
                    confidence=_trust(cells.get("confidence"), where),
                    source=_optional_text(cells.get("source"), where, "source"),
                    technology_id=_optional_text(
                        cells.get("technology_term"), where, "technology_term"
                    ),
                )
            )
        except InvalidKnowledgeError as exc:
            raise _located(where, exc) from exc


def _sheet_constraints(workbook: Any, systems: dict[str, dict[str, Any]]) -> None:
    for number, cells in _rows(workbook, CONSTRAINTS):
        where = f"{CONSTRAINTS} row {number}"
        _owning_system(systems, cells, where)["constraints"].append(
            _text(cells.get("constraint"), where, "constraint")
        )


def _sheet_mastery_notes(workbook: Any, systems: dict[str, dict[str, Any]]) -> None:
    for number, cells in _rows(workbook, MASTERY_NOTES):
        where = f"{MASTERY_NOTES} row {number}"
        try:
            _owning_system(systems, cells, where)["mastery_notes"].append(
                MasteryNote(
                    _text(cells.get("entity_id"), where, "entity_id"),
                    _text(cells.get("note"), where, "note"),
                )
            )
        except InvalidKnowledgeError as exc:
            raise _located(where, exc) from exc


def _system_definitions(systems: dict[str, dict[str, Any]]) -> list[SystemDefinition]:
    definitions = []
    for system_id, data in systems.items():
        try:
            definitions.append(
                SystemDefinition(
                    id=system_id,
                    name=data["name"],
                    aliases=data["aliases"],
                    capabilities=tuple(data["capabilities"]),
                    constraints=tuple(data["constraints"]),
                    components=tuple(data["components"]),
                    description=data["description"],
                    landscape_domain_id=data["landscape_domain_id"],
                    owner=data["owner"],
                    external=data["external"],
                    roadmap=data["roadmap"],
                    placement_from=data["placement_from"],
                    placement_reason=data["placement_reason"],
                    confidence=data["confidence"],
                    source=data["source"],
                    masters=data["masters"],
                    reads=data["reads"],
                    deprecated=data["deprecated"],
                    replaced_by=data["replaced_by"],
                    mastery_notes=tuple(data["mastery_notes"]),
                )
            )
        except InvalidKnowledgeError as exc:
            raise _located(data["where"], exc) from exc
    return definitions


def _sheet_relationships(workbook: Any) -> list[SystemRelationship]:
    relationships = []
    for number, cells in _rows(workbook, RELATIONSHIPS):
        where = f"{RELATIONSHIPS} row {number}"
        try:
            relationships.append(
                SystemRelationship(
                    _text(cells.get("source_system_id"), where, "source_system_id"),
                    _text(cells.get("target_system_id"), where, "target_system_id"),
                    _text(cells.get("description"), where, "description"),
                    _kind(cells.get("kind"), where),
                    confidence=_trust(cells.get("confidence"), where),
                    source=_optional_text(cells.get("source"), where, "source"),
                )
            )
        except InvalidKnowledgeError as exc:
            raise _located(where, exc) from exc
    return relationships


# The sheets read as the YAML/JSON entries of their kind, each entry knowing its row.
def _sheet_channels(workbook: Any) -> list[dict[str, Any]]:
    return [
        {
            **cells,
            "_where": f"{CHANNELS} row {number}",
            "id": cells.get("channel_id"),
            "entry_system": cells.get("entry_system_id"),
        }
        for number, cells in _rows(workbook, CHANNELS)
    ]


def _sheet_sources(workbook: Any) -> list[dict[str, Any]]:
    return [
        {**cells, "_where": f"{SOURCES} row {number}", "id": cells.get("source_id")}
        for number, cells in _rows(workbook, SOURCES)
    ]


def _sheet_portfolio(workbook: Any) -> list[dict[str, Any]]:
    return [
        {
            **cells,
            "_where": f"{PORTFOLIO} row {number}",
            "id": cells.get("node_id"),
            "parent": cells.get("parent_id"),
        }
        for number, cells in _rows(workbook, PORTFOLIO)
    ]


def _sheet_concepts(workbook: Any) -> list[dict[str, Any]]:
    return [
        {
            **cells,
            "_where": f"{CONCEPTS} row {number}",
            "id": cells.get("concept_id"),
            "alt_labels": list(
                _split(cells.get("alt_labels"), f"{CONCEPTS} row {number}", "alt_labels")
            ),
            "close_match": list(
                _split(cells.get("close_match"), f"{CONCEPTS} row {number}", "close_match")
            ),
            "broader": cells.get("broader_id"),
            "domain": cells.get("domain_id"),
        }
        for number, cells in _rows(workbook, CONCEPTS)
    ]


def _sheet_terms(workbook: Any) -> list[dict[str, Any]]:
    return [
        {
            **cells,
            "_where": f"{VOCABULARY} row {number}",
            "id": cells.get("term_id"),
            "alt_labels": list(
                _split(cells.get("alt_labels"), f"{VOCABULARY} row {number}", "alt_labels")
            ),
            "close_match": list(
                _split(cells.get("close_match"), f"{VOCABULARY} row {number}", "close_match")
            ),
            "broader": cells.get("broader_id"),
        }
        for number, cells in _rows(workbook, VOCABULARY)
    ]


def _sheet_schemes(workbook: Any) -> list[dict[str, Any]]:
    return [
        {
            **cells,
            "_where": f"{SCHEMES} row {number}",
            "id": cells.get("scheme_id"),
            "fields": list(_split(cells.get("fields"), f"{SCHEMES} row {number}", "fields")),
        }
        for number, cells in _rows(workbook, SCHEMES)
    ]


def _sheet_interfaces(workbook: Any) -> list[dict[str, Any]]:
    return [
        {
            **cells,
            "_where": f"{INTERFACES} row {number}",
            "id": cells.get("interface_id"),
            "system": cells.get("system_id"),
            **{
                key: list(_split(cells.get(key), f"{INTERFACES} row {number}", key))
                for key in ("consumers", "open_apis", "entities", "relays")
            },
        }
        for number, cells in _rows(workbook, INTERFACES)
    ]


def _sheet_realisations(workbook: Any) -> list[dict[str, Any]]:
    return [
        {
            **cells,
            "_where": f"{REALISATION_RECORDS} row {number}",
            "id": cells.get("record_id"),
            **{
                key: list(_split(cells.get(key), f"{REALISATION_RECORDS} row {number}", key))
                for key in ("aliases", "systems", "realised_by")
            },
        }
        for number, cells in _rows(workbook, REALISATION_RECORDS)
    ]


def _sheet_acceptances(workbook: Any) -> list[dict[str, Any]]:
    return [
        {
            **cells,
            "_where": f"{ACCEPTANCES} row {number}",
            "key": _key_cell(cells.get("key"), f"{ACCEPTANCES} row {number}"),
        }
        for number, cells in _rows(workbook, ACCEPTANCES)
    ]


def _sheet_journeys(workbook: Any) -> list[dict[str, Any]]:
    """The journey sheets gathered into the YAML/JSON shape, each entry knowing its row."""
    journeys: dict[str, dict[str, Any]] = {}
    for number, cells in _rows(workbook, JOURNEYS):
        where = f"{JOURNEYS} row {number}"
        journey_id = _text(cells.get("journey_id"), where, "journey_id")
        if journey_id in journeys:
            raise InvalidKnowledgeError(f"{where}: journey {journey_id!r} is listed twice.")
        journeys[journey_id] = {
            "_where": where,
            "id": journey_id,
            "name": cells.get("name"),
            "product": cells.get("product_id"),
            "order_type": cells.get("order_type_code"),
            "description": cells.get("description"),
            "confidence": cells.get("confidence"),
            "source": cells.get("source"),
            "former_ids": list(_split(cells.get("former_ids"), where, "former_ids")),
            "activities": [],
            "flow_rules": [],
            "integrations": [],
        }

    def journey(cells: dict[str, object], where: str) -> dict[str, Any]:
        journey_id = _text(cells.get("journey_id"), where, "journey_id")
        if journey_id not in journeys:
            raise InvalidKnowledgeError(
                f"{where}: journey {journey_id!r} is not listed on the {JOURNEYS} sheet."
            )
        return journeys[journey_id]

    for number, cells in _rows(workbook, ACTIVITIES):
        where = f"{ACTIVITIES} row {number}"
        journey(cells, where)["activities"].append(
            {
                **cells,
                "_where": where,
                "system": cells.get("performing_system_id"),
                "supporting": list(
                    _split(cells.get("supporting_system_ids"), where, "supporting_system_ids")
                ),
                "components": list(_split(cells.get("component_ids"), where, "component_ids")),
                "channels": list(_split(cells.get("channels"), where, "channels")),
            }
        )
    for number, cells in _rows(workbook, FLOW_RULES):
        where = f"{FLOW_RULES} row {number}"
        journey(cells, where)["flow_rules"].append(
            {
                **cells,
                "_where": where,
                "from": cells.get("from_activity"),
                "to": cells.get("to_activity"),
                "channel": cells.get("channel_id"),
            }
        )
    for number, cells in _rows(workbook, ACTIVITY_INTEGRATIONS):
        where = f"{ACTIVITY_INTEGRATIONS} row {number}"
        journey(cells, where)["integrations"].append(
            {
                **cells,
                "_where": where,
                "from": cells.get("from_activity"),
                "to": cells.get("to_activity"),
                "open_apis": list(_split(cells.get("open_apis"), where, "open_apis")),
            }
        )
    return list(journeys.values())


def _sheet_conflicts(workbook: Any) -> list[dict[str, Any]]:
    """The conflict sheets in the YAML/JSON shape, each entry knowing its row."""
    conflicts: dict[str, dict[str, Any]] = {}
    for number, cells in _rows(workbook, CONFLICTS):
        where = f"{CONFLICTS} row {number}"
        conflict_id = _text(cells.get("conflict_id"), where, "conflict_id")
        if conflict_id in conflicts:
            raise InvalidKnowledgeError(f"{where}: conflict {conflict_id!r} is listed twice.")
        conflicts[conflict_id] = {
            **cells,
            "_where": where,
            "id": conflict_id,
            "a": {
                "source": cells.get("a_source_id"),
                "reference": cells.get("a_reference"),
                "statement": cells.get("a_statement"),
            },
            "b": {
                "source": cells.get("b_source_id"),
                "reference": cells.get("b_reference"),
                "statement": cells.get("b_statement"),
            },
            "scope": [],
        }
    for number, cells in _rows(workbook, CONFLICT_SCOPES):
        where = f"{CONFLICT_SCOPES} row {number}"
        conflict_id = _text(cells.get("conflict_id"), where, "conflict_id")
        if conflict_id not in conflicts:
            raise InvalidKnowledgeError(
                f"{where}: conflict {conflict_id!r} is not listed on the {CONFLICTS} sheet."
            )
        conflicts[conflict_id]["scope"].append(
            {
                "product": cells.get("product_id"),
                "order_types": list(_split(cells.get("order_types"), where, "order_types")),
                "question": cells.get("question_id"),
            }
        )
    return list(conflicts.values())


def _sheet_domains[DomainT: (CapabilityDomain, LandscapeDomain)](
    workbook: Any, sheet: str, kind: type[DomainT]
) -> list[DomainT]:
    domains = []
    for number, cells in _rows(workbook, sheet):
        where = f"{sheet} row {number}"
        try:
            domain = kind(
                id=_text(cells.get("domain_id"), where, "domain_id"),
                name=_text(cells.get("name"), where, "name"),
                parent_id=_optional_text(cells.get("parent_id"), where, "parent_id"),
                description=_optional_text(cells.get("description"), where, "description"),
            )
            domains.append(domain)
        except InvalidKnowledgeError as exc:
            raise _located(where, exc) from exc
    return domains


def _sheet_history(workbook: Any) -> list[dict[str, Any]]:
    """The change history sheets gathered into the YAML/JSON shape."""
    records: dict[str, dict[str, Any]] = {}
    for number, cells in _rows(workbook, CHANGE_HISTORY):
        where = f"{CHANGE_HISTORY} row {number}"
        record_id = _text(cells.get("change_request_id"), where, "change_request_id")
        if record_id in records:
            raise InvalidKnowledgeError(f"{where}: change request {record_id!r} is listed twice.")
        features = [
            dict(zip(("id", "name"), (part.strip() for part in item.split(":", 1)), strict=False))
            for item in _split(cells.get("features"), where, "features")
        ]
        trace = (
            {
                "requirement": cells.get("requirement_id"),
                "revision": cells.get("breakdown_revision"),
                "approval": cells.get("approval_id"),
                "approved_by": cells.get("approved_by"),
                "approved_at": cells.get("approved_at"),
                "epic": cells.get("epic_id"),
                "epic_name": cells.get("epic_name"),
                "features": features,
                "export_schema": cells.get("export_schema"),
                "knowledge_version": cells.get("knowledge_version"),
            }
            if cells.get("requirement_id")
            else None
        )
        records[record_id] = {
            **cells,
            "_where": where,
            "id": record_id,
            "product": cells.get("product_id"),
            "gaps": list(_split(cells.get("gaps"), where, "gaps")),
            "trace": trace,
            "items": [],
        }
    for number, cells in _rows(workbook, CHANGE_ITEMS):
        where = f"{CHANGE_ITEMS} row {number}"
        record_id = _text(cells.get("change_request_id"), where, "change_request_id")
        if record_id not in records:
            raise InvalidKnowledgeError(
                f"{where}: change request {record_id!r} is not on {CHANGE_HISTORY}."
            )
        records[record_id]["items"].append(
            {**cells, "feature": cells.get("feature_id"), "_where": where}
        )
    return list(records.values())
