"""Excel, YAML and JSON catalogue files, all mapping to one catalogue structure.

YAML and JSON share the packaged catalogue's shape (``systems`` and
``dependencies``). Excel spreads the same content over one sheet per kind so
people can fill it in without learning a data format. Every parse error names
where it happened (sheet and row, or entry), because a maintainer has to find
and fix it in the file.
"""

from __future__ import annotations

import io
import json
import zipfile
from collections.abc import Callable, Iterator, Sequence
from datetime import UTC, datetime
from functools import partial
from typing import Any

import yaml
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils.exceptions import InvalidFileException
from smb_kernel.documents.extraction_base import (
    MAX_OFFICE_PARTS,
    MAX_OFFICE_UNCOMPRESSED_BYTES,
    MAX_WORKBOOK_ROWS,
)

from knowledge_portal.application.ports.catalogue_file import (
    CatalogueContent,
    CatalogueFileFormat,
)
from knowledge_portal.domain.architecture.change_requests import (
    ChangeItem,
    ChangeItemStatus,
    ChangeOrigin,
    ChangeRequestRecord,
    RequirementTrace,
    TracedFeature,
)
from knowledge_portal.domain.architecture.channels import Channel
from knowledge_portal.domain.architecture.governance import (
    ArchitectureDecision,
    ConflictScope,
    ConflictSide,
    KnowledgeSource,
    OpenQuestion,
    SourceConflict,
)
from knowledge_portal.domain.architecture.journeys import (
    Activity,
    ActivityIntegration,
    FlowRule,
    Journey,
    flow_rule_kind,
)
from knowledge_portal.domain.architecture.knowledge import (
    ArchitectureKnowledge,
    CapabilityDomain,
    InvalidKnowledgeError,
    KnowledgeCapability,
    LandscapeDomain,
    RelationshipKind,
    SystemComponent,
    SystemDefinition,
    SystemRelationship,
)
from knowledge_portal.domain.architecture.lifecycle import (
    MAX_TABLE_COLUMNS,
    LifecycleNote,
    NoteBlock,
    NoteBlockKind,
)
from knowledge_portal.domain.architecture.portfolio import PortfolioNode
from knowledge_portal.domain.architecture.products import (
    BusinessRule,
    ComponentResponsibility,
    OfferingComponent,
    OfferingNfr,
    OfferingPlan,
    OfferingPoint,
    OrderType,
    PlanCharacteristic,
    ProductOffering,
    Realisation,
    SourceConfidence,
)
from knowledge_portal.domain.architecture.tracking import (
    FalloutCase,
    OrderTracking,
    TrackingChannel,
    TrackingEvent,
    TrackingFlow,
)

SYSTEMS = "Systems"
COMPONENTS = "Components"
CAPABILITIES = "Capabilities"
CONSTRAINTS = "Constraints"
RELATIONSHIPS = "Relationships"
DOMAINS = "Domains"
LANDSCAPE = "LandscapeDomains"
PRODUCTS = "Products"
ORDER_TYPES = "OrderTypes"
OFFERING_COMPONENTS = "OfferingComponents"
RESPONSIBILITIES = "Responsibilities"
PRODUCT_POINTS = "ProductPoints"
JOURNEYS = "Journeys"
ACTIVITIES = "Activities"
FLOW_RULES = "FlowRules"
ACTIVITY_INTEGRATIONS = "ActivityIntegrations"
CHANNELS = "Channels"
REALISATION = "Realisation"
NFRS = "NFRs"
TRACKING = "Tracking"
TRACKING_FLOWS = "TrackingFlows"
TRACKING_CHANNELS = "TrackingChannels"
TRACKING_EVENTS = "TrackingEvents"
LIFECYCLE_NOTES = "LifecycleNotes"
LIFECYCLE_BLOCKS = "LifecycleBlocks"
SOURCES = "Sources"
CONFLICTS = "Conflicts"
CONFLICT_SCOPES = "ConflictScopes"
QUESTIONS = "Questions"
DECISIONS = "Decisions"
BOUNDARIES = "Boundaries"
CHANGE_HISTORY = "ChangeHistory"
CHANGE_ITEMS = "ChangeItems"
PORTFOLIO = "Portfolio"
PLANS = "Plans"
PLAN_CHARACTERISTICS = "PlanCharacteristics"
BUSINESS_RULES = "BusinessRules"
_CELLS = tuple(f"cell_{number}" for number in range(1, MAX_TABLE_COLUMNS + 1))
INSTRUCTIONS = "Instructions"
_DOMAIN_HEADERS = ("domain_id", "name", "name_ar", "parent_id", "description")
_HEADERS: dict[str, tuple[str, ...]] = {
    SYSTEMS: (
        "system_id",
        "name",
        "name_ar",
        "aliases",
        "description",
        "landscape_domain_id",
        "owner",
        "external",
        "roadmap",
        "placement_from",
        "placement_reason",
        "confidence",
        "source",
    ),
    COMPONENTS: (
        "system_id",
        "component_id",
        "name",
        "name_ar",
        "aliases",
        "technology",
        "description",
    ),
    CAPABILITIES: (
        "system_id",
        "capability_id",
        "name",
        "triggers",
        "domain_id",
        "component_id",
    ),
    CONSTRAINTS: ("system_id", "constraint"),
    RELATIONSHIPS: ("source_system_id", "target_system_id", "description", "kind"),
    DOMAINS: _DOMAIN_HEADERS,
    LANDSCAPE: _DOMAIN_HEADERS,
    PRODUCTS: (
        "product_id",
        "name",
        "code",
        "family",
        "version",
        "lifecycle",
        "proposition",
        "rules",
        "confidence",
        "source",
        "sources",
        "primary_source",
        "portfolio",
    ),
    ORDER_TYPES: (
        "product_id",
        "code",
        "name",
        "enabled",
        "description",
        "confidence",
        "source",
        "channels",
    ),
    OFFERING_COMPONENTS: (
        "product_id",
        "component_id",
        "name",
        "code",
        "type",
        "mandatory",
        "customer_visible",
        "description",
        "commercial_spec",
        "technical_spec",
        "technical_details",
        "confidence",
        "source",
    ),
    RESPONSIBILITIES: (
        "product_id",
        "component_id",
        "system_id",
        "role",
        "description",
        "order_types",
        "confidence",
        "source",
    ),
    PRODUCT_POINTS: ("product_id", "kind", "name", "description", "confidence", "source"),
    JOURNEYS: (
        "journey_id",
        "name",
        "product_id",
        "order_type_code",
        "description",
        "confidence",
        "source",
    ),
    ACTIVITIES: (
        "journey_id",
        "number",
        "name",
        "phase",
        "track",
        "performing_system_id",
        "supporting_system_ids",
        "system_function",
        "mode",
        "customer_visible",
        "description",
        "component_ids",
        "input",
        "output",
        "etom",
        "confidence",
        "source",
        "channels",
        "channel_entry",
        "performer",
        "point_of_no_return",
        "role",
    ),
    FLOW_RULES: (
        "journey_id",
        "kind",
        "from_activity",
        "to_activity",
        "condition",
        "branch",
        "parallel_group",
        "rejoin_at",
        "confidence",
        "source",
    ),
    ACTIVITY_INTEGRATIONS: (
        "journey_id",
        "from_activity",
        "to_activity",
        "interaction",
        "interface",
        "payload",
        "timing",
        "correlation_key",
        "confidence",
        "source",
        "from_system",
        "to_system",
        "via",
        "purpose",
        "style",
        "tmf_equivalent",
    ),
    CHANNELS: (
        "channel_id",
        "name",
        "kind",
        "entry_system_id",
        "description",
        "confidence",
        "source",
    ),
    REALISATION: ("product_id", "component_id", "layer", "name", "confidence", "source"),
    NFRS: ("product_id", "quality", "coverage", "statement", "confidence", "source"),
    TRACKING: (
        "product_id",
        "order_types",
        "scope_note",
        "not_applicable_note",
        "confidence",
        "source",
    ),
    TRACKING_FLOWS: (
        "product_id",
        "from_system_id",
        "to_system_id",
        "label",
        "interface",
        "confidence",
        "source",
    ),
    TRACKING_CHANNELS: (
        "product_id",
        "channel_id",
        "correlation_key",
        "ui_system_id",
        "story",
        "read_system_id",
        "read_interface",
        "ui_note",
        "confidence",
        "source",
    ),
    TRACKING_EVENTS: (
        "product_id",
        "kind",
        "label",
        "detail",
        "system_id",
        "confidence",
        "source",
    ),
    LIFECYCLE_NOTES: (
        "product_id",
        "note_id",
        "title",
        "kind",
        "summary",
        "order_types",
        "channels",
        "confidence",
        "source",
    ),
    LIFECYCLE_BLOCKS: (
        "product_id",
        "note_id",
        "kind",
        "title",
        "text",
        *_CELLS,
        "confidence",
        "source",
        "to_verify",
    ),
    SOURCES: (
        "source_id",
        "level",
        "title",
        "short",
        "version",
        "file",
        "supplied",
        "authority",
        "scope",
        "boundary",
    ),
    CONFLICTS: (
        "conflict_id",
        "title",
        "a_source_id",
        "a_reference",
        "a_statement",
        "b_source_id",
        "b_reference",
        "b_statement",
        "difference",
        "impact",
        "decision",
        "confidence",
        "source",
    ),
    CONFLICT_SCOPES: ("conflict_id", "product_id", "order_types", "question_id"),
    QUESTIONS: (
        "product_id",
        "question_id",
        "text",
        "impact",
        "confidence",
        "source",
        "order_types",
    ),
    DECISIONS: ("product_id", "decision_id", "title", "text", "confidence", "source"),
    BOUNDARIES: ("product_id", "kind", "text"),
    CHANGE_HISTORY: (
        "change_request_id",
        "title",
        "origin",
        "product_id",
        "requester",
        "reason",
        "priority",
        "target_date",
        "created_at",
        "applied_at",
        "gaps",
        "requirement_id",
        "breakdown_revision",
        "approval_id",
        "approved_by",
        "approved_at",
        "epic_id",
        "epic_name",
        "features",
        "export_schema",
        "knowledge_version",
    ),
    CHANGE_ITEMS: ("change_request_id", "kind", "summary", "status", "feature_id"),
    PORTFOLIO: ("node_id", "name", "level", "parent_id", "description", "confidence", "source"),
    PLANS: ("product_id", "plan", "description", "confidence", "source"),
    PLAN_CHARACTERISTICS: ("product_id", "plan", "characteristic", "value"),
    BUSINESS_RULES: ("product_id", "rule_id", "statement", "kind", "confidence", "source"),
}
# Headers a sheet cannot do without. Columns added later stay optional, so
# workbooks filled from an older template still import.
_REQUIRED_HEADERS: dict[str, tuple[str, ...]] = {
    SYSTEMS: ("system_id", "name"),
    COMPONENTS: ("system_id", "component_id", "name"),
    CAPABILITIES: ("system_id", "capability_id", "name", "triggers"),
    CONSTRAINTS: _HEADERS[CONSTRAINTS],
    RELATIONSHIPS: ("source_system_id", "target_system_id", "description"),
    DOMAINS: ("domain_id", "name"),
    LANDSCAPE: ("domain_id", "name"),
    PRODUCTS: ("product_id", "name"),
    ORDER_TYPES: ("product_id", "code", "name"),
    OFFERING_COMPONENTS: ("product_id", "component_id", "name"),
    RESPONSIBILITIES: ("product_id", "component_id", "system_id", "role", "description"),
    PRODUCT_POINTS: ("product_id", "kind", "name"),
    JOURNEYS: ("journey_id", "name"),
    ACTIVITIES: ("journey_id", "number", "name"),
    FLOW_RULES: ("journey_id", "kind", "from_activity", "to_activity"),
    ACTIVITY_INTEGRATIONS: ("journey_id", "from_activity", "to_activity"),
    CHANNELS: ("channel_id", "name"),
    REALISATION: ("product_id", "component_id", "layer", "name"),
    NFRS: ("product_id", "quality", "coverage"),
    TRACKING: ("product_id",),
    TRACKING_FLOWS: ("product_id", "from_system_id", "to_system_id", "label"),
    TRACKING_CHANNELS: ("product_id", "channel_id"),
    TRACKING_EVENTS: ("product_id", "kind", "label"),
    LIFECYCLE_NOTES: ("product_id", "note_id", "title"),
    LIFECYCLE_BLOCKS: ("product_id", "note_id", "kind"),
    SOURCES: ("source_id", "level", "title"),
    CONFLICTS: (
        "conflict_id",
        "title",
        "a_source_id",
        "a_statement",
        "b_source_id",
        "b_statement",
    ),
    CONFLICT_SCOPES: ("conflict_id", "product_id"),
    QUESTIONS: ("product_id", "question_id", "text"),
    DECISIONS: ("product_id", "decision_id", "title"),
    BOUNDARIES: ("product_id", "kind", "text"),
    CHANGE_HISTORY: ("change_request_id", "title"),
    CHANGE_ITEMS: ("change_request_id", "kind", "summary"),
    PORTFOLIO: ("node_id", "name", "level"),
    PLANS: ("product_id", "plan"),
    PLAN_CHARACTERISTICS: ("product_id", "plan", "characteristic", "value"),
    BUSINESS_RULES: ("product_id", "rule_id", "statement"),
}
_KINDS = ", ".join(kind.value for kind in RelationshipKind)
_REQUIRED_SHEETS = (SYSTEMS,)
_LIST_SEPARATOR = ";"
_INSTRUCTIONS = (
    ("How to fill this catalogue",),
    ("Systems: one row per system. system_id is a short stable key such as bcrm.",),
    ("aliases and triggers: several values separated by semicolons (;).",),
    ("Capabilities: one row per capability; triggers are the phrases that point to it.",),
    ("Constraints: one row per constraint of a system.",),
    ("Relationships: one row per dependency between two systems listed on Systems.",),
    (f"kind (optional): how the source depends on the target, one of {_KINDS}.",),
    ("Domains (optional): business areas such as Order capture; parent_id nests one in another.",),
    ("Capabilities domain_id (optional): the domain a capability belongs to.",),
    ("Components (optional): the parts of a system, such as a module or service.",),
    ("Capabilities component_id (optional): the component of its system that delivers it.",),
    (
        "LandscapeDomains (optional): where systems sit, such as Customer; parent_id makes a "
        "sub-domain such as Customer > Assisted.",
    ),
    ("Systems description and landscape_domain_id (optional): what it is for, and where.",),
    ("Products (optional): product offerings such as Business Pro Plus; rules separated by ;.",),
    ("OrderTypes: how each offering is ordered; enabled is yes or no.",),
    ("OfferingComponents: the parts of each offering; mandatory and customer_visible yes or no.",),
    (
        "Responsibilities: which system does what for a component, in which role; order_types "
        "names order type codes separated by ;.",
    ),
    ("ProductPoints: kind value (what it gives customers) or audience (who it is for).",),
    ("confidence (optional): confirmed, inferred or gap, as the source says.",),
    ("Journeys (optional): the activities that fulfil an offering's order type.",),
    (
        "Activities: one row per numbered activity; track is MAIN or a side track; systems and "
        "component_ids separated by ;.",
    ),
    ("FlowRules: kind decision, loop or parallel, from one activity number to another.",),
    ("ActivityIntegrations: how one activity hands over to another.",),
    (
        "Channels (optional): where orders are placed, such as B2B Web or a shop; "
        "entry_system_id is the system that takes the order in.",
    ),
    ("OrderTypes channels (optional): the channel ids it can be ordered through, by ;.",),
    (
        "Activities channels (optional): the channel ids the step happens in, by ; (none: every "
        "channel); channel_entry yes when the order's channel entry system performs it.",
    ),
    (
        "Realisation (optional): how a component is realised; layer is CFS (what the customer "
        "is sold), RFS (what delivers it) or resource (what it runs on).",
    ),
    (
        "NFRs (optional): an offering's non-functional requirements, one row per quality such "
        "as Availability; coverage is defined, partial or missing.",
    ),
    (
        "Tracking (optional): one row per product whose order tracking its sources describe; "
        "order_types are the codes it is specified for, by ; (none: every order type).",
    ),
    (
        "TrackingFlows: order and milestone events from one system to another (the same "
        "system twice: it logs them); interface is the API or message used.",
    ),
    (
        "TrackingChannels: per channel, the correlation_key tying its order to the fulfilment "
        "order, the ui_system_id the customer follows progress in, and what it reads from.",
    ),
    (
        "TrackingEvents: kind milestone (seen by the customer), status (internal) or fallout "
        "(label is what makes the order fall out, detail how it is handled).",
    ),
    (
        "LifecycleNotes (optional): what happens to a product over its life, one row per "
        "topic; order_types and channels by ; (none: every one).",
    ),
    (
        "LifecycleBlocks: each note's content in order. kind text (text), list then one item "
        "row per item (text), table (column heads in cell_1…, caption in text) then one row "
        "row per table row (cells in cell_1…). to_verify yes: carried over from another "
        "source, to re-verify.",
    ),
    (
        "Sources (optional): the sources the catalogue is read from; level is L1 (the "
        "canonical landscape), L2 (a primary source for its scope) or L3 (a baseline carried "
        "forward); supplied no when it was not supplied and is carried forward unread.",
    ),
    (
        "Conflicts: two sources contradicting each other, a and b each a source_id, where it "
        "says so (reference) and what it says (statement); never reconciled by the catalogue.",
    ),
    (
        "ConflictScopes: the products a conflict affects, its order types by ; (none: every "
        "one), and the question_id it raises for that product.",
    ),
    (
        "Products sources and primary_source (optional): the source_ids a product is read "
        "from, by ;. Questions, Decisions: its open questions and architecture decisions. "
        "Boundaries: kind boundary (what its sources cover) or not_used (what it no longer uses).",
    ),
    (
        "ChangeHistory (optional): the change requests applied to this version, one row each, "
        "with the approved requirement they come from; features as id: name by ;, gaps by ;. "
        "ChangeItems: what each asked for. Leave both out to keep the version's own history.",
    ),
    (
        "Portfolio (optional): where offerings sit, such as Enterprise > Fixed > SMB; level "
        "is the operator's own word for it, parent_id nests one node under another. Products "
        "portfolio: the node_id an offering sits in.",
    ),
    (
        "Plans (optional): an offering's plans as its sources describe them; "
        "PlanCharacteristics: one row per stated fact of a plan, such as Download 200 Mbps. "
        "Prices are read live from the product catalog, never kept here.",
    ),
    ("BusinessRules (optional): the rules an offering is sold and fulfilled by, with a kind.",),
    (
        "Systems owner, external (yes or no), roadmap, and placement_from with "
        "placement_reason when a system sits elsewhere than a source put it.",
    ),
    (
        "Activities performer (a team or the customer, when no system performs it), "
        "point_of_no_return and role. ActivityIntegrations from_system, to_system, via, "
        "purpose, style and tmf_equivalent; a call to a system with no step of its own has "
        "to_activity equal to from_activity.",
    ),
    ("Row 1 of each sheet holds the headers; keep them as they are.",),
)


class CatalogueFileAdapter:
    def read(self, file_format: CatalogueFileFormat, content: bytes) -> CatalogueContent:
        if file_format is CatalogueFileFormat.XLSX:
            return _read_workbook(content)
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise InvalidKnowledgeError("The catalogue file must be UTF-8 text.") from exc
        try:
            raw = (
                json.loads(text)
                if file_format is CatalogueFileFormat.JSON
                else yaml.safe_load(text)
            )
        except (json.JSONDecodeError, yaml.YAMLError) as exc:
            raise InvalidKnowledgeError(
                f"The {file_format.value.upper()} file could not be parsed: {_position(exc)}."
            ) from exc
        return content_from_mapping(raw)

    def write(self, file_format: CatalogueFileFormat, release: ArchitectureKnowledge) -> bytes:
        if file_format is CatalogueFileFormat.XLSX:
            return _workbook(release)
        mapping = release_to_mapping(release)
        if file_format is CatalogueFileFormat.JSON:
            return json.dumps(mapping, ensure_ascii=False, indent=2).encode()
        return yaml.safe_dump(mapping, allow_unicode=True, sort_keys=False).encode()

    def template(self) -> bytes:
        return _workbook(None)


def _position(exc: json.JSONDecodeError | yaml.YAMLError) -> str:
    if isinstance(exc, json.JSONDecodeError):
        return f"line {exc.lineno}, column {exc.colno}"
    mark = getattr(exc, "problem_mark", None)
    return f"line {mark.line + 1}" if mark is not None else "invalid syntax"


def _text(value: object, where: str, field: str) -> str:
    if value is None or (isinstance(value, str) and not value.strip()):
        raise InvalidKnowledgeError(f"{where}: {field} is required.")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if not isinstance(value, str | int | float):
        raise InvalidKnowledgeError(f"{where}: {field} must be text.")
    return str(value).strip()


def _optional_text(value: object, where: str, field: str) -> str | None:
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    return _text(value, where, field)


def _kind(value: object, where: str) -> RelationshipKind:
    """An optional relationship kind; "Calls API" and "calls_api" both read as calls_api."""
    text = _optional_text(value, where, "kind")
    if text is None:
        return RelationshipKind.UNSPECIFIED
    try:
        return RelationshipKind(text.strip().casefold().replace(" ", "_"))
    except ValueError as exc:
        raise InvalidKnowledgeError(f"{where}: kind must be one of {_KINDS}.") from exc


def _text_list(value: object, where: str, field: str) -> tuple[str, ...]:
    if value is None:
        return ()
    if isinstance(value, list):
        return tuple(_text(item, where, field) for item in value)
    raise InvalidKnowledgeError(f"{where}: {field} must be a list.")


def _entries(raw: dict[str, Any], key: str) -> list[dict[str, Any]]:
    value = raw.get(key, [])
    if not isinstance(value, list) or not all(isinstance(item, dict) for item in value):
        raise InvalidKnowledgeError(f"'{key}' must be a list of entries.")
    return value


def content_from_mapping(raw: object) -> CatalogueContent:
    """The YAML/JSON shape: ``systems`` with nested capabilities and components, plus
    ``dependencies``."""
    if not isinstance(raw, dict):
        raise InvalidKnowledgeError("The catalogue file must contain an object at the top level.")
    systems = []
    for number, item in enumerate(_entries(raw, "systems"), start=1):
        where = f"systems entry {number}"
        capabilities = []
        raw_capabilities = item.get("capabilities") or []
        if not isinstance(raw_capabilities, list):
            raise InvalidKnowledgeError(f"{where}: capabilities must be a list.")
        for position, capability in enumerate(raw_capabilities, start=1):
            place = f"{where}, capability {position}"
            if not isinstance(capability, dict):
                raise InvalidKnowledgeError(f"{place} must be an object.")
            try:
                capabilities.append(
                    KnowledgeCapability(
                        _text(capability.get("id"), place, "id"),
                        _text(capability.get("name"), place, "name"),
                        _text_list(capability.get("triggers"), place, "triggers"),
                        _optional_text(capability.get("domain"), place, "domain"),
                        _optional_text(capability.get("component"), place, "component"),
                    )
                )
            except InvalidKnowledgeError as exc:
                raise _located(place, exc) from exc
        components = []
        try:
            raw_components = _entries(item, "components")
        except InvalidKnowledgeError as exc:
            raise _located(where, exc) from exc
        for position, component in enumerate(raw_components, start=1):
            place = f"{where}, component {position}"
            try:
                components.append(
                    SystemComponent(
                        id=_text(component.get("id"), place, "id"),
                        name=_text(component.get("name"), place, "name"),
                        name_ar=_optional_text(component.get("name_ar"), place, "name_ar"),
                        description=_optional_text(
                            component.get("description"), place, "description"
                        ),
                        aliases=_text_list(component.get("aliases"), place, "aliases"),
                        technology=_optional_text(component.get("technology"), place, "technology"),
                    )
                )
            except InvalidKnowledgeError as exc:
                raise _located(place, exc) from exc
        try:
            systems.append(
                SystemDefinition(
                    id=_text(item.get("id"), where, "id"),
                    name=_text(item.get("name"), where, "name"),
                    aliases=_text_list(item.get("aliases"), where, "aliases"),
                    capabilities=tuple(capabilities),
                    constraints=_text_list(item.get("constraints"), where, "constraints"),
                    name_ar=_optional_text(item.get("name_ar"), where, "name_ar"),
                    components=tuple(components),
                    description=_optional_text(item.get("description"), where, "description"),
                    landscape_domain_id=_optional_text(
                        item.get("landscape_domain"), where, "landscape_domain"
                    ),
                    owner=_optional_text(item.get("owner"), where, "owner"),
                    external=bool(_flag(item.get("external"), where, "external", False)),
                    roadmap=_optional_text(item.get("roadmap"), where, "roadmap"),
                    placement_from=_optional_text(
                        item.get("placement_from"), where, "placement_from"
                    ),
                    placement_reason=_optional_text(
                        item.get("placement_reason"), where, "placement_reason"
                    ),
                    confidence=_trust(item.get("confidence"), where),
                    source=_optional_text(item.get("source"), where, "source"),
                )
            )
        except InvalidKnowledgeError as exc:
            raise _located(where, exc) from exc
    relationships = []
    for number, item in enumerate(_entries(raw, "dependencies"), start=1):
        where = f"dependencies entry {number}"
        try:
            relationships.append(
                SystemRelationship(
                    _text(item.get("source_system_id"), where, "source_system_id"),
                    _text(item.get("target_system_id"), where, "target_system_id"),
                    _text(item.get("description"), where, "description"),
                    _kind(item.get("kind"), where),
                )
            )
        except InvalidKnowledgeError as exc:
            raise _located(where, exc) from exc
    return _content(
        systems,
        relationships,
        _mapped_domains(raw, "capability_domains", CapabilityDomain),
        _mapped_domains(raw, "landscape_domains", LandscapeDomain),
        _offerings(_entries(raw, "products")),
        _journeys(_entries(raw, "journeys")),
        _channels(_entries(raw, "channels")),
        _sources(_entries(raw, "sources")),
        _conflicts(_entries(raw, "conflicts")),
        _history(_entries(raw, "change_history")) if "change_history" in raw else None,
        _portfolio(_entries(raw, "portfolio")),
    )


class _LocatedError(InvalidKnowledgeError):
    """An error that already says which row or entry it is about."""


def _part[T](place: str, build: Callable[[], T]) -> T:
    try:
        return build()
    except _LocatedError:
        raise
    except InvalidKnowledgeError as exc:
        raise _LocatedError(str(_located(place, exc))) from exc


def _activity(step: dict[str, Any], place: str) -> Activity:
    return Activity(
        number=_text(step.get("number"), place, "number"),
        name=_text(step.get("name"), place, "name"),
        phase=_optional_text(step.get("phase"), place, "phase"),
        track=_optional_text(step.get("track"), place, "track"),
        performing_system_id=_optional_text(step.get("system"), place, "system"),
        supporting_system_ids=_text_list(step.get("supporting"), place, "supporting"),
        system_function=_optional_text(step.get("system_function"), place, "system_function"),
        mode=_optional_text(step.get("mode"), place, "mode"),
        customer_visible=_flag(step.get("customer_visible"), place, "customer_visible", None),
        description=_optional_text(step.get("description"), place, "description"),
        component_ids=_text_list(step.get("components"), place, "components"),
        input=_optional_text(step.get("input"), place, "input"),
        output=_optional_text(step.get("output"), place, "output"),
        etom=_optional_text(step.get("etom"), place, "etom"),
        confidence=_trust(step.get("confidence"), place),
        source=_optional_text(step.get("source"), place, "source"),
        channels=_text_list(step.get("channels"), place, "channels"),
        channel_entry=bool(_flag(step.get("channel_entry"), place, "channel_entry", False)),
        performer=_optional_text(step.get("performer"), place, "performer"),
        point_of_no_return=_optional_text(
            step.get("point_of_no_return"), place, "point_of_no_return"
        ),
        role=_optional_text(step.get("role"), place, "role"),
    )


def _rule(rule: dict[str, Any], place: str) -> FlowRule:
    return FlowRule(
        kind=flow_rule_kind(_text(rule.get("kind"), place, "kind")),
        from_activity=_text(rule.get("from"), place, "from"),
        to_activity=_text(rule.get("to"), place, "to"),
        condition=_optional_text(rule.get("condition"), place, "condition"),
        branch=_optional_text(rule.get("branch"), place, "branch"),
        parallel_group=_optional_text(rule.get("parallel_group"), place, "parallel_group"),
        rejoin_at=_optional_text(rule.get("rejoin_at"), place, "rejoin_at"),
        confidence=_trust(rule.get("confidence"), place),
        source=_optional_text(rule.get("source"), place, "source"),
    )


def _link(link: dict[str, Any], place: str) -> ActivityIntegration:
    return ActivityIntegration(
        from_activity=_text(link.get("from"), place, "from"),
        to_activity=_text(link.get("to"), place, "to"),
        interaction=_optional_text(link.get("interaction"), place, "interaction"),
        interface=_optional_text(link.get("interface"), place, "interface"),
        payload=_optional_text(link.get("payload"), place, "payload"),
        timing=_optional_text(link.get("timing"), place, "timing"),
        correlation_key=_optional_text(link.get("correlation_key"), place, "correlation_key"),
        confidence=_trust(link.get("confidence"), place),
        source=_optional_text(link.get("source"), place, "source"),
        from_system_id=_optional_text(link.get("from_system"), place, "from_system"),
        to_system_id=_optional_text(link.get("to_system"), place, "to_system"),
        via_system_id=_optional_text(link.get("via"), place, "via"),
        purpose=_optional_text(link.get("purpose"), place, "purpose"),
        style=_optional_text(link.get("style"), place, "style"),
        tmf_equivalent=_optional_text(link.get("tmf_equivalent"), place, "tmf_equivalent"),
    )


# Journeys (ADR-0096), read the same way as offerings; each activity, rule and
# integration says which row it is when it is refused.
def _journeys(entries: list[dict[str, Any]]) -> list[Journey]:
    journeys = []
    for number, item in enumerate(entries, start=1):
        where = _where(item, f"journeys entry {number}")
        try:
            activities = [
                _part(place, partial(_activity, step, place))
                for position, step in enumerate(_sub_entries(item, "activities", where), 1)
                for place in [_where(step, f"{where}, activity {position}")]
            ]
            rules = [
                _part(place, partial(_rule, rule, place))
                for position, rule in enumerate(_sub_entries(item, "flow_rules", where), 1)
                for place in [_where(rule, f"{where}, flow rule {position}")]
            ]
            links = [
                _part(place, partial(_link, link, place))
                for position, link in enumerate(_sub_entries(item, "integrations", where), 1)
                for place in [_where(link, f"{where}, integration {position}")]
            ]
            journeys.append(
                Journey(
                    id=_text(item.get("id"), where, "id"),
                    name=_text(item.get("name"), where, "name"),
                    product_id=_optional_text(item.get("product"), where, "product"),
                    order_type_code=_optional_text(item.get("order_type"), where, "order_type"),
                    description=_optional_text(item.get("description"), where, "description"),
                    activities=tuple(activities),
                    flow_rules=tuple(rules),
                    integrations=tuple(links),
                    confidence=_trust(item.get("confidence"), where),
                    source=_optional_text(item.get("source"), where, "source"),
                )
            )
        except _LocatedError:
            raise
        except InvalidKnowledgeError as exc:
            raise _located(where, exc) from exc
    return journeys


def _journey_mapping(journey: Journey) -> dict[str, Any]:
    return {
        "id": journey.id,
        "name": journey.name,
        **_present(
            product=journey.product_id,
            order_type=journey.order_type_code,
            description=journey.description,
        ),
        **_sourced(journey),
        **_present(
            activities=[
                {
                    "number": step.number,
                    "name": step.name,
                    **_present(
                        phase=step.phase,
                        track=step.track,
                        system=step.performing_system_id,
                        supporting=list(step.supporting_system_ids),
                        system_function=step.system_function,
                        mode=step.mode,
                        customer_visible=step.customer_visible,
                        description=step.description,
                        components=list(step.component_ids),
                        input=step.input,
                        output=step.output,
                        etom=step.etom,
                        channels=list(step.channels),
                        channel_entry=step.channel_entry or None,
                        performer=step.performer,
                        point_of_no_return=step.point_of_no_return,
                        role=step.role,
                    ),
                    **_sourced(step),
                }
                for step in journey.activities
            ],
            flow_rules=[
                {
                    "kind": rule.kind.value,
                    "from": rule.from_activity,
                    "to": rule.to_activity,
                    **_present(
                        condition=rule.condition,
                        branch=rule.branch,
                        parallel_group=rule.parallel_group,
                        rejoin_at=rule.rejoin_at,
                    ),
                    **_sourced(rule),
                }
                for rule in journey.flow_rules
            ],
            integrations=[
                {
                    "from": link.from_activity,
                    "to": link.to_activity,
                    **_present(
                        interaction=link.interaction,
                        interface=link.interface,
                        payload=link.payload,
                        timing=link.timing,
                        correlation_key=link.correlation_key,
                        from_system=link.from_system_id,
                        to_system=link.to_system_id,
                        via=link.via_system_id,
                        purpose=link.purpose,
                        style=link.style,
                        tmf_equivalent=link.tmf_equivalent,
                    ),
                    **_sourced(link),
                }
                for link in journey.integrations
            ],
        ),
    }


# Product offerings (ADR-0095). Workbook rows are gathered into the YAML/JSON shape and
# read by the same code, each entry keeping where it came from for its errors.
_YES = frozenset({"yes", "y", "true", "1"})
_NO = frozenset({"no", "n", "false", "0"})
_VALUE_KINDS = frozenset({"value", "values", "customer value"})
_EVENT_KINDS = {
    "milestone": "milestones",
    "status": "statuses",
    "internal status": "statuses",
    "fallout": "fallout",
}
_AUDIENCE_KINDS = frozenset({"audience", "audiences", "who it is for"})


def _flag(value: object, where: str, field: str, default: bool | None) -> bool | None:
    if value is None or (isinstance(value, str) and not value.strip()):
        return default
    if isinstance(value, bool):
        return value
    text = str(value).strip().casefold()
    if text in _YES:
        return True
    if text in _NO:
        return False
    raise InvalidKnowledgeError(f"{where}: {field} must be yes or no.")


def _trust(value: object, where: str) -> SourceConfidence | None:
    """How sure the source is: confirmed, inferred or gap, in any case."""
    text = _optional_text(value, where, "confidence")
    if text is None:
        return None
    try:
        return SourceConfidence(text.casefold())
    except ValueError as exc:
        raise InvalidKnowledgeError(
            f"{where}: confidence must be confirmed, inferred or gap."
        ) from exc


def _where(item: dict[str, Any], fallback: str) -> str:
    return str(item.get("_where") or fallback)


def _sub_entries(item: dict[str, Any], key: str, where: str) -> list[dict[str, Any]]:
    try:
        return _entries(item, key)
    except InvalidKnowledgeError as exc:
        raise _located(where, exc) from exc


def _point(item: dict[str, Any], where: str) -> OfferingPoint:
    return OfferingPoint(
        _text(item.get("name"), where, "name"),
        _optional_text(item.get("description"), where, "description"),
        _trust(item.get("confidence"), where),
        _optional_text(item.get("source"), where, "source"),
    )


def _offerings(entries: list[dict[str, Any]]) -> list[ProductOffering]:
    offerings = []
    for number, item in enumerate(entries, start=1):
        where = _where(item, f"products entry {number}")
        try:
            order_types = [
                OrderType(
                    _text(order.get("code"), place, "code"),
                    _text(order.get("name"), place, "name"),
                    bool(_flag(order.get("enabled"), place, "enabled", True)),
                    _optional_text(order.get("description"), place, "description"),
                    _trust(order.get("confidence"), place),
                    _optional_text(order.get("source"), place, "source"),
                    _text_list(order.get("channels"), place, "channels"),
                )
                for position, order in enumerate(_sub_entries(item, "order_types", where), 1)
                for place in [_where(order, f"{where}, order type {position}")]
            ]
            components = []
            for position, part in enumerate(_sub_entries(item, "components", where), 1):
                place = _where(part, f"{where}, component {position}")
                responsibilities = [
                    ComponentResponsibility(
                        _text(duty.get("system"), spot, "system"),
                        _text(duty.get("role"), spot, "role"),
                        _text(duty.get("description"), spot, "description"),
                        _text_list(duty.get("order_types"), spot, "order_types"),
                        _trust(duty.get("confidence"), spot),
                        _optional_text(duty.get("source"), spot, "source"),
                    )
                    for index, duty in enumerate(_sub_entries(part, "responsibilities", place), 1)
                    for spot in [_where(duty, f"{place}, responsibility {index}")]
                ]
                realisation = [
                    Realisation(
                        _text(layer.get("layer"), spot, "layer"),  # type: ignore[arg-type]
                        _text(layer.get("name"), spot, "name"),
                        _trust(layer.get("confidence"), spot),
                        _optional_text(layer.get("source"), spot, "source"),
                    )
                    for index, layer in enumerate(_sub_entries(part, "realisation", place), 1)
                    for spot in [_where(layer, f"{place}, realisation {index}")]
                ]
                components.append(
                    OfferingComponent(
                        id=_text(part.get("id"), place, "id"),
                        name=_text(part.get("name"), place, "name"),
                        code=_optional_text(part.get("code"), place, "code"),
                        kind=_optional_text(part.get("type"), place, "type"),
                        mandatory=_flag(part.get("mandatory"), place, "mandatory", None),
                        customer_visible=_flag(
                            part.get("customer_visible"), place, "customer_visible", None
                        ),
                        description=_optional_text(part.get("description"), place, "description"),
                        commercial_spec=_optional_text(
                            part.get("commercial_spec"), place, "commercial_spec"
                        ),
                        technical_spec=_optional_text(
                            part.get("technical_spec"), place, "technical_spec"
                        ),
                        technical_details=_optional_text(
                            part.get("technical_details"), place, "technical_details"
                        ),
                        responsibilities=tuple(responsibilities),
                        confidence=_trust(part.get("confidence"), place),
                        source=_optional_text(part.get("source"), place, "source"),
                        realisation=tuple(realisation),
                    )
                )
            offerings.append(
                ProductOffering(
                    id=_text(item.get("id"), where, "id"),
                    name=_text(item.get("name"), where, "name"),
                    code=_optional_text(item.get("code"), where, "code"),
                    family=_optional_text(item.get("family"), where, "family"),
                    version=_optional_text(item.get("version"), where, "version"),
                    lifecycle=_optional_text(item.get("lifecycle"), where, "lifecycle"),
                    proposition=_optional_text(item.get("proposition"), where, "proposition"),
                    rules=_text_list(item.get("rules"), where, "rules"),
                    order_types=tuple(order_types),
                    components=tuple(components),
                    values=tuple(
                        _point(point, _where(point, f"{where}, value {index}"))
                        for index, point in enumerate(_sub_entries(item, "values", where), 1)
                    ),
                    audiences=tuple(
                        _point(point, _where(point, f"{where}, audience {index}"))
                        for index, point in enumerate(_sub_entries(item, "audiences", where), 1)
                    ),
                    confidence=_trust(item.get("confidence"), where),
                    source=_optional_text(item.get("source"), where, "source"),
                    nfrs=tuple(
                        _nfr(nfr, _where(nfr, f"{where}, NFR {index}"))
                        for index, nfr in enumerate(_sub_entries(item, "nfrs", where), 1)
                    ),
                    tracking=_tracking(item, where),
                    lifecycle_notes=_lifecycle_notes(item, where),
                    sources=_text_list(item.get("sources"), where, "sources"),
                    primary_source=_optional_text(
                        item.get("primary_source"), where, "primary_source"
                    ),
                    questions=tuple(
                        _part(spot, partial(_question, raw, spot))
                        for index, raw in enumerate(_sub_entries(item, "questions", where), 1)
                        for spot in [_where(raw, f"{where}, question {index}")]
                    ),
                    decisions=tuple(
                        _part(spot, partial(_decision, raw, spot))
                        for index, raw in enumerate(_sub_entries(item, "decisions", where), 1)
                        for spot in [_where(raw, f"{where}, decision {index}")]
                    ),
                    boundaries=_text_list(item.get("boundaries"), where, "boundaries"),
                    not_used=_text_list(item.get("not_used"), where, "not_used"),
                    portfolio_node_id=_optional_text(item.get("portfolio"), where, "portfolio"),
                    plans=tuple(
                        _part(spot, partial(_plan, raw, spot))
                        for index, raw in enumerate(_sub_entries(item, "plans", where), 1)
                        for spot in [_where(raw, f"{where}, plan {index}")]
                    ),
                    business_rules=tuple(
                        _part(spot, partial(_business_rule, raw, spot))
                        for index, raw in enumerate(_sub_entries(item, "business_rules", where), 1)
                        for spot in [_where(raw, f"{where}, business rule {index}")]
                    ),
                )
            )
        except InvalidKnowledgeError as exc:
            raise _located(where, exc) from exc
    return offerings


def _plan(raw: dict[str, Any], where: str) -> OfferingPlan:
    """A plan: its name and characteristics, given as a mapping of name to value."""
    characteristics = raw.get("characteristics") or {}
    if not isinstance(characteristics, dict):
        raise InvalidKnowledgeError(f"{where}: characteristics must map each name to its value.")
    return OfferingPlan(
        name=_text(raw.get("name"), where, "name"),
        characteristics=tuple(
            PlanCharacteristic(_text(name, where, "characteristic"), _text(value, where, str(name)))
            for name, value in characteristics.items()
        ),
        description=_optional_text(raw.get("description"), where, "description"),
        confidence=_trust(raw.get("confidence"), where),
        source=_optional_text(raw.get("source"), where, "source"),
    )


def _business_rule(raw: dict[str, Any], where: str) -> BusinessRule:
    return BusinessRule(
        id=_text(raw.get("id"), where, "id"),
        statement=_text(raw.get("statement"), where, "statement"),
        kind=_optional_text(raw.get("kind"), where, "kind"),
        confidence=_trust(raw.get("confidence"), where),
        source=_optional_text(raw.get("source"), where, "source"),
    )


def _portfolio(entries: list[dict[str, Any]]) -> list[PortfolioNode]:
    nodes = []
    for number, item in enumerate(entries, start=1):
        where = _where(item, f"portfolio entry {number}")
        try:
            nodes.append(
                PortfolioNode(
                    id=_text(item.get("id"), where, "id"),
                    name=_text(item.get("name"), where, "name"),
                    level=_text(item.get("level"), where, "level"),
                    parent_id=_optional_text(item.get("parent"), where, "parent"),
                    description=_optional_text(item.get("description"), where, "description"),
                    confidence=_trust(item.get("confidence"), where),
                    source=_optional_text(item.get("source"), where, "source"),
                )
            )
        except InvalidKnowledgeError as exc:
            raise _located(where, exc) from exc
    return nodes


def _portfolio_mapping(node: PortfolioNode) -> dict[str, Any]:
    return {
        "id": node.id,
        "name": node.name,
        "level": node.level,
        **_present(parent=node.parent_id, description=node.description),
        **_sourced(node),
    }


def _question(raw: dict[str, Any], where: str) -> OpenQuestion:
    order_types = raw.get("order_types")
    return OpenQuestion(
        _text(raw.get("id"), where, "id"),
        _text(raw.get("text"), where, "text"),
        _optional_text(raw.get("impact"), where, "impact"),
        _trust(raw.get("confidence"), where),
        _optional_text(raw.get("source"), where, "source"),
        _text_list(order_types, where, "order_types")
        if isinstance(order_types, list)
        else _split(order_types, where, "order_types"),
    )


def _decision(raw: dict[str, Any], where: str) -> ArchitectureDecision:
    return ArchitectureDecision(
        _text(raw.get("id"), where, "id"),
        _text(raw.get("title"), where, "title"),
        _optional_text(raw.get("text"), where, "text"),
        _trust(raw.get("confidence"), where),
        _optional_text(raw.get("source"), where, "source"),
    )


def _tracking(item: dict[str, Any], where: str) -> OrderTracking | None:
    raw = item.get("tracking")
    if raw is None:
        return None
    if not isinstance(raw, dict):
        raise InvalidKnowledgeError(f"{where}: tracking must be an entry.")
    place = _where(raw, f"{where}, tracking")

    def each(key: str, label: str) -> list[tuple[dict[str, Any], str]]:
        return [
            (entry, _where(entry, f"{place}, {label} {index}"))
            for index, entry in enumerate(_sub_entries(raw, key, place), 1)
        ]

    return OrderTracking(
        order_types=_text_list(raw.get("order_types"), place, "order_types"),
        scope_note=_optional_text(raw.get("scope_note"), place, "scope_note"),
        not_applicable_note=_optional_text(
            raw.get("not_applicable_note"), place, "not_applicable_note"
        ),
        flows=tuple(
            TrackingFlow(
                _text(flow.get("from_system"), spot, "from_system"),
                _text(flow.get("to_system"), spot, "to_system"),
                _text(flow.get("label"), spot, "label"),
                _optional_text(flow.get("interface"), spot, "interface"),
                _trust(flow.get("confidence"), spot),
                _optional_text(flow.get("source"), spot, "source"),
            )
            for flow, spot in each("flows", "flow")
        ),
        channels=tuple(
            TrackingChannel(
                _text(channel.get("channel"), spot, "channel"),
                _optional_text(channel.get("correlation_key"), spot, "correlation_key"),
                _optional_text(channel.get("ui_system"), spot, "ui_system"),
                _optional_text(channel.get("story"), spot, "story"),
                _optional_text(channel.get("read_system"), spot, "read_system"),
                _optional_text(channel.get("read_interface"), spot, "read_interface"),
                _optional_text(channel.get("ui_note"), spot, "ui_note"),
                _trust(channel.get("confidence"), spot),
                _optional_text(channel.get("source"), spot, "source"),
            )
            for channel, spot in each("channels", "channel")
        ),
        milestones=tuple(_event(event, spot) for event, spot in each("milestones", "milestone")),
        statuses=tuple(_event(event, spot) for event, spot in each("statuses", "status")),
        fallout=tuple(
            FalloutCase(
                _text(case.get("trigger"), spot, "trigger"),
                _optional_text(case.get("handling"), spot, "handling"),
                _trust(case.get("confidence"), spot),
                _optional_text(case.get("source"), spot, "source"),
            )
            for case, spot in each("fallout", "fallout")
        ),
        confidence=_trust(raw.get("confidence"), place),
        source=_optional_text(raw.get("source"), place, "source"),
    )


def _event(item: dict[str, Any], where: str) -> TrackingEvent:
    return TrackingEvent(
        _text(item.get("label"), where, "label"),
        _optional_text(item.get("detail"), where, "detail"),
        _optional_text(item.get("system"), where, "system"),
        _trust(item.get("confidence"), where),
        _optional_text(item.get("source"), where, "source"),
    )


def _lifecycle_notes(item: dict[str, Any], where: str) -> tuple[LifecycleNote, ...]:
    notes = []
    for index, note in enumerate(_sub_entries(item, "lifecycle_notes", where), 1):
        place = _where(note, f"{where}, lifecycle note {index}")
        try:
            blocks = []
            for position, block in enumerate(_sub_entries(note, "blocks", place), 1):
                spot = _where(block, f"{place}, block {position}")
                rows = block.get("rows") or []
                if not isinstance(rows, list) or not all(isinstance(row, list) for row in rows):
                    raise InvalidKnowledgeError(f"{spot}: rows must be a list of lists.")
                blocks.append(
                    NoteBlock(
                        kind=_text(block.get("kind"), spot, "kind"),  # type: ignore[arg-type]
                        title=_optional_text(block.get("title"), spot, "title"),
                        text=_optional_text(block.get("text"), spot, "text"),
                        items=_text_list(block.get("items"), spot, "items"),
                        columns=tuple(
                            _optional_text(cell, spot, "columns") or ""
                            for cell in block.get("columns") or []
                        ),
                        rows=tuple(
                            tuple(_optional_text(cell, spot, "rows") or "" for cell in row)
                            for row in rows
                        ),
                        caption=_optional_text(block.get("caption"), spot, "caption"),
                        confidence=_trust(block.get("confidence"), spot),
                        source=_optional_text(block.get("source"), spot, "source"),
                        to_verify=bool(_flag(block.get("to_verify"), spot, "to_verify", False)),
                    )
                )
            notes.append(
                LifecycleNote(
                    id=_text(note.get("id"), place, "id"),
                    title=_text(note.get("title"), place, "title"),
                    kind=_optional_text(note.get("kind"), place, "kind"),
                    summary=_optional_text(note.get("summary"), place, "summary"),
                    order_types=_text_list(note.get("order_types"), place, "order_types"),
                    channels=_text_list(note.get("channels"), place, "channels"),
                    blocks=tuple(blocks),
                    confidence=_trust(note.get("confidence"), place),
                    source=_optional_text(note.get("source"), place, "source"),
                )
            )
        except InvalidKnowledgeError as exc:
            raise _located(place, exc) from exc
    return tuple(notes)


def _nfr(item: dict[str, Any], where: str) -> OfferingNfr:
    return OfferingNfr(
        _text(item.get("quality"), where, "quality"),
        _text(item.get("coverage"), where, "coverage"),  # type: ignore[arg-type]
        _optional_text(item.get("statement"), where, "statement"),
        _trust(item.get("confidence"), where),
        _optional_text(item.get("source"), where, "source"),
    )


def _present(**values: object) -> dict[str, Any]:
    """The keys that say something; files leave out what is unknown."""
    return {key: value for key, value in values.items() if value not in (None, "", [], ())}


def _sourced(item: Any) -> dict[str, Any]:
    return _present(
        confidence=item.confidence.value if item.confidence else None, source=item.source
    )


def _offering_mapping(offering: ProductOffering) -> dict[str, Any]:
    return {
        "id": offering.id,
        "name": offering.name,
        **_present(
            code=offering.code,
            family=offering.family,
            version=offering.version,
            lifecycle=offering.lifecycle,
            proposition=offering.proposition,
            rules=list(offering.rules),
        ),
        **_sourced(offering),
        **_present(
            order_types=[
                {
                    "code": order.code,
                    "name": order.name,
                    "enabled": order.enabled,
                    **_present(description=order.description, channels=list(order.channels)),
                    **_sourced(order),
                }
                for order in offering.order_types
            ],
            components=[
                {
                    "id": part.id,
                    "name": part.name,
                    **_present(
                        code=part.code,
                        type=part.kind,
                        mandatory=part.mandatory,
                        customer_visible=part.customer_visible,
                        description=part.description,
                        commercial_spec=part.commercial_spec,
                        technical_spec=part.technical_spec,
                        technical_details=part.technical_details,
                    ),
                    **_sourced(part),
                    **_present(
                        responsibilities=[
                            {
                                "system": duty.system_id,
                                "role": duty.role,
                                "description": duty.description,
                                **_present(order_types=list(duty.order_types)),
                                **_sourced(duty),
                            }
                            for duty in part.responsibilities
                        ],
                        realisation=[
                            {"layer": item.layer.value, "name": item.name, **_sourced(item)}
                            for item in part.realisation
                        ],
                    ),
                }
                for part in offering.components
            ],
            values=[
                {"name": point.name, **_present(description=point.description), **_sourced(point)}
                for point in offering.values
            ],
            audiences=[
                {"name": point.name, **_present(description=point.description), **_sourced(point)}
                for point in offering.audiences
            ],
            nfrs=[
                {
                    "quality": nfr.quality,
                    "coverage": nfr.coverage.value,
                    **_present(statement=nfr.statement),
                    **_sourced(nfr),
                }
                for nfr in offering.nfrs
            ],
        ),
        **({"tracking": _tracking_mapping(offering.tracking)} if offering.tracking else {}),
        **_present(lifecycle_notes=[_note_mapping(note) for note in offering.lifecycle_notes]),
        **_present(
            sources=list(offering.sources),
            primary_source=offering.primary_source,
            questions=[
                {
                    "id": item.id,
                    "text": item.text,
                    **_present(impact=item.impact, order_types=list(item.order_types)),
                    **_sourced(item),
                }
                for item in offering.questions
            ],
            decisions=[
                {"id": item.id, "title": item.title, **_present(text=item.text), **_sourced(item)}
                for item in offering.decisions
            ],
            boundaries=list(offering.boundaries),
            not_used=list(offering.not_used),
            portfolio=offering.portfolio_node_id,
            plans=[
                {
                    "name": plan.name,
                    **_present(
                        characteristics={item.name: item.value for item in plan.characteristics},
                        description=plan.description,
                    ),
                    **_sourced(plan),
                }
                for plan in offering.plans
            ],
            business_rules=[
                {
                    "id": rule.id,
                    "statement": rule.statement,
                    **_present(kind=rule.kind),
                    **_sourced(rule),
                }
                for rule in offering.business_rules
            ],
        ),
    }


def _note_mapping(note: LifecycleNote) -> dict[str, Any]:
    return {
        "id": note.id,
        "title": note.title,
        **_present(
            kind=note.kind,
            summary=note.summary,
            order_types=list(note.order_types),
            channels=list(note.channels),
        ),
        **_sourced(note),
        **_present(
            blocks=[
                {
                    "kind": block.kind.value,
                    **_present(
                        title=block.title,
                        text=block.text,
                        items=list(block.items),
                        columns=list(block.columns),
                        rows=[list(row) for row in block.rows],
                        caption=block.caption,
                    ),
                    **_sourced(block),
                    **({"to_verify": True} if block.to_verify else {}),
                }
                for block in note.blocks
            ]
        ),
    }


def _tracking_mapping(tracking: OrderTracking) -> dict[str, Any]:
    def event(item: TrackingEvent) -> dict[str, Any]:
        return {
            "label": item.label,
            **_present(detail=item.detail, system=item.system_id),
            **_sourced(item),
        }

    return {
        **_present(
            order_types=list(tracking.order_types),
            scope_note=tracking.scope_note,
            not_applicable_note=tracking.not_applicable_note,
        ),
        **_sourced(tracking),
        **_present(
            flows=[
                {
                    "from_system": flow.from_system_id,
                    "to_system": flow.to_system_id,
                    "label": flow.label,
                    **_present(interface=flow.interface),
                    **_sourced(flow),
                }
                for flow in tracking.flows
            ],
            channels=[
                {
                    "channel": channel.channel_id,
                    **_present(
                        correlation_key=channel.correlation_key,
                        ui_system=channel.ui_system_id,
                        story=channel.story,
                        read_system=channel.read_system_id,
                        read_interface=channel.read_interface,
                        ui_note=channel.ui_note,
                    ),
                    **_sourced(channel),
                }
                for channel in tracking.channels
            ],
            milestones=[event(item) for item in tracking.milestones],
            statuses=[event(item) for item in tracking.statuses],
            fallout=[
                {
                    "trigger": case.trigger,
                    **_present(handling=case.handling),
                    **_sourced(case),
                }
                for case in tracking.fallout
            ],
        ),
    }


def _mapped_domains[DomainT: (CapabilityDomain, LandscapeDomain)](
    raw: dict[str, Any], key: str, kind: type[DomainT]
) -> list[DomainT]:
    domains = []
    for number, item in enumerate(_entries(raw, key), start=1):
        where = f"{key} entry {number}"
        try:
            domains.append(
                kind(
                    _text(item.get("id"), where, "id"),
                    _text(item.get("name"), where, "name"),
                    _optional_text(item.get("name_ar"), where, "name_ar"),
                    _optional_text(item.get("parent_id"), where, "parent_id"),
                    _optional_text(item.get("description"), where, "description"),
                )
            )
        except InvalidKnowledgeError as exc:
            raise _located(where, exc) from exc
    return domains


def _content(
    systems: list[SystemDefinition],
    relationships: list[SystemRelationship],
    domains: list[CapabilityDomain],
    landscape: list[LandscapeDomain],
    offerings: list[ProductOffering],
    journeys: list[Journey],
    channels: list[Channel] | None = None,
    sources: list[KnowledgeSource] | None = None,
    conflicts: list[SourceConflict] | None = None,
    history: list[ChangeRequestRecord] | None = None,
    portfolio: list[PortfolioNode] | None = None,
) -> CatalogueContent:
    """The file's content, refusing a placement in a domain the file does not list."""
    known = {item.id for item in domains}
    places = {item.id for item in landscape}
    for system in systems:
        for capability in system.capabilities:
            if capability.domain_id is not None and capability.domain_id not in known:
                raise InvalidKnowledgeError(
                    f"{system.id}, capability {capability.id}: domain {capability.domain_id!r} "
                    "is not listed among the capability domains."
                )
        if system.landscape_domain_id is not None and system.landscape_domain_id not in places:
            raise InvalidKnowledgeError(
                f"{system.id}: landscape domain {system.landscape_domain_id!r} is not listed "
                "among the landscape domains."
            )
    return CatalogueContent(
        tuple(systems),
        tuple(relationships),
        tuple(domains),
        tuple(landscape),
        tuple(offerings),
        tuple(journeys),
        tuple(channels or ()),
        tuple(sources or ()),
        tuple(conflicts or ()),
        None if history is None else tuple(history),
        tuple(portfolio or ()),
    )


# Channels (requirement-portal ADR-0101, step 3): where orders are placed.
def _channels(entries: list[dict[str, Any]]) -> list[Channel]:
    channels = []
    for number, item in enumerate(entries, start=1):
        where = _where(item, f"channels entry {number}")
        try:
            channels.append(
                Channel(
                    id=_text(item.get("id"), where, "id"),
                    name=_text(item.get("name"), where, "name"),
                    kind=_optional_text(item.get("kind"), where, "kind"),
                    entry_system_id=_optional_text(item.get("entry_system"), where, "entry_system"),
                    description=_optional_text(item.get("description"), where, "description"),
                    confidence=_trust(item.get("confidence"), where),
                    source=_optional_text(item.get("source"), where, "source"),
                )
            )
        except InvalidKnowledgeError as exc:
            raise _located(where, exc) from exc
    return channels


def _channel_mapping(channel: Channel) -> dict[str, Any]:
    return {
        "id": channel.id,
        "name": channel.name,
        **_present(
            kind=channel.kind,
            entry_system=channel.entry_system_id,
            description=channel.description,
        ),
        **_sourced(channel),
    }


# Governance (requirement-portal ADR-0101, step 5): the source register and the conflicts
# between sources.
def _sources(entries: list[dict[str, Any]]) -> list[KnowledgeSource]:
    sources = []
    for number, item in enumerate(entries, start=1):
        where = _where(item, f"sources entry {number}")
        sources.append(_part(where, partial(_source, item, where)))
    return sources


def _source(item: dict[str, Any], where: str) -> KnowledgeSource:
    return KnowledgeSource(
        id=_text(item.get("id"), where, "id"),
        title=_text(item.get("title"), where, "title"),
        level=_text(item.get("level"), where, "level"),  # type: ignore[arg-type]
        short=_optional_text(item.get("short"), where, "short"),
        version=_optional_text(item.get("version"), where, "version"),
        file=_optional_text(item.get("file"), where, "file"),
        supplied=bool(_flag(item.get("supplied"), where, "supplied", True)),
        authority=_optional_text(item.get("authority"), where, "authority"),
        scope=_optional_text(item.get("scope"), where, "scope"),
        boundary=_optional_text(item.get("boundary"), where, "boundary"),
    )


def _side(raw: object, where: str, field: str) -> ConflictSide:
    if not isinstance(raw, dict):
        raise InvalidKnowledgeError(f"{where}: {field} must be an object.")
    return ConflictSide(
        _text(raw.get("source"), where, f"{field}.source"),
        _text(raw.get("statement"), where, f"{field}.statement"),
        _optional_text(raw.get("reference"), where, f"{field}.reference"),
    )


def _conflicts(entries: list[dict[str, Any]]) -> list[SourceConflict]:
    conflicts = []
    for number, item in enumerate(entries, start=1):
        where = _where(item, f"conflicts entry {number}")
        conflicts.append(_part(where, partial(_conflict, item, where)))
    return conflicts


def _conflict(item: dict[str, Any], where: str) -> SourceConflict:
    return SourceConflict(
        id=_text(item.get("id"), where, "id"),
        title=_text(item.get("title"), where, "title"),
        a=_side(item.get("a"), where, "a"),
        b=_side(item.get("b"), where, "b"),
        scope=tuple(
            ConflictScope(
                _text(scope.get("product"), spot, "product"),
                _text_list(scope.get("order_types"), spot, "order_types"),
                _optional_text(scope.get("question"), spot, "question"),
            )
            for index, scope in enumerate(_sub_entries(item, "scope", where), 1)
            for spot in [f"{where}, scope {index}"]
        ),
        difference=_optional_text(item.get("difference"), where, "difference"),
        impact=_optional_text(item.get("impact"), where, "impact"),
        decision=_optional_text(item.get("decision"), where, "decision"),
        confidence=_trust(item.get("confidence"), where),
        source=_optional_text(item.get("source"), where, "source"),
    )


def _source_mapping(source: KnowledgeSource) -> dict[str, Any]:
    return {
        "id": source.id,
        "title": source.title,
        "level": source.level.value,
        **_present(short=source.short, version=source.version, file=source.file),
        **({} if source.supplied else {"supplied": False}),
        **_present(authority=source.authority, scope=source.scope, boundary=source.boundary),
    }


def _conflict_mapping(conflict: SourceConflict) -> dict[str, Any]:
    def side(item: ConflictSide) -> dict[str, Any]:
        return {
            "source": item.source_id,
            **_present(reference=item.reference),
            "statement": item.statement,
        }

    return {
        "id": conflict.id,
        "title": conflict.title,
        "a": side(conflict.a),
        "b": side(conflict.b),
        **_present(
            scope=[
                {
                    "product": item.product_id,
                    **_present(order_types=list(item.order_types), question=item.question_id),
                }
                for item in conflict.scope
            ],
            difference=conflict.difference,
            impact=conflict.impact,
            decision=conflict.decision,
        ),
        **_sourced(conflict),
    }


def _domain_mapping(item: CapabilityDomain | LandscapeDomain) -> dict[str, Any]:
    return {
        "id": item.id,
        "name": item.name,
        **({"name_ar": item.name_ar} if item.name_ar else {}),
        **({"parent_id": item.parent_id} if item.parent_id else {}),
        **({"description": item.description} if item.description else {}),
    }


def _component_mapping(component: SystemComponent) -> dict[str, Any]:
    return {
        "id": component.id,
        "name": component.name,
        **({"name_ar": component.name_ar} if component.name_ar else {}),
        **({"aliases": list(component.aliases)} if component.aliases else {}),
        **({"technology": component.technology} if component.technology else {}),
        **({"description": component.description} if component.description else {}),
    }


def release_to_mapping(release: ArchitectureKnowledge) -> dict[str, Any]:
    return {
        "version": release.id,
        "systems": [
            {
                "id": item.id,
                "name": item.name,
                **({"name_ar": item.name_ar} if item.name_ar else {}),
                "aliases": list(item.aliases),
                **({"description": item.description} if item.description else {}),
                **(
                    {"landscape_domain": item.landscape_domain_id}
                    if item.landscape_domain_id
                    else {}
                ),
                "capabilities": [
                    {
                        "id": cap.id,
                        "name": cap.name,
                        "triggers": list(cap.triggers),
                        **({"domain": cap.domain_id} if cap.domain_id else {}),
                        **({"component": cap.component_id} if cap.component_id else {}),
                    }
                    for cap in item.capabilities
                ],
                **(
                    {"components": [_component_mapping(part) for part in item.components]}
                    if item.components
                    else {}
                ),
                **({"constraints": list(item.constraints)} if item.constraints else {}),
                **_present(
                    owner=item.owner,
                    external=item.external or None,
                    roadmap=item.roadmap,
                    placement_from=item.placement_from,
                    placement_reason=item.placement_reason,
                ),
                **_sourced(item),
            }
            for item in release.systems
        ],
        "dependencies": [
            {
                "source_system_id": item.source_system_id,
                "target_system_id": item.target_system_id,
                "description": item.description,
                **(
                    {"kind": item.kind.value}
                    if item.kind is not RelationshipKind.UNSPECIFIED
                    else {}
                ),
            }
            for item in release.relationships
        ],
        **(
            {"capability_domains": [_domain_mapping(item) for item in release.capability_domains]}
            if release.capability_domains
            else {}
        ),
        **(
            {"landscape_domains": [_domain_mapping(item) for item in release.landscape_domains]}
            if release.landscape_domains
            else {}
        ),
        **(
            {"products": [_offering_mapping(item) for item in release.products]}
            if release.products
            else {}
        ),
        **(
            {"journeys": [_journey_mapping(item) for item in release.journeys]}
            if release.journeys
            else {}
        ),
        **(
            {"channels": [_channel_mapping(item) for item in release.channels]}
            if release.channels
            else {}
        ),
        **_present(
            sources=[_source_mapping(item) for item in release.sources],
            conflicts=[_conflict_mapping(item) for item in release.conflicts],
            change_history=[_history_mapping(item) for item in release.change_history],
            portfolio=[_portfolio_mapping(item) for item in release.portfolio],
        ),
    }


def _check_archive(content: bytes) -> None:
    """Refuse archive bombs before openpyxl inflates anything."""
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            parts = archive.infolist()
    except zipfile.BadZipFile as exc:
        raise InvalidKnowledgeError("The Excel file is damaged or not an .xlsx workbook.") from exc
    if len(parts) > MAX_OFFICE_PARTS or (
        sum(item.file_size for item in parts) > MAX_OFFICE_UNCOMPRESSED_BYTES
    ):
        raise InvalidKnowledgeError("The Excel file is too large to import.")


def _rows(workbook: Any, sheet: str) -> Iterator[tuple[int, dict[str, object]]]:
    if sheet not in workbook.sheetnames:
        if sheet in _REQUIRED_SHEETS:
            raise InvalidKnowledgeError(f"The workbook needs a '{sheet}' sheet.")
        return
    rows = workbook[sheet].iter_rows(values_only=True)
    header_row = next(rows, None) or ()
    headers = [str(value).strip().casefold() if value is not None else "" for value in header_row]
    if any(name not in headers for name in _REQUIRED_HEADERS[sheet]):
        raise InvalidKnowledgeError(
            f"{sheet} row 1 must have the headers {', '.join(_HEADERS[sheet])}."
        )
    for number, values in enumerate(rows, start=2):
        if number > MAX_WORKBOOK_ROWS:
            raise InvalidKnowledgeError(f"{sheet} has more rows than can be imported.")
        cells: dict[str, object] = {
            header: value for header, value in zip(headers, values, strict=False) if header
        }
        if all(value is None or str(value).strip() == "" for value in cells.values()):
            continue
        yield number, cells


def _split(value: object, where: str, field: str) -> tuple[str, ...]:
    text = _optional_text(value, where, field)
    if text is None:
        return ()
    return tuple(part.strip() for part in text.split(_LIST_SEPARATOR) if part.strip())


def _read_workbook(content: bytes) -> CatalogueContent:
    _check_archive(content)
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except (InvalidFileException, KeyError, OSError, ValueError) as exc:
        raise InvalidKnowledgeError("The Excel file could not be opened.") from exc
    try:
        systems: dict[str, dict[str, Any]] = {}
        for number, cells in _rows(workbook, SYSTEMS):
            where = f"{SYSTEMS} row {number}"
            system_id = _text(cells.get("system_id"), where, "system_id")
            if system_id in systems:
                raise InvalidKnowledgeError(f"{where}: system {system_id!r} is listed twice.")
            systems[system_id] = {
                "where": where,
                "name": _text(cells.get("name"), where, "name"),
                "name_ar": _optional_text(cells.get("name_ar"), where, "name_ar"),
                "aliases": _split(cells.get("aliases"), where, "aliases"),
                "description": _optional_text(cells.get("description"), where, "description"),
                "landscape_domain_id": _optional_text(
                    cells.get("landscape_domain_id"), where, "landscape_domain_id"
                ),
                "owner": _optional_text(cells.get("owner"), where, "owner"),
                "external": bool(_flag(cells.get("external"), where, "external", False)),
                "roadmap": _optional_text(cells.get("roadmap"), where, "roadmap"),
                "placement_from": _optional_text(
                    cells.get("placement_from"), where, "placement_from"
                ),
                "placement_reason": _optional_text(
                    cells.get("placement_reason"), where, "placement_reason"
                ),
                "confidence": _trust(cells.get("confidence"), where),
                "source": _optional_text(cells.get("source"), where, "source"),
                "capabilities": [],
                "constraints": [],
                "components": [],
            }

        def owner(cells: dict[str, object], where: str) -> dict[str, Any]:
            system_id = _text(cells.get("system_id"), where, "system_id")
            if system_id not in systems:
                raise InvalidKnowledgeError(
                    f"{where}: system {system_id!r} is not listed on the {SYSTEMS} sheet."
                )
            return systems[system_id]

        for number, cells in _rows(workbook, CAPABILITIES):
            where = f"{CAPABILITIES} row {number}"
            try:
                owner(cells, where)["capabilities"].append(
                    KnowledgeCapability(
                        _text(cells.get("capability_id"), where, "capability_id"),
                        _text(cells.get("name"), where, "name"),
                        _split(cells.get("triggers"), where, "triggers"),
                        _optional_text(cells.get("domain_id"), where, "domain_id"),
                        _optional_text(cells.get("component_id"), where, "component_id"),
                    )
                )
            except InvalidKnowledgeError as exc:
                raise _located(where, exc) from exc
        for number, cells in _rows(workbook, COMPONENTS):
            where = f"{COMPONENTS} row {number}"
            try:
                owner(cells, where)["components"].append(
                    SystemComponent(
                        id=_text(cells.get("component_id"), where, "component_id"),
                        name=_text(cells.get("name"), where, "name"),
                        name_ar=_optional_text(cells.get("name_ar"), where, "name_ar"),
                        description=_optional_text(cells.get("description"), where, "description"),
                        aliases=_split(cells.get("aliases"), where, "aliases"),
                        technology=_optional_text(cells.get("technology"), where, "technology"),
                    )
                )
            except InvalidKnowledgeError as exc:
                raise _located(where, exc) from exc
        for number, cells in _rows(workbook, CONSTRAINTS):
            where = f"{CONSTRAINTS} row {number}"
            owner(cells, where)["constraints"].append(
                _text(cells.get("constraint"), where, "constraint")
            )
        relationships = []
        for number, cells in _rows(workbook, RELATIONSHIPS):
            where = f"{RELATIONSHIPS} row {number}"
            try:
                relationships.append(
                    SystemRelationship(
                        _text(cells.get("source_system_id"), where, "source_system_id"),
                        _text(cells.get("target_system_id"), where, "target_system_id"),
                        _text(cells.get("description"), where, "description"),
                        _kind(cells.get("kind"), where),
                    )
                )
            except InvalidKnowledgeError as exc:
                raise _located(where, exc) from exc
        domains = _sheet_domains(workbook, DOMAINS, CapabilityDomain)
        landscape = _sheet_domains(workbook, LANDSCAPE, LandscapeDomain)
        offerings = _offerings(_sheet_offerings(workbook))
        journeys = _journeys(_sheet_journeys(workbook))
        channels = _channels(
            [
                {
                    **cells,
                    "_where": f"{CHANNELS} row {number}",
                    "id": cells.get("channel_id"),
                    "entry_system": cells.get("entry_system_id"),
                }
                for number, cells in _rows(workbook, CHANNELS)
            ]
        )
        sources = _sources(
            [
                {**cells, "_where": f"{SOURCES} row {number}", "id": cells.get("source_id")}
                for number, cells in _rows(workbook, SOURCES)
            ]
        )
        conflicts = _conflicts(_sheet_conflicts(workbook))
        history = (
            _history(_sheet_history(workbook)) if CHANGE_HISTORY in workbook.sheetnames else None
        )
        portfolio = _portfolio(
            [
                {
                    **cells,
                    "_where": f"{PORTFOLIO} row {number}",
                    "id": cells.get("node_id"),
                    "parent": cells.get("parent_id"),
                }
                for number, cells in _rows(workbook, PORTFOLIO)
            ]
        )
    finally:
        workbook.close()
    definitions = []
    for system_id, data in systems.items():
        try:
            definitions.append(
                SystemDefinition(
                    id=system_id,
                    name=data["name"],
                    aliases=data["aliases"],
                    capabilities=tuple(data["capabilities"]),
                    constraints=tuple(data["constraints"]),
                    name_ar=data["name_ar"],
                    components=tuple(data["components"]),
                    description=data["description"],
                    landscape_domain_id=data["landscape_domain_id"],
                    owner=data["owner"],
                    external=data["external"],
                    roadmap=data["roadmap"],
                    placement_from=data["placement_from"],
                    placement_reason=data["placement_reason"],
                    confidence=data["confidence"],
                    source=data["source"],
                )
            )
        except InvalidKnowledgeError as exc:
            raise _located(data["where"], exc) from exc
    return _content(
        definitions,
        relationships,
        domains,
        landscape,
        offerings,
        journeys,
        channels,
        sources,
        conflicts,
        history,
        portfolio,
    )


def _sheet_journeys(workbook: Any) -> list[dict[str, Any]]:
    """The journey sheets gathered into the YAML/JSON shape, each entry knowing its row."""
    journeys: dict[str, dict[str, Any]] = {}
    for number, cells in _rows(workbook, JOURNEYS):
        where = f"{JOURNEYS} row {number}"
        journey_id = _text(cells.get("journey_id"), where, "journey_id")
        if journey_id in journeys:
            raise InvalidKnowledgeError(f"{where}: journey {journey_id!r} is listed twice.")
        journeys[journey_id] = {
            "_where": where,
            "id": journey_id,
            "name": cells.get("name"),
            "product": cells.get("product_id"),
            "order_type": cells.get("order_type_code"),
            "description": cells.get("description"),
            "confidence": cells.get("confidence"),
            "source": cells.get("source"),
            "activities": [],
            "flow_rules": [],
            "integrations": [],
        }

    def journey(cells: dict[str, object], where: str) -> dict[str, Any]:
        journey_id = _text(cells.get("journey_id"), where, "journey_id")
        if journey_id not in journeys:
            raise InvalidKnowledgeError(
                f"{where}: journey {journey_id!r} is not listed on the {JOURNEYS} sheet."
            )
        return journeys[journey_id]

    for number, cells in _rows(workbook, ACTIVITIES):
        where = f"{ACTIVITIES} row {number}"
        journey(cells, where)["activities"].append(
            {
                **cells,
                "_where": where,
                "system": cells.get("performing_system_id"),
                "supporting": list(
                    _split(cells.get("supporting_system_ids"), where, "supporting_system_ids")
                ),
                "components": list(_split(cells.get("component_ids"), where, "component_ids")),
                "channels": list(_split(cells.get("channels"), where, "channels")),
            }
        )
    for number, cells in _rows(workbook, FLOW_RULES):
        where = f"{FLOW_RULES} row {number}"
        journey(cells, where)["flow_rules"].append(
            {
                **cells,
                "_where": where,
                "from": cells.get("from_activity"),
                "to": cells.get("to_activity"),
            }
        )
    for number, cells in _rows(workbook, ACTIVITY_INTEGRATIONS):
        where = f"{ACTIVITY_INTEGRATIONS} row {number}"
        journey(cells, where)["integrations"].append(
            {
                **cells,
                "_where": where,
                "from": cells.get("from_activity"),
                "to": cells.get("to_activity"),
            }
        )
    return list(journeys.values())


def _sheet_offerings(workbook: Any) -> list[dict[str, Any]]:
    """The offering sheets gathered into the YAML/JSON shape, each entry knowing its row."""
    offerings: dict[str, dict[str, Any]] = {}
    for number, cells in _rows(workbook, PRODUCTS):
        where = f"{PRODUCTS} row {number}"
        product_id = _text(cells.get("product_id"), where, "product_id")
        if product_id in offerings:
            raise InvalidKnowledgeError(f"{where}: product {product_id!r} is listed twice.")
        offerings[product_id] = {
            **{
                key: cells.get(key)
                for key in _HEADERS[PRODUCTS][1:]
                if key not in {"rules", "sources"}
            },
            "_where": where,
            "id": product_id,
            "rules": list(_split(cells.get("rules"), where, "rules")),
            "sources": list(_split(cells.get("sources"), where, "sources")),
            "questions": [],
            "decisions": [],
            "boundaries": [],
            "not_used": [],
            "order_types": [],
            "components": [],
            "values": [],
            "audiences": [],
            "nfrs": [],
            "lifecycle_notes": [],
            "plans": [],
            "business_rules": [],
        }

    def offering(cells: dict[str, object], where: str) -> dict[str, Any]:
        product_id = _text(cells.get("product_id"), where, "product_id")
        if product_id not in offerings:
            raise InvalidKnowledgeError(
                f"{where}: product {product_id!r} is not listed on the {PRODUCTS} sheet."
            )
        return offerings[product_id]

    for number, cells in _rows(workbook, ORDER_TYPES):
        where = f"{ORDER_TYPES} row {number}"
        offering(cells, where)["order_types"].append(
            {
                **cells,
                "_where": where,
                "channels": list(_split(cells.get("channels"), where, "channels")),
            }
        )
    parts: dict[tuple[str, str], dict[str, Any]] = {}
    for number, cells in _rows(workbook, OFFERING_COMPONENTS):
        where = f"{OFFERING_COMPONENTS} row {number}"
        owner = offering(cells, where)
        part = {
            **cells,
            "_where": where,
            "id": cells.get("component_id"),
            "responsibilities": [],
            "realisation": [],
        }
        parts[(owner["id"], _text(cells.get("component_id"), where, "component_id"))] = part
        owner["components"].append(part)
    for number, cells in _rows(workbook, RESPONSIBILITIES):
        where = f"{RESPONSIBILITIES} row {number}"
        key = (
            offering(cells, where)["id"],
            _text(cells.get("component_id"), where, "component_id"),
        )
        if key not in parts:
            raise InvalidKnowledgeError(
                f"{where}: component {key[1]!r} is not listed on the {OFFERING_COMPONENTS} sheet."
            )
        parts[key]["responsibilities"].append(
            {
                **cells,
                "_where": where,
                "system": cells.get("system_id"),
                "order_types": list(_split(cells.get("order_types"), where, "order_types")),
            }
        )
    for number, cells in _rows(workbook, REALISATION):
        where = f"{REALISATION} row {number}"
        key = (
            offering(cells, where)["id"],
            _text(cells.get("component_id"), where, "component_id"),
        )
        if key not in parts:
            raise InvalidKnowledgeError(
                f"{where}: component {key[1]!r} is not listed on the {OFFERING_COMPONENTS} sheet."
            )
        parts[key]["realisation"].append({**cells, "_where": where})
    for number, cells in _rows(workbook, NFRS):
        where = f"{NFRS} row {number}"
        offering(cells, where)["nfrs"].append({**cells, "_where": where})
    plans: dict[tuple[str, str], dict[str, Any]] = {}
    for number, cells in _rows(workbook, PLANS):
        where = f"{PLANS} row {number}"
        owner = offering(cells, where)
        plan = {**cells, "_where": where, "name": cells.get("plan"), "characteristics": {}}
        plans[(owner["id"], _text(cells.get("plan"), where, "plan"))] = plan
        owner["plans"].append(plan)
    for number, cells in _rows(workbook, PLAN_CHARACTERISTICS):
        where = f"{PLAN_CHARACTERISTICS} row {number}"
        key = (offering(cells, where)["id"], _text(cells.get("plan"), where, "plan"))
        if key not in plans:
            raise InvalidKnowledgeError(f"{where}: plan {key[1]!r} is not listed on {PLANS}.")
        name = _text(cells.get("characteristic"), where, "characteristic")
        if name in plans[key]["characteristics"]:
            raise InvalidKnowledgeError(f"{where}: {name!r} is given twice for {key[1]!r}.")
        plans[key]["characteristics"][name] = cells.get("value")
    for number, cells in _rows(workbook, BUSINESS_RULES):
        where = f"{BUSINESS_RULES} row {number}"
        offering(cells, where)["business_rules"].append(
            {**cells, "_where": where, "id": cells.get("rule_id")}
        )

    def tracking(cells: dict[str, object], where: str) -> dict[str, Any]:
        owner = offering(cells, where)
        if owner.get("tracking") is None:
            owner["tracking"] = {
                "_where": where,
                "flows": [],
                "channels": [],
                "milestones": [],
                "statuses": [],
                "fallout": [],
            }
        found: dict[str, Any] = owner["tracking"]
        return found

    notes: dict[tuple[str, str], dict[str, Any]] = {}
    for number, cells in _rows(workbook, LIFECYCLE_NOTES):
        where = f"{LIFECYCLE_NOTES} row {number}"
        owner = offering(cells, where)
        note = {
            **cells,
            "_where": where,
            "id": cells.get("note_id"),
            "order_types": list(_split(cells.get("order_types"), where, "order_types")),
            "channels": list(_split(cells.get("channels"), where, "channels")),
            "blocks": [],
        }
        notes[(owner["id"], _text(cells.get("note_id"), where, "note_id"))] = note
        owner["lifecycle_notes"].append(note)
    for number, cells in _rows(workbook, LIFECYCLE_BLOCKS):
        where = f"{LIFECYCLE_BLOCKS} row {number}"
        key = (offering(cells, where)["id"], _text(cells.get("note_id"), where, "note_id"))
        if key not in notes:
            raise InvalidKnowledgeError(
                f"{where}: note {key[1]!r} is not listed on the {LIFECYCLE_NOTES} sheet."
            )
        blocks = notes[key]["blocks"]
        kind = _text(cells.get("kind"), where, "kind").casefold()
        row = [cells.get(name) for name in _CELLS]
        while row and (row[-1] is None or str(row[-1]).strip() == ""):
            row.pop()
        if kind in {"item", "row"}:
            opener = "list" if kind == "item" else "table"
            if not blocks or blocks[-1].get("kind") != opener:
                what = "a list item" if kind == "item" else "a table row"
                raise InvalidKnowledgeError(f"{where}: {what} must follow its {opener}.")
            if kind == "item":
                blocks[-1]["items"].append(cells.get("text"))
            else:
                blocks[-1]["rows"].append(["" if cell is None else cell for cell in row])
            continue
        if kind not in {"text", "list", "table"}:
            raise InvalidKnowledgeError(f"{where}: kind must be text, list, item, table or row.")
        blocks.append(
            {
                **cells,
                "_where": where,
                "kind": kind,
                "text": cells.get("text") if kind == "text" else None,
                "caption": cells.get("text") if kind == "table" else None,
                "items": [],
                "columns": row if kind == "table" else [],
                "rows": [],
            }
        )
    for number, cells in _rows(workbook, TRACKING):
        where = f"{TRACKING} row {number}"
        if offering(cells, where).get("tracking") is not None:
            raise InvalidKnowledgeError(f"{where}: a product's tracking is one row.")
        tracking(cells, where).update(
            {
                **cells,
                "order_types": list(_split(cells.get("order_types"), where, "order_types")),
            }
        )
    for number, cells in _rows(workbook, TRACKING_FLOWS):
        where = f"{TRACKING_FLOWS} row {number}"
        tracking(cells, where)["flows"].append(
            {
                **cells,
                "_where": where,
                "from_system": cells.get("from_system_id"),
                "to_system": cells.get("to_system_id"),
            }
        )
    for number, cells in _rows(workbook, TRACKING_CHANNELS):
        where = f"{TRACKING_CHANNELS} row {number}"
        tracking(cells, where)["channels"].append(
            {
                **cells,
                "_where": where,
                "channel": cells.get("channel_id"),
                "ui_system": cells.get("ui_system_id"),
                "read_system": cells.get("read_system_id"),
            }
        )
    for number, cells in _rows(workbook, TRACKING_EVENTS):
        where = f"{TRACKING_EVENTS} row {number}"
        kind = _text(cells.get("kind"), where, "kind").casefold()
        if kind not in _EVENT_KINDS:
            raise InvalidKnowledgeError(f"{where}: kind must be milestone, status or fallout.")
        entry = {**cells, "_where": where, "system": cells.get("system_id")}
        if kind == "fallout":
            entry = {**entry, "trigger": cells.get("label"), "handling": cells.get("detail")}
        tracking(cells, where)[_EVENT_KINDS[kind]].append(entry)
    for number, cells in _rows(workbook, PRODUCT_POINTS):
        where = f"{PRODUCT_POINTS} row {number}"
        kind = (_text(cells.get("kind"), where, "kind")).casefold()
        if kind not in _VALUE_KINDS | _AUDIENCE_KINDS:
            raise InvalidKnowledgeError(f"{where}: kind must be value or audience.")
        bucket = "values" if kind in _VALUE_KINDS else "audiences"
        offering(cells, where)[bucket].append({**cells, "_where": where})
    for number, cells in _rows(workbook, QUESTIONS):
        where = f"{QUESTIONS} row {number}"
        offering(cells, where)["questions"].append(
            {**cells, "_where": where, "id": cells.get("question_id")}
        )
    for number, cells in _rows(workbook, DECISIONS):
        where = f"{DECISIONS} row {number}"
        offering(cells, where)["decisions"].append(
            {**cells, "_where": where, "id": cells.get("decision_id")}
        )
    for number, cells in _rows(workbook, BOUNDARIES):
        where = f"{BOUNDARIES} row {number}"
        kind = _text(cells.get("kind"), where, "kind").casefold().replace(" ", "_")
        if kind not in {"boundary", "not_used"}:
            raise InvalidKnowledgeError(f"{where}: kind must be boundary or not_used.")
        bucket = "boundaries" if kind == "boundary" else "not_used"
        offering(cells, where)[bucket].append(_text(cells.get("text"), where, "text"))
    return list(offerings.values())


def _sheet_conflicts(workbook: Any) -> list[dict[str, Any]]:
    """The conflict sheets in the YAML/JSON shape, each entry knowing its row."""
    conflicts: dict[str, dict[str, Any]] = {}
    for number, cells in _rows(workbook, CONFLICTS):
        where = f"{CONFLICTS} row {number}"
        conflict_id = _text(cells.get("conflict_id"), where, "conflict_id")
        if conflict_id in conflicts:
            raise InvalidKnowledgeError(f"{where}: conflict {conflict_id!r} is listed twice.")
        conflicts[conflict_id] = {
            **cells,
            "_where": where,
            "id": conflict_id,
            "a": {
                "source": cells.get("a_source_id"),
                "reference": cells.get("a_reference"),
                "statement": cells.get("a_statement"),
            },
            "b": {
                "source": cells.get("b_source_id"),
                "reference": cells.get("b_reference"),
                "statement": cells.get("b_statement"),
            },
            "scope": [],
        }
    for number, cells in _rows(workbook, CONFLICT_SCOPES):
        where = f"{CONFLICT_SCOPES} row {number}"
        conflict_id = _text(cells.get("conflict_id"), where, "conflict_id")
        if conflict_id not in conflicts:
            raise InvalidKnowledgeError(
                f"{where}: conflict {conflict_id!r} is not listed on the {CONFLICTS} sheet."
            )
        conflicts[conflict_id]["scope"].append(
            {
                "product": cells.get("product_id"),
                "order_types": list(_split(cells.get("order_types"), where, "order_types")),
                "question": cells.get("question_id"),
            }
        )
    return list(conflicts.values())


def _sheet_domains[DomainT: (CapabilityDomain, LandscapeDomain)](
    workbook: Any, sheet: str, kind: type[DomainT]
) -> list[DomainT]:
    domains = []
    for number, cells in _rows(workbook, sheet):
        where = f"{sheet} row {number}"
        try:
            domains.append(
                kind(
                    _text(cells.get("domain_id"), where, "domain_id"),
                    _text(cells.get("name"), where, "name"),
                    _optional_text(cells.get("name_ar"), where, "name_ar"),
                    _optional_text(cells.get("parent_id"), where, "parent_id"),
                    _optional_text(cells.get("description"), where, "description"),
                )
            )
        except InvalidKnowledgeError as exc:
            raise _located(where, exc) from exc
    return domains


def _located(where: str, exc: InvalidKnowledgeError) -> InvalidKnowledgeError:
    message = str(exc)
    return exc if message.startswith(where) else InvalidKnowledgeError(f"{where}: {message}")


def _append(sheet: Any, values: Sequence[object], *, header: bool = False) -> None:
    sheet.append(list(values))
    if header:
        for cell in sheet[sheet.max_row]:
            cell.font = Font(bold=True)


def _yes(value: bool | None) -> str | None:
    return None if value is None else "yes" if value else "no"


def _confidence(item: Any) -> str | None:
    return item.confidence.value if item.confidence else None


def _write_offering(sheets: dict[str, Any], product: ProductOffering) -> None:
    _append(
        sheets[PRODUCTS],
        (
            product.id,
            product.name,
            product.code,
            product.family,
            product.version,
            product.lifecycle,
            product.proposition,
            f"{_LIST_SEPARATOR} ".join(product.rules),
            _confidence(product),
            product.source,
            f"{_LIST_SEPARATOR} ".join(product.sources),
            product.primary_source,
            product.portfolio_node_id,
        ),
    )
    for plan in product.plans:
        _append(
            sheets[PLANS],
            (product.id, plan.name, plan.description, _confidence(plan), plan.source),
        )
        for fact in plan.characteristics:
            _append(sheets[PLAN_CHARACTERISTICS], (product.id, plan.name, fact.name, fact.value))
    for rule in product.business_rules:
        _append(
            sheets[BUSINESS_RULES],
            (product.id, rule.id, rule.statement, rule.kind, _confidence(rule), rule.source),
        )
    for question in product.questions:
        _append(
            sheets[QUESTIONS],
            (
                product.id,
                question.id,
                question.text,
                question.impact,
                _confidence(question),
                question.source,
                f"{_LIST_SEPARATOR} ".join(question.order_types),
            ),
        )
    for decision in product.decisions:
        _append(
            sheets[DECISIONS],
            (
                product.id,
                decision.id,
                decision.title,
                decision.text,
                _confidence(decision),
                decision.source,
            ),
        )
    for kind, texts in (("boundary", product.boundaries), ("not_used", product.not_used)):
        for text in texts:
            _append(sheets[BOUNDARIES], (product.id, kind, text))
    for order in product.order_types:
        _append(
            sheets[ORDER_TYPES],
            (
                product.id,
                order.code,
                order.name,
                _yes(order.enabled),
                order.description,
                _confidence(order),
                order.source,
                f"{_LIST_SEPARATOR} ".join(order.channels),
            ),
        )
    for part in product.components:
        _append(
            sheets[OFFERING_COMPONENTS],
            (
                product.id,
                part.id,
                part.name,
                part.code,
                part.kind,
                _yes(part.mandatory),
                _yes(part.customer_visible),
                part.description,
                part.commercial_spec,
                part.technical_spec,
                part.technical_details,
                _confidence(part),
                part.source,
            ),
        )
        for duty in part.responsibilities:
            _append(
                sheets[RESPONSIBILITIES],
                (
                    product.id,
                    part.id,
                    duty.system_id,
                    duty.role,
                    duty.description,
                    f"{_LIST_SEPARATOR} ".join(duty.order_types),
                    _confidence(duty),
                    duty.source,
                ),
            )
        for item in part.realisation:
            _append(
                sheets[REALISATION],
                (product.id, part.id, item.layer.value, item.name, _confidence(item), item.source),
            )
    if product.tracking is not None:
        _write_tracking(sheets, product.id, product.tracking)
    for note in product.lifecycle_notes:
        _write_note(sheets, product.id, note)
    for nfr in product.nfrs:
        _append(
            sheets[NFRS],
            (
                product.id,
                nfr.quality,
                nfr.coverage.value,
                nfr.statement,
                _confidence(nfr),
                nfr.source,
            ),
        )
    for kind, points in (("value", product.values), ("audience", product.audiences)):
        for point in points:
            _append(
                sheets[PRODUCT_POINTS],
                (product.id, kind, point.name, point.description, _confidence(point), point.source),
            )


def _write_note(sheets: dict[str, Any], product_id: str, note: LifecycleNote) -> None:
    joined = f"{_LIST_SEPARATOR} ".join
    _append(
        sheets[LIFECYCLE_NOTES],
        (
            product_id,
            note.id,
            note.title,
            note.kind,
            note.summary,
            joined(note.order_types),
            joined(note.channels),
            _confidence(note),
            note.source,
        ),
    )
    blank = (None,) * len(_CELLS)

    def row(kind: str, title: str | None, text: str | None, cells: tuple[object, ...]) -> None:
        padded = (*cells, *blank)[: len(_CELLS)]
        _append(
            sheets[LIFECYCLE_BLOCKS],
            (product_id, note.id, kind, title, text, *padded, None, None, None),
        )

    for block in note.blocks:
        if block.kind is NoteBlockKind.TEXT:
            cells: tuple[object, ...] = ()
            text = block.text
        elif block.kind is NoteBlockKind.LIST:
            cells, text = (), None
        else:
            cells, text = block.columns, block.caption
        _append(
            sheets[LIFECYCLE_BLOCKS],
            (
                product_id,
                note.id,
                block.kind.value,
                block.title,
                text,
                *(*cells, *blank)[: len(_CELLS)],
                _confidence(block),
                block.source,
                "yes" if block.to_verify else None,
            ),
        )
        for item in block.items:
            row("item", None, item, ())
        for table_row in block.rows:
            row("row", None, None, table_row)


def _write_tracking(sheets: dict[str, Any], product_id: str, tracking: OrderTracking) -> None:
    _append(
        sheets[TRACKING],
        (
            product_id,
            f"{_LIST_SEPARATOR} ".join(tracking.order_types),
            tracking.scope_note,
            tracking.not_applicable_note,
            _confidence(tracking),
            tracking.source,
        ),
    )
    for flow in tracking.flows:
        _append(
            sheets[TRACKING_FLOWS],
            (
                product_id,
                flow.from_system_id,
                flow.to_system_id,
                flow.label,
                flow.interface,
                _confidence(flow),
                flow.source,
            ),
        )
    for channel in tracking.channels:
        _append(
            sheets[TRACKING_CHANNELS],
            (
                product_id,
                channel.channel_id,
                channel.correlation_key,
                channel.ui_system_id,
                channel.story,
                channel.read_system_id,
                channel.read_interface,
                channel.ui_note,
                _confidence(channel),
                channel.source,
            ),
        )
    for kind, events in (("milestone", tracking.milestones), ("status", tracking.statuses)):
        for event in events:
            _append(
                sheets[TRACKING_EVENTS],
                (
                    product_id,
                    kind,
                    event.label,
                    event.detail,
                    event.system_id,
                    _confidence(event),
                    event.source,
                ),
            )
    for case in tracking.fallout:
        _append(
            sheets[TRACKING_EVENTS],
            (
                product_id,
                "fallout",
                case.trigger,
                case.handling,
                None,
                _confidence(case),
                case.source,
            ),
        )


def _write_journey(sheets: dict[str, Any], journey: Journey) -> None:
    joined = f"{_LIST_SEPARATOR} ".join
    _append(
        sheets[JOURNEYS],
        (
            journey.id,
            journey.name,
            journey.product_id,
            journey.order_type_code,
            journey.description,
            _confidence(journey),
            journey.source,
        ),
    )
    for step in journey.activities:
        _append(
            sheets[ACTIVITIES],
            (
                journey.id,
                step.number,
                step.name,
                step.phase,
                step.track,
                step.performing_system_id,
                joined(step.supporting_system_ids),
                step.system_function,
                step.mode,
                _yes(step.customer_visible),
                step.description,
                joined(step.component_ids),
                step.input,
                step.output,
                step.etom,
                _confidence(step),
                step.source,
                joined(step.channels),
                _yes(step.channel_entry) if step.channel_entry else None,
                step.performer,
                step.point_of_no_return,
                step.role,
            ),
        )
    for rule in journey.flow_rules:
        _append(
            sheets[FLOW_RULES],
            (
                journey.id,
                rule.kind.value,
                rule.from_activity,
                rule.to_activity,
                rule.condition,
                rule.branch,
                rule.parallel_group,
                rule.rejoin_at,
                _confidence(rule),
                rule.source,
            ),
        )
    for link in journey.integrations:
        _append(
            sheets[ACTIVITY_INTEGRATIONS],
            (
                journey.id,
                link.from_activity,
                link.to_activity,
                link.interaction,
                link.interface,
                link.payload,
                link.timing,
                link.correlation_key,
                _confidence(link),
                link.source,
                link.from_system_id,
                link.to_system_id,
                link.via_system_id,
                link.purpose,
                link.style,
                link.tmf_equivalent,
            ),
        )


def _workbook(release: ArchitectureKnowledge | None) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.worksheets[0])
    instructions = workbook.create_sheet(INSTRUCTIONS)
    for line in _INSTRUCTIONS:
        instructions.append(list(line))
    instructions["A1"].font = Font(bold=True)
    instructions.column_dimensions["A"].width = 90
    sheets = {name: workbook.create_sheet(name) for name in _HEADERS}
    for name, sheet in sheets.items():
        _append(sheet, _HEADERS[name], header=True)
        sheet.freeze_panes = "A2"
        for column in "ABCDEFGHIJKLMNOPQRS":
            sheet.column_dimensions[column].width = 32
    systems = release.systems if release is not None else ()
    for system in systems:
        _append(
            sheets[SYSTEMS],
            (
                system.id,
                system.name,
                system.name_ar,
                f"{_LIST_SEPARATOR} ".join(system.aliases),
                system.description,
                system.landscape_domain_id,
                system.owner,
                _yes(system.external) if system.external else None,
                system.roadmap,
                system.placement_from,
                system.placement_reason,
                _confidence(system),
                system.source,
            ),
        )
        for capability in system.capabilities:
            _append(
                sheets[CAPABILITIES],
                (
                    system.id,
                    capability.id,
                    capability.name,
                    f"{_LIST_SEPARATOR} ".join(capability.triggers),
                    capability.domain_id,
                    capability.component_id,
                ),
            )
        for component in system.components:
            _append(
                sheets[COMPONENTS],
                (
                    system.id,
                    component.id,
                    component.name,
                    component.name_ar,
                    f"{_LIST_SEPARATOR} ".join(component.aliases),
                    component.technology,
                    component.description,
                ),
            )
        for constraint in system.constraints:
            _append(sheets[CONSTRAINTS], (system.id, constraint))
    for relationship in release.relationships if release is not None else ():
        _append(
            sheets[RELATIONSHIPS],
            (
                relationship.source_system_id,
                relationship.target_system_id,
                relationship.description,
                relationship.kind.value,
            ),
        )
    for sheet, tree in (
        (DOMAINS, release.capability_domains if release is not None else ()),
        (LANDSCAPE, release.landscape_domains if release is not None else ()),
    ):
        for domain in tree:
            _append(
                sheets[sheet],
                (domain.id, domain.name, domain.name_ar, domain.parent_id, domain.description),
            )
    for node in release.portfolio if release is not None else ():
        _append(
            sheets[PORTFOLIO],
            (
                node.id,
                node.name,
                node.level,
                node.parent_id,
                node.description,
                _confidence(node),
                node.source,
            ),
        )
    for product in release.products if release is not None else ():
        _write_offering(sheets, product)
    for journey in release.journeys if release is not None else ():
        _write_journey(sheets, journey)
    for channel in release.channels if release is not None else ():
        _append(
            sheets[CHANNELS],
            (
                channel.id,
                channel.name,
                channel.kind,
                channel.entry_system_id,
                channel.description,
                _confidence(channel),
                channel.source,
            ),
        )
    for source in release.sources if release is not None else ():
        _append(
            sheets[SOURCES],
            (
                source.id,
                source.level.value,
                source.title,
                source.short,
                source.version,
                source.file,
                _yes(source.supplied),
                source.authority,
                source.scope,
                source.boundary,
            ),
        )
    for conflict in release.conflicts if release is not None else ():
        _append(
            sheets[CONFLICTS],
            (
                conflict.id,
                conflict.title,
                conflict.a.source_id,
                conflict.a.reference,
                conflict.a.statement,
                conflict.b.source_id,
                conflict.b.reference,
                conflict.b.statement,
                conflict.difference,
                conflict.impact,
                conflict.decision,
                _confidence(conflict),
                conflict.source,
            ),
        )
        for scope in conflict.scope:
            _append(
                sheets[CONFLICT_SCOPES],
                (
                    conflict.id,
                    scope.product_id,
                    f"{_LIST_SEPARATOR} ".join(scope.order_types),
                    scope.question_id,
                ),
            )
    for record in release.change_history if release is not None else ():
        trace = record.trace
        _append(
            sheets[CHANGE_HISTORY],
            (
                record.id,
                record.title,
                record.origin.value,
                record.product_id,
                record.requester,
                record.reason,
                record.priority,
                record.target_date,
                _iso(record.created_at),
                _iso(record.applied_at),
                f"{_LIST_SEPARATOR} ".join(record.gaps),
                trace.requirement_id if trace else None,
                trace.breakdown_revision if trace else None,
                trace.approval_id if trace else None,
                trace.approved_by if trace else None,
                _iso(trace.approved_at) if trace else None,
                trace.epic_id if trace else None,
                trace.epic_name if trace else None,
                f"{_LIST_SEPARATOR} ".join(f"{item.id}: {item.name}" for item in trace.features)
                if trace
                else None,
                trace.export_schema if trace else None,
                trace.knowledge_version if trace else None,
            ),
        )
        for item in record.items:
            _append(
                sheets[CHANGE_ITEMS],
                (record.id, item.kind, item.summary, item.status.value, item.feature_id),
            )
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _sheet_history(workbook: Any) -> list[dict[str, Any]]:
    """The change history sheets gathered into the YAML/JSON shape."""
    records: dict[str, dict[str, Any]] = {}
    for number, cells in _rows(workbook, CHANGE_HISTORY):
        where = f"{CHANGE_HISTORY} row {number}"
        record_id = _text(cells.get("change_request_id"), where, "change_request_id")
        if record_id in records:
            raise InvalidKnowledgeError(f"{where}: change request {record_id!r} is listed twice.")
        features = [
            dict(zip(("id", "name"), (part.strip() for part in item.split(":", 1)), strict=False))
            for item in _split(cells.get("features"), where, "features")
        ]
        trace = (
            {
                "requirement": cells.get("requirement_id"),
                "revision": cells.get("breakdown_revision"),
                "approval": cells.get("approval_id"),
                "approved_by": cells.get("approved_by"),
                "approved_at": cells.get("approved_at"),
                "epic": cells.get("epic_id"),
                "epic_name": cells.get("epic_name"),
                "features": features,
                "export_schema": cells.get("export_schema"),
                "knowledge_version": cells.get("knowledge_version"),
            }
            if cells.get("requirement_id")
            else None
        )
        records[record_id] = {
            **cells,
            "_where": where,
            "id": record_id,
            "product": cells.get("product_id"),
            "gaps": list(_split(cells.get("gaps"), where, "gaps")),
            "trace": trace,
            "items": [],
        }
    for number, cells in _rows(workbook, CHANGE_ITEMS):
        where = f"{CHANGE_ITEMS} row {number}"
        record_id = _text(cells.get("change_request_id"), where, "change_request_id")
        if record_id not in records:
            raise InvalidKnowledgeError(
                f"{where}: change request {record_id!r} is not on {CHANGE_HISTORY}."
            )
        records[record_id]["items"].append(
            {**cells, "feature": cells.get("feature_id"), "_where": where}
        )
    return list(records.values())


def _iso(value: datetime | None) -> str | None:
    return value.isoformat() if value is not None else None


def _moment(value: object, where: str, field: str) -> datetime | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value if value.tzinfo else value.replace(tzinfo=UTC)
    text = _text(value, where, field)
    try:
        moment = datetime.fromisoformat(text.replace("Z", "+00:00"))
    except ValueError as exc:
        raise InvalidKnowledgeError(f"{where}: {field} must be a date and time.") from exc
    return moment if moment.tzinfo else moment.replace(tzinfo=UTC)


# Change history (requirement-portal ADR-0101, step 7): the change requests applied, each
# with the approved requirement it comes from.


def _history(entries: list[dict[str, Any]]) -> list[ChangeRequestRecord]:
    return [
        _part(where, partial(_record, item, where))
        for number, item in enumerate(entries, start=1)
        for where in (_where(item, f"change_history entry {number}"),)
    ]


def _record(item: dict[str, Any], where: str) -> ChangeRequestRecord:
    raw_trace = item.get("trace")
    trace = None
    if isinstance(raw_trace, dict):
        place = f"{where}, trace"
        revision = raw_trace.get("revision")
        trace = RequirementTrace(
            requirement_id=_text(raw_trace.get("requirement"), place, "requirement"),
            breakdown_revision=int(str(revision)) if revision not in (None, "") else 1,
            approval_id=_text(raw_trace.get("approval"), place, "approval"),
            epic_id=_text(raw_trace.get("epic"), place, "epic"),
            epic_name=_text(raw_trace.get("epic_name"), place, "epic_name"),
            approved_by=_optional_text(raw_trace.get("approved_by"), place, "approved_by"),
            approved_at=_moment(raw_trace.get("approved_at"), place, "approved_at"),
            features=tuple(
                TracedFeature(
                    _text(feature.get("id"), place, "feature id"),
                    _text(feature.get("name"), place, "feature name"),
                )
                for feature in raw_trace.get("features") or []
                if isinstance(feature, dict)
            ),
            export_schema=_optional_text(raw_trace.get("export_schema"), place, "export_schema"),
            knowledge_version=_optional_text(
                raw_trace.get("knowledge_version"), place, "knowledge_version"
            ),
        )
    items = []
    for position, raw in enumerate(item.get("items") or [], start=1):
        place = _where(raw, f"{where}, item {position}") if isinstance(raw, dict) else where
        if not isinstance(raw, dict):
            raise InvalidKnowledgeError(f"{place}: an item must be an object.")
        status = _optional_text(raw.get("status"), place, "status") or "recorded"
        try:
            known = ChangeItemStatus(status.casefold())
        except ValueError as exc:
            raise InvalidKnowledgeError(
                f"{place}: status must be recorded, inferred, gap or conflict."
            ) from exc
        items.append(
            ChangeItem(
                _text(raw.get("kind"), place, "kind"),
                _text(raw.get("summary"), place, "summary"),
                known,
                _optional_text(raw.get("feature"), place, "feature"),
            )
        )
    origin = _optional_text(item.get("origin"), where, "origin") or "requirement-ai"
    try:
        source = ChangeOrigin(origin.casefold())
    except ValueError as exc:
        raise InvalidKnowledgeError(f"{where}: origin must be requirement-ai or explorer.") from exc
    return ChangeRequestRecord(
        id=_text(item.get("id"), where, "id"),
        title=_text(item.get("title"), where, "title"),
        origin=source,
        product_id=_optional_text(item.get("product"), where, "product"),
        requester=_optional_text(item.get("requester"), where, "requester"),
        reason=_optional_text(item.get("reason"), where, "reason"),
        priority=_optional_text(item.get("priority"), where, "priority"),
        target_date=_optional_text(item.get("target_date"), where, "target_date"),
        created_at=_moment(item.get("created_at"), where, "created_at"),
        applied_at=_moment(item.get("applied_at"), where, "applied_at"),
        trace=trace,
        items=tuple(items),
        gaps=_text_list(item.get("gaps"), where, "gaps"),
    )


def _history_mapping(record: ChangeRequestRecord) -> dict[str, Any]:
    trace = record.trace
    return {
        "id": record.id,
        "title": record.title,
        "origin": record.origin.value,
        **_present(
            product=record.product_id,
            requester=record.requester,
            reason=record.reason,
            priority=record.priority,
            target_date=record.target_date,
            created_at=_iso(record.created_at),
            applied_at=_iso(record.applied_at),
        ),
        **(
            {
                "trace": {
                    "requirement": trace.requirement_id,
                    "revision": trace.breakdown_revision,
                    "approval": trace.approval_id,
                    "epic": trace.epic_id,
                    "epic_name": trace.epic_name,
                    **_present(
                        approved_by=trace.approved_by,
                        approved_at=_iso(trace.approved_at),
                        features=[{"id": item.id, "name": item.name} for item in trace.features],
                        export_schema=trace.export_schema,
                        knowledge_version=trace.knowledge_version,
                    ),
                }
            }
            if trace
            else {}
        ),
        **_present(
            items=[
                {
                    "kind": item.kind,
                    "summary": item.summary,
                    "status": item.status.value,
                    **_present(feature=item.feature_id),
                }
                for item in record.items
            ],
            gaps=list(record.gaps),
        ),
    }
